"""Rodada 05 — MIP v04 (versão final: D17 + D18), variante E4 (T14), D18 (T16) e material do artigo (T13, T15).

Os testes com a especificação real são pulados se os parâmetros v04 ainda não tiverem sido gerados
(scripts/rodada05_parametros.py); os das simulações e do artigo, se as saídas ainda não existirem.
"""
from dataclasses import replace

import numpy as np
import pandas as pd
import pytest

from mip.caminhos import PASTA_EXTERNOS, PASTA_INTERMEDIARIOS, PASTA_PARAMETROS
from mip.desagregacao import (
    EspecificacaoInvalida,
    agregar,
    desagregar,
    ler_especificacao,
    realocar_emissoes_original,
)
from mip.estrutura import ITENS_SATELITE
from mip.leitura import ler_mip_original
from mip.modelo import ModeloLeontief

TOL = 1e-6
ARRAYS = ["Z", "F", "X", "total_linha", "V", "V_df", "X_df"]
VARIANTES = ["", "_P2", "_E2", "_E3", "_E4", "_gas4dist", "_rap_media", "_k060", "_eta005"]
K_CO2 = ITENS_SATELITE.index("co2e")
GER_RS = ["S38_HID", "S38_EOL", "S38_BIO", "S38_CAR", "S38_GN"]


@pytest.fixture(scope="module")
def mip():
    return ler_mip_original()


def _espec(suf=""):
    caminho = PASTA_PARAMETROS / f"especificacao_v04{suf}.yaml"
    if not caminho.exists():
        pytest.skip(f"{caminho.name} ainda não foi gerado (scripts/rodada05_parametros.py)")
    return ler_especificacao(caminho)


@pytest.fixture(scope="module")
def v04(mip):
    return desagregar(mip, _espec(""))


@pytest.fixture(scope="module")
def v04_e4(mip):
    return desagregar(mip, _espec("_E4"), estrito=False)


@pytest.mark.parametrize("suf", VARIANTES)
def test_t2_reagregacao_v04(mip, suf):
    """T2 nas 9 especificações; na E3 e na E4 o CO₂e só vale pelo total (E3: níveis; E4: total da economia)."""
    res = desagregar(mip, _espec(suf), estrito=False)
    ag = agregar(res.mip, res.mapa, mip.setores)
    for nome in ARRAYS:
        np.testing.assert_allclose(getattr(ag, nome), getattr(mip, nome), rtol=0, atol=TOL, err_msg=f"{suf}: {nome}")
    outros = [k for k, it in enumerate(ITENS_SATELITE) if it != "co2e" or suf not in ("_E3", "_E4")]
    np.testing.assert_allclose(ag.S[outros], mip.S[outros], rtol=0, atol=TOL)
    if suf == "_E4":
        assert ag.S[K_CO2].sum() == pytest.approx(mip.S[K_CO2].sum(), rel=1e-12)
    assert not (res.checagens["status"] == "FALHA").any()
    assert res.especificacao.piso_vbp.get("viabilidade_eta") is None


def test_d17_transmissao_ons(mip, v04):
    d = pd.read_csv(PASTA_EXTERNOS / "valores_documentais_v03.csv").set_index("parametro")["valor"]
    ons = d["ons_encargos_uso_ano_civil_2019"]
    assert ons == pytest.approx(22935.6, abs=0.05)
    xs = {q: mip.X[mip.indice(q, "S38")] for q in ("RS", "RB")}
    tra = {q: v04.mip.X[v04.mip.indice(q, "S38_TRA")] for q in ("RS", "RB")}
    for q in ("RS", "RB"):
        assert tra[q] == pytest.approx(ons * xs[q] / sum(xs.values()), rel=1e-9)
    assert sum(tra.values()) == pytest.approx(ons, rel=1e-12)
    media = desagregar(mip, _espec("_rap_media"), estrito=False)
    assert sum(media.mip.X[media.mip.indice(q, "S38_TRA")] for q in ("RS", "RB")) == pytest.approx(24883.43, abs=0.01)
    comp = pd.read_csv(PASTA_PARAMETROS / "componentes_v04.csv")
    f = comp.loc[comp["componente"] == "vbp_transmissao", "fonte"]
    assert f.str.startswith("D17: encargos de uso do sistema de transmissão faturados pelo ONS").all()


def test_t16_d18_tust_latentes(v04):
    """T16: a TUST por MWh das colunas latentes, ao p_k do espelho, é igual a τ_RS (TUST por MWh das fontes do RS)."""
    m = v04.mip
    tra = [m.indice(q, "S38_TRA") for q in ("RS", "RB")]
    pesos = v04.especificacao.pesos_tust_geracao
    pr = pesos[pesos["regiao"] == "RS"].set_index("subsetor")["valor"].astype(float)
    tust = sum(m.Z[tra, m.indice("RS", s)].sum() for s in GER_RS)
    tau = tust / sum(pr[s] for s in GER_RS)  # R$ mi por GWh
    d8 = pd.read_csv(PASTA_INTERMEDIARIOS / "geracao_2019_regiao_subsetor_v03.csv").set_index("regiao")
    cl = m.coef_latentes
    for s in ("S38_SOL", "S38_PET"):
        p = m.X[m.indice("RB", s)] / float(d8.loc["RB", s])
        sel = (cl["regiao"] == "RS") & (cl["subsetor"] == s) & (cl["tipo_linha"] == "setor") & (cl["codigo_linha"] == "S38_TRA")
        c = cl.loc[sel, "coeficiente"].sum()
        assert c * p == pytest.approx(tau, rel=1e-9), s
        assert cl.loc[sel, "fonte"].str.contains("HIPÓTESE (D18)", regex=False).all()
    t = v04.extras["tust_latentes"].set_index("subsetor")
    assert t["tau (R$/MWh)"].iloc[0] == pytest.approx(tau * 1000, rel=1e-9)
    assert (t["dispersao_relativa_tau_entre_fontes"] < 1e-9).all()  # todas as fontes do RS pagam o mesmo τ (V1)


def test_d18_mantem_reparticao_e_demais_coeficientes(mip, v04):
    sem = desagregar(mip, replace(v04.especificacao, tust_latentes={}), estrito=False)
    a = sem.mip.coef_latentes.set_index(["regiao", "subsetor", "tipo_linha", "regiao_linha", "codigo_linha"])["coeficiente"]
    b = v04.mip.coef_latentes.set_index(["regiao", "subsetor", "tipo_linha", "regiao_linha", "codigo_linha"])["coeficiente"]
    a, b = a.sort_index(), b.sort_index()
    tra = b.index.get_level_values("codigo_linha") == "S38_TRA"
    np.testing.assert_allclose(b[~tra].to_numpy(), a.reindex(b.index)[~tra].to_numpy(), rtol=1e-12, atol=0)
    for s in ("S38_SOL", "S38_PET"):
        x, y = a.loc[("RS", s, "setor")], b.loc[("RS", s, "setor")]
        xt = x[x.index.get_level_values("codigo_linha") == "S38_TRA"]
        yt = y[y.index.get_level_values("codigo_linha") == "S38_TRA"]
        np.testing.assert_allclose((yt / yt.sum()).to_numpy(), (xt / xt.sum()).reindex(yt.index).to_numpy(), rtol=1e-12)
    # a MIP monetária não muda (só o modelo, pelas colunas latentes)
    np.testing.assert_allclose(sem.mip.Z, v04.mip.Z, rtol=1e-12, atol=1e-9)


def test_d18_validacao(mip):
    e = _espec("")
    with pytest.raises(EspecificacaoInvalida, match="falta o GWh do espelho"):
        desagregar(mip, replace(e, tust_latentes={"gwh_espelho": {"S38_SOL": 4978.0}}), estrito=False)
    with pytest.raises(EspecificacaoInvalida, match="precisa ser > 0"):
        desagregar(mip, replace(e, tust_latentes={"gwh_espelho": {"S38_SOL": 4978.0, "S38_PET": 0}}), estrito=False)


def test_t14_e4(mip, v04, v04_e4):
    """T14: valores movidos exatos; total da economia inalterado; S38 reduzido no valor movido; nada negativo."""
    ra = v04_e4.especificacao.realocacao_emissoes.assign(valor=lambda d: d["valor"].astype(float))
    ref = {("RS", "S19"): 1812.1, ("RS", "S04"): 88.1, ("RB", "S19"): 18088.8, ("RB", "S05"): 23499.2,
           ("RB", "S04"): 1155.4}  # Seção 3.2 das instruções (mil t)
    assert len(ra) == len(ref)
    for x in ra.itertuples(index=False):
        assert x.valor / 1000 == pytest.approx(ref[(x.regiao, x.setor_destino)], abs=0.1)  # soma de parcelas arredondadas
    a, b = v04.mip.S[K_CO2], v04_e4.mip.S[K_CO2]
    m = v04.mip
    for x in ra.itertuples(index=False):
        assert b[m.indice(x.regiao, x.setor_destino)] - a[m.indice(x.regiao, x.setor_destino)] == pytest.approx(x.valor, abs=1e-6)
    assert b.sum() == pytest.approx(a.sum(), rel=1e-13)
    for r in ("RS", "RB"):
        idx = [m.indice(r, s) for s in m.setores.loc[m.setores["regiao"] == r, "codigo"] if s.startswith("S38_")]
        movido = ra.loc[ra["regiao"] == r, "valor"].sum()
        assert a[idx].sum() - b[idx].sum() == pytest.approx(movido, abs=1e-6)
        # sai só de DIS e GAS, na proporção do resíduo (= do VBP)
        dis, gas = m.indice(r, "S38_DIS"), m.indice(r, "S38_GAS")
        d_dis, d_gas = a[dis] - b[dis], a[gas] - b[gas]
        assert d_dis + d_gas == pytest.approx(movido, abs=1e-6)
        assert d_dis / d_gas == pytest.approx(m.X[dis] / m.X[gas], rel=1e-9)
    assert (b >= -1e-9).all()
    np.testing.assert_array_equal(v04.mip.Z, v04_e4.mip.Z)  # a MIP monetária é a da versão final
    c = v04_e4.checagens.set_index("checagem")["status"]
    for nome in ("valores_movidos", "setor_alvo", "total_economia", "sem_negativos"):
        assert c[f"realocacao_co2e_{nome}"] == "OK"
    assert c["satelite_co2e_preservado"] == "ALERTA"


def test_t14_a0_e4(mip, v04_e4):
    ra = v04_e4.especificacao.realocacao_emissoes.assign(valor=lambda d: d["valor"].astype(float))
    m0 = realocar_emissoes_original(mip, ra, "S38")
    for r in ("RS", "RB"):
        movido = ra.loc[ra["regiao"] == r, "valor"].sum()
        assert mip.S[K_CO2, mip.indice(r, "S38")] - m0.S[K_CO2, mip.indice(r, "S38")] == pytest.approx(movido, abs=1e-6)
    assert m0.S[K_CO2].sum() == pytest.approx(mip.S[K_CO2].sum(), rel=1e-13)
    assert (m0.S[K_CO2] >= 0).all()
    np.testing.assert_array_equal(m0.Z, mip.Z)
    assert mip.S[K_CO2].sum() == pytest.approx(ler_mip_original().S[K_CO2].sum())  # a original não foi alterada


def test_e4_validacao(mip):
    e = _espec("_E4")
    ra = e.realocacao_emissoes.copy()
    ra.loc[0, "setor_destino"] = "S38"
    with pytest.raises(EspecificacaoInvalida, match="destino"):
        desagregar(mip, replace(e, realocacao_emissoes=ra), estrito=False)
    ra = e.realocacao_emissoes.copy()
    ra["valor"] = "1e12"
    with pytest.raises(EspecificacaoInvalida, match="passa do resíduo"):
        desagregar(mip, replace(e, realocacao_emissoes=ra), estrito=False)


def test_latentes_entram_no_modelo(v04):
    mod = ModeloLeontief(v04.mip)
    j = v04.mip.indice("RS", "S38_SOL")
    assert mod.A[:, j].sum() > 0 and mod.L[j, j] >= 1


# ---------------------------------------------------------------------------------------------
# material do artigo (T13, T15)
# ---------------------------------------------------------------------------------------------
def _artigo():
    import importlib.util

    from mip.caminhos import RAIZ

    caminho = RAIZ / "scripts" / "rodada05_artigo.py"
    spec = importlib.util.spec_from_file_location("rodada05_artigo", caminho)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    if not mod.SIM.exists():
        pytest.skip("simulacoes_v04.xlsx ainda não foi gerado (scripts/rodada05_simulacoes.py)")
    return mod


@pytest.fixture(scope="module")
def artigo():
    mod = _artigo()
    return mod, {t.numero: t for t in mod.construir_tabelas(mod.carregar())}


def _painel(t, nome):
    return next(df for n, df in t.paineis if n == nome).reset_index(drop=True)


def test_t13_painel_b_reproduz_o_artigo(artigo):
    """T13: os Painéis B das Tabelas 1, 3 e 4 e a Tabela A2 reproduzem os valores publicados (uma casa decimal)."""
    mod, tabs = artigo
    t1 = _painel(tabs["Tabela 1"], mod.PAINEL_B)
    esperado1 = [(1925.0, 428113.5, 222.4, 4.5), (1959.4, 216407.0, 110.4, 9.1), (2120.6, 356045.7, 167.9, 6.0)]
    for i, linha in enumerate(esperado1):
        assert tuple(round(float(v), 1) for v in t1.iloc[i, 1:5]) == linha, t1.iloc[i, 0]
    assert round(float(t1.iloc[3]["Emissões totais (t CO₂e)"]), 1) == 139638.8
    assert round(float(t1.iloc[4]["Intensidade carbônica (t CO₂e/R$ milhão)"]), 1) == 34.2
    t3 = _painel(tabs["Tabela 3"], mod.PAINEL_B)
    esperado3 = [(1357.0, 339326.4, 79.3), (568.0, 88787.1, 20.7), (1391.3, 142390.0, 65.8), (568.0, 74016.9, 34.2),
                 (1373.2, 277141.4, 77.8), (747.3, 78904.3, 22.2)]
    for i, linha in enumerate(esperado3):
        assert tuple(round(float(v), 1) for v in t3.iloc[i, 2:5]) == linha, t3.iloc[i, :2].tolist()
    t4 = _painel(tabs["Tabela 4"], mod.PAINEL_B)
    esperado4 = [(1959.4, 216407.0, 110.4, 51.5, 34.2), (1920.4, 216987.2, 113.0, 46.1, 29.1),
                 (1969.2, 169090.0, 85.9, 62.7, 43.5), (1959.5, 309794.9, 158.1, 39.0, 23.8),
                 (2120.6, 356045.7, 167.9, 35.3, 22.2), (2110.5, 381702.9, 180.9, 32.9, 19.7),
                 (2073.9, 375427.3, 181.0, 35.0, 21.5)]
    for i, linha in enumerate(esperado4):
        assert tuple(round(float(v), 1) for v in t4.iloc[i, 1:6]) == linha, t4.iloc[i, 0]
    a2 = tabs["Tabela A2"].paineis[0][1].reset_index(drop=True)
    cols = ["Produção induzida (R$ milhões)", "Emissões induzidas (t CO₂e)", "Participação (%)"]
    assert (a2.iloc[0]["Região"], a2.iloc[0]["Código"]) == ("RS", "S38")
    assert tuple(round(float(a2.iloc[0][c]), 1) for c in cols) == (357.1, 99554.6, 46.0)
    assert (a2.iloc[1]["Região"], a2.iloc[1]["Código"]) == ("RS", "S43")
    assert tuple(round(float(a2.iloc[1][c]), 1) for c in cols) == (84.1, 35489.9, 16.4)
    assert (a2.iloc[9]["Região"], a2.iloc[9]["Código"]) == ("RB", "S01")
    assert tuple(round(float(a2.iloc[9][c]), 1) for c in cols) == (3.5, 2649.5, 1.2)
    assert a2.iloc[10]["Setor"] == "Demais setores"
    assert tuple(round(float(a2.iloc[10][c]), 1) for c in cols[1:]) == (16285.6, 7.5)


def test_t15_artigo_deterministico(artigo):
    """T15: gerar o material duas vezes dá os mesmos números em todas as tabelas e nos números-chave."""
    mod, tabs = artigo
    d = mod.carregar()
    outra = {t.numero: t for t in mod.construir_tabelas(d)}
    assert set(outra) == set(tabs)
    for k, t in tabs.items():
        for (n1, a), (n2, b) in zip(t.paineis, outra[k].paineis):
            assert n1 == n2
            pd.testing.assert_frame_equal(a, b, check_exact=True)
    assert mod.numeros_chave(d) == mod.numeros_chave(mod.carregar())


def test_artigo_formato_br_e_saidas(artigo):
    mod, tabs = artigo
    assert mod.br(1925.04) == "1.925,0" and mod.br(-6.4441) == "−6,4" and mod.br(0.0) == "0,0"
    assert mod.br(139638.75, 1) == "139.638,8" and mod.pct(46.0034) == "46,0%"
    assert {"Quadro 1"} | {f"Tabela {i}" for i in range(1, 8)} | {f"Tabela A{i}" for i in range(1, 9)} <= set(tabs)
    if (mod.TAB / "tabelas_artigo.docx").exists():
        from docx import Document
        from docx.shared import Cm

        doc = Document(str(mod.TAB / "tabelas_artigo.docx"))
        s = doc.sections[0]
        assert abs(s.page_width - Cm(21.0)) < Cm(0.01) and abs(s.left_margin - Cm(2.5)) < Cm(0.01)
        run = doc.tables[0].cell(1, 0).paragraphs[0].runs[0]
        assert run.font.name == "Times New Roman" and run.font.size.pt == 8
    for nome in ("figura_1_top10_cesta_renovavel", "figura_2_diretas_indiretas", "figura_3_emissoes_por_TWh",
                 "figura_4_emissoes_por_cesta", "figura_A1_top10_cesta_renovavel_original"):
        if (mod.FIG / f"{nome}.png").exists():
            assert (mod.FIG / f"{nome}.pdf").exists()
