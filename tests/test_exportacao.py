"""Exportação: o .xlsx sai com o mesmo layout da original e é relido sem perda."""
import numpy as np
import openpyxl
import pytest

from mip.caminhos import CAMINHO_ORIGINAL, PASTA_PARAMETROS
from mip.desagregacao import desagregar, ler_especificacao
from mip.exportacao import exportar_longo, exportar_xlsx
from mip.leitura import ler_mip, ler_mip_original

ARRAYS = ["Z", "F", "X", "total_linha", "V", "S", "V_df", "X_df"]


@pytest.fixture(scope="module")
def mip():
    return ler_mip_original()


def test_exportar_original_reproduz_a_planilha_celula_a_celula(mip, tmp_path):
    destino = exportar_xlsx(mip, tmp_path / "copia.xlsx")
    a = openpyxl.load_workbook(CAMINHO_ORIGINAL).active
    b = openpyxl.load_workbook(destino).active
    assert (a.max_row, a.max_column) == (b.max_row, b.max_column)
    diferencas = []
    for r in range(1, a.max_row + 1):
        for c in range(1, a.max_column + 1):
            va, vb = a.cell(r, c).value, b.cell(r, c).value
            if isinstance(va, (int, float)) and isinstance(vb, (int, float)):
                if va != vb:
                    diferencas.append((r, c, va, vb))
            elif va != vb:
                diferencas.append((r, c, va, vb))
    assert diferencas == []
    assert {str(m) for m in a.merged_cells.ranges} == {str(m) for m in b.merged_cells.ranges}


def test_reler_exportacao_da_mip_desagregada(mip, tmp_path):
    res = desagregar(mip, ler_especificacao(PASTA_PARAMETROS / "TESTE_especificacao.yaml"))
    xlsx = tmp_path / "teste.xlsx"
    csv = tmp_path / "teste.csv"
    res.exportar(xlsx, csv)
    lida = ler_mip(xlsx)
    assert lida.n == res.mip.n == 136 + 2 * 7
    assert list(lida.setores["codigo"]) == list(res.mip.setores["codigo"])
    # o openpyxl grava os números com 16 algarismos significativos: diferença relativa ≤ ~1e-15
    for nome in ARRAYS:
        np.testing.assert_allclose(getattr(lida, nome), getattr(res.mip, nome), rtol=1e-14, atol=1e-12, err_msg=nome)
    assert lida.coef_latentes is not None and len(lida.coef_latentes) == len(res.mip.coef_latentes)
    wb = openpyxl.load_workbook(xlsx)
    assert wb.sheetnames == ["MIP SS RS_RBR", "LEIAME", "mapa_setores", "parametros", "checagens", "ras", "coef_latentes"]
    ws = wb["MIP SS RS_RBR"]
    assert ws["A1"].value.startswith("TESTE — VALORES FICTÍCIOS")
    assert ws["A3"].value == "TESTE — VALORES FICTÍCIOS"
    assert wb["LEIAME"]["A1"].value == "TESTE — VALORES FICTÍCIOS"
    # linhas de totais com fórmulas, deslocadas pelo número de setores
    n = res.mip.n
    assert ws.cell(7 + n, 5).value == f"=SUM(E7:E{6 + n})"
    assert ws.cell(7 + n + 16, 5).value == f"=E{7 + n + 14}-E{7 + n + 13}"


def test_exportar_longo(mip, tmp_path):
    import pandas as pd

    caminho = exportar_longo(mip, tmp_path / "longo.csv")
    df = pd.read_csv(caminho)
    assert set(df.columns) == {"bloco", "regiao_origem", "setor_origem", "regiao_destino",
                               "setor_ou_coluna_destino", "valor", "unidade"}
    assert df.loc[df["bloco"] == "Z", "valor"].sum() == pytest.approx(mip.Z.sum())
