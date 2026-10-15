"""Exportação: tabelas em .xlsx, MIP em .xlsx com o layout da original e MIP em CSV formato longo."""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .estrutura import (
    COLUNAS_DF,
    ITENS_PRIMARIOS,
    ITENS_PRIMARIOS_DF,
    ITENS_SATELITE,
    MIP,
    ROTULOS_LINHAS_INFERIORES,
    ROTULOS_REGIAO,
)
from .leitura import ABA_PRINCIPAL, formato_longo

NEGRITO = Font(bold=True)
DESTAQUE = PatternFill("solid", fgColor="FFF2CC")
ALERTA = PatternFill("solid", fgColor="F8CBAD")
FORMATO_NUM = "#,##0.00"


# ---------------------------------------------------------------------------------------------
# Tabelas genéricas
# ---------------------------------------------------------------------------------------------
def _ajustar_larguras(ws, df: pd.DataFrame, max_largura: int = 60):
    for k, col in enumerate(df.columns, start=1):
        valores = [str(col)] + [str(v) for v in df[col].head(200).tolist()]
        largura = min(max_largura, max(len(v) for v in valores) + 2)
        ws.column_dimensions[get_column_letter(k)].width = max(8, largura)


def salvar_tabelas(
    caminho: str | Path,
    abas: dict[str, pd.DataFrame],
    leiame: list[str] | None = None,
    destacar: dict[str, pd.Series] | None = None,
):
    """Salva DataFrames em abas de um .xlsx (primeira aba = LEIAME, se houver texto).

    `destacar` = {aba: série booleana} pinta as linhas marcadas (ex.: insumos energéticos).
    """
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(caminho, engine="openpyxl") as xw:
        if leiame:
            pd.DataFrame({"LEIAME": leiame}).to_excel(xw, sheet_name="LEIAME", index=False)
            xw.sheets["LEIAME"].column_dimensions["A"].width = 140
        for nome, df in abas.items():
            df.to_excel(xw, sheet_name=nome, index=False)
            ws = xw.sheets[nome]
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = NEGRITO
                c.alignment = Alignment(wrap_text=True, vertical="top")
            _ajustar_larguras(ws, df)
            for k, col in enumerate(df.columns, start=1):
                if pd.api.types.is_float_dtype(df[col]):
                    for (cel,) in ws.iter_rows(min_row=2, min_col=k, max_col=k):
                        cel.number_format = FORMATO_NUM
            if destacar and nome in destacar:
                marcas = destacar[nome].to_numpy()
                for i in np.flatnonzero(marcas):
                    for cel in ws[i + 2]:
                        cel.fill = DESTAQUE


# ---------------------------------------------------------------------------------------------
# MIP com o layout da planilha original
# ---------------------------------------------------------------------------------------------
def exportar_longo(mip: MIP, caminho: str | Path) -> Path:
    """Salva a MIP em CSV formato longo (ver leitura.formato_longo)."""
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    formato_longo(mip).to_csv(caminho, index=False)
    return caminho


def escrever_aba_mip(ws, mip: MIP, aviso_topo: str | None = None):
    """Escreve a MIP numa aba, com o mesmo layout da planilha original.

    Linhas de totais (CI, CI+importações, total de impostos, TOTAL VBP, CHECK) saem como fórmulas,
    como na original. `aviso_topo` (opcional) vai na célula A3, que é vazia na original, e é
    prefixado ao título — assim o aviso fica no topo sem deslocar o layout.
    """
    lay = mip.layout
    n = mip.n
    ult_col = lay.col_total
    L0, L1 = lay.linha_ini, lay.linha_fim
    letra = get_column_letter

    # títulos
    t1, t2 = mip.titulo
    ws.cell(1, 1, f"{aviso_topo} — {t1}" if aviso_topo else t1)
    ws.cell(2, 1, t2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ult_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ult_col)
    if aviso_topo:
        ws.cell(3, 1, aviso_topo).font = Font(bold=True, color="C00000", size=14)

    # cabeçalhos dos setores (linhas 3–6) e rótulos das linhas (A:D)
    ws.cell(6, 1, "Região")
    ws.cell(6, 3, "Setor / Setor")
    regiao_anterior = None
    for k, s in enumerate(mip.setores.itertuples(index=False)):
        c = lay.col_ini + k
        r = L0 + k
        rot = ROTULOS_REGIAO[s.regiao]
        if s.regiao != regiao_anterior:
            ws.cell(3, c, rot["nome"])
            ws.cell(r, 1, rot["nome"])
            regiao_anterior = s.regiao
        ws.cell(4, c, rot["cab"])
        ws.cell(5, c, s.nome)
        ws.cell(6, c, s.codigo)
        ws.cell(r, 2, rot["linha"])
        ws.cell(r, 3, s.codigo)
        ws.cell(r, 4, s.nome)

    # cabeçalhos da demanda final e do total
    ws.cell(3, lay.col_df_ini, "FINAL DEMAND")
    ws.cell(3, lay.col_total, "TOTAL")
    for k, d in enumerate(COLUNAS_DF.itertuples(index=False)):
        c = lay.col_df_ini + k
        if d.cab4:
            ws.cell(4, c, d.cab4)
        if d.cab5:
            ws.cell(5, c, d.cab5)
        if d.regiao:
            ws.cell(6, c, ROTULOS_REGIAO[d.regiao]["cab"])

    # Z, demanda final e total da linha
    for i in range(n):
        r = L0 + i
        for j in range(n):
            ws.cell(r, lay.col_ini + j, float(mip.Z[i, j]))
        for j in range(len(COLUNAS_DF)):
            ws.cell(r, lay.col_df_ini + j, float(mip.F[i, j]))
        ws.cell(r, lay.col_total, float(mip.total_linha[i]))

    # rótulos das linhas inferiores
    for item, rotulo in ROTULOS_LINHAS_INFERIORES.items():
        ws.cell(lay.linha(item), 4, rotulo)

    # linhas inferiores: valores e fórmulas
    for j in range(n):
        c = lay.col_ini + j
        L = letra(c)
        ws.cell(lay.linha("ci_total"), c, f"=SUM({L}{L0}:{L}{L1})")
        ws.cell(lay.linha("ci_importacoes"), c, f"={L}{lay.linha('ci_total')}+{L}{lay.linha('importacoes')}")
        ws.cell(
            lay.linha("total_impostos"), c, f"=SUM({L}{lay.linha('ipi')}:{L}{lay.linha('outros_impostos')})"
        )
        ws.cell(
            lay.linha("total_vbp"),
            c,
            f"={L}{lay.linha('ci_importacoes')}+{L}{lay.linha('total_impostos')}+{L}{lay.linha('va')}",
        )
        ws.cell(lay.linha("check"), c, f"={L}{lay.linha('total_vbp')}-{L}{lay.linha('vbp')}")
        for k, item in enumerate(ITENS_PRIMARIOS):
            ws.cell(lay.linha(item), c, float(mip.V[k, j]))
        ws.cell(lay.linha("vbp"), c, float(mip.X[j]))
        for k, item in enumerate(ITENS_SATELITE):
            ws.cell(lay.linha(item), c, float(mip.S[k, j]))
    for j in range(len(COLUNAS_DF)):
        c = lay.col_df_ini + j
        for k, item in enumerate(ITENS_PRIMARIOS_DF):
            ws.cell(lay.linha(item), c, float(mip.V_df[k, j]))
        ws.cell(lay.linha("vbp"), c, float(mip.X_df[j]))
    for item in ITENS_PRIMARIOS + ["vbp"]:
        if item in mip.totais_ev:
            ws.cell(lay.linha(item), lay.col_total, float(mip.totais_ev[item]))

    # formatação leve (opcional)
    ws.freeze_panes = ws.cell(L0, lay.col_ini)
    ws.column_dimensions["D"].width = 45
    for r in (3, 4, 5, 6):
        for c in range(1, ult_col + 1):
            ws.cell(r, c).font = NEGRITO


def exportar_xlsx(
    mip: MIP,
    caminho: str | Path,
    aviso_topo: str | None = None,
    leiame: list[str] | None = None,
    abas_extras: dict[str, pd.DataFrame] | None = None,
    destacar_setores: list[str] | None = None,
) -> Path:
    """Exporta a MIP com o mesmo layout da original (aba principal) e abas extras opcionais.

    Abas extras típicas: LEIAME (texto), mapa_setores, parametros, checagens, coef_latentes.
    `destacar_setores`: ids (ex.: 'RS_S38_HID') cujas linhas/colunas recebem destaque visual.
    """
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = ABA_PRINCIPAL
    escrever_aba_mip(ws, mip, aviso_topo=aviso_topo)
    if destacar_setores:
        lay = mip.layout
        ids = list(mip.setores["id"])
        for sid in destacar_setores:
            k = ids.index(sid)
            for c in range(1, 5):
                ws.cell(lay.linha_ini + k, c).fill = DESTAQUE
            for r in range(3, 7):
                ws.cell(r, lay.col_ini + k).fill = DESTAQUE
    if leiame:
        wl = wb.create_sheet("LEIAME")
        for k, linha in enumerate(leiame, start=1):
            wl.cell(k, 1, linha)
        wl.column_dimensions["A"].width = 140
        if aviso_topo:
            wl.cell(1, 1).font = Font(bold=True, color="C00000", size=14)
    for nome, df in (abas_extras or {}).items():
        we = wb.create_sheet(nome)
        for c, col in enumerate(df.columns, start=1):
            we.cell(1, c, str(col)).font = NEGRITO
        for r, linha in enumerate(df.itertuples(index=False), start=2):
            for c, v in enumerate(linha, start=1):
                if isinstance(v, (np.floating, np.integer)):
                    v = v.item()
                elif isinstance(v, (np.bool_, bool)):
                    v = bool(v)
                elif v is not None and not isinstance(v, (int, float, str)):
                    v = str(v)
                if isinstance(v, float) and np.isnan(v):
                    v = None
                we.cell(r, c, v)
        we.freeze_panes = "A2"
        _ajustar_larguras(we, df)
    wb.save(caminho)
    return caminho
