"""Leitura da MIP a partir da planilha (original ou exportada com o mesmo layout).

Usa as coordenadas fixas descritas em `estrutura` (Seção 5 das instruções da Rodada 01).
Na original: Z em E7:EJ142, demanda final em EK7:EU142, total em EV, insumos primários nas
linhas 144/147–149/151–154, VBP na linha 156, ocupações na 158 e contas ambientais nas 161–163.
As linhas de fórmulas (143, 145, 150, 157, 159) não são lidas: são recalculadas em Python.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from .caminhos import CAMINHO_ORIGINAL
from .estrutura import (
    COLUNAS_DF,
    ITENS_PRIMARIOS,
    ITENS_PRIMARIOS_DF,
    ITENS_SATELITE,
    LINHA_PRIMEIRO_SETOR,
    MIP,
    N_ORIGINAL,
    REGIAO_POR_ROTULO_CAB,
    REGIAO_POR_ROTULO_LINHA,
    ROTULOS_LINHAS_INFERIORES,
    UNIDADES,
    Layout,
)

ABA_PRINCIPAL = "MIP SS RS_RBR"


def _num(v) -> float:
    """Converte o conteúdo de uma célula em número (vazia → 0). Fórmulas não são aceitas aqui."""
    if v is None:
        return 0.0
    if isinstance(v, str):
        raise ValueError(f"célula de dados com texto/fórmula: {v!r}")
    return float(v)


def _detectar_n(grade: list) -> int:
    """Número de setores-região: a linha 'Importações internacionais' (coluna D) fica em 8 + n."""
    for i, linha in enumerate(grade, start=1):
        if len(linha) >= 4 and linha[3] == ROTULOS_LINHAS_INFERIORES["importacoes"]:
            return i - 8
    raise ValueError("rótulo 'Importações internacionais' não encontrado na coluna D")


def ler_mip(caminho: str | Path = CAMINHO_ORIGINAL, aba: str | None = None) -> MIP:
    """Lê a planilha e devolve um objeto `MIP`.

    Se o arquivo tiver uma aba `coef_latentes` (MIP desagregada exportada por este pacote),
    os coeficientes dos subsetores latentes também são lidos.
    """
    caminho = Path(caminho)
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=False)
    ws = wb[aba] if aba else wb[wb.sheetnames[0]]
    grade = [list(r) for r in ws.iter_rows(values_only=True)]
    coef_latentes = None
    if "coef_latentes" in wb.sheetnames:
        coef_latentes = pd.read_excel(caminho, sheet_name="coef_latentes")
    wb.close()

    n = _detectar_n(grade)
    lay = Layout(n)

    def cel(linha: int, col: int):
        lin = grade[linha - 1]
        return lin[col - 1] if col - 1 < len(lin) else None

    def bloco(l0: int, l1: int, c0: int, c1: int) -> np.ndarray:
        return np.array(
            [[_num(cel(r, c)) for c in range(c0, c1 + 1)] for r in range(l0, l1 + 1)], dtype=float
        )

    # --- rótulos das linhas (A:D) e conferência com os cabeçalhos das colunas (linhas 4 e 6) ---
    regioes, codigos, nomes = [], [], []
    for k, r in enumerate(range(lay.linha_ini, lay.linha_fim + 1)):
        rot_reg = cel(r, 2)
        if rot_reg not in REGIAO_POR_ROTULO_LINHA:
            raise ValueError(f"linha {r}: rótulo de região desconhecido {rot_reg!r}")
        reg = REGIAO_POR_ROTULO_LINHA[rot_reg]
        cod = cel(r, 3)
        c = lay.col_ini + k
        if cel(6, c) != cod or REGIAO_POR_ROTULO_CAB.get(cel(4, c)) != reg:
            raise ValueError(
                f"linha {r} ({reg} {cod}) não corresponde à coluna {c} "
                f"({cel(4, c)} {cel(6, c)}): a matriz não é quadrada na mesma ordem"
            )
        regioes.append(reg)
        codigos.append(str(cod))
        nomes.append(cel(r, 4))
    setores = pd.DataFrame({"regiao": regioes, "codigo": codigos, "nome": nomes})

    # --- blocos numéricos ---------------------------------------------------------------------
    L0, L1 = lay.linha_ini, lay.linha_fim
    Z = bloco(L0, L1, lay.col_ini, lay.col_fim)
    F = bloco(L0, L1, lay.col_df_ini, lay.col_df_fim)
    total_linha = bloco(L0, L1, lay.col_total, lay.col_total).ravel()

    V = np.vstack(
        [bloco(lay.linha(it), lay.linha(it), lay.col_ini, lay.col_fim).ravel() for it in ITENS_PRIMARIOS]
    )
    V_df = np.vstack(
        [
            bloco(lay.linha(it), lay.linha(it), lay.col_df_ini, lay.col_df_fim).ravel()
            for it in ITENS_PRIMARIOS_DF
        ]
    )
    X = bloco(lay.linha("vbp"), lay.linha("vbp"), lay.col_ini, lay.col_fim).ravel()
    X_df = bloco(lay.linha("vbp"), lay.linha("vbp"), lay.col_df_ini, lay.col_df_fim).ravel()
    S = np.vstack(
        [bloco(lay.linha(it), lay.linha(it), lay.col_ini, lay.col_fim).ravel() for it in ITENS_SATELITE]
    )
    totais_ev = {
        it: _num(cel(lay.linha(it), lay.col_total)) for it in ITENS_PRIMARIOS + ["vbp"]
    }

    # conferência dos rótulos das linhas inferiores (garante que as coordenadas estão certas)
    for item, rotulo in ROTULOS_LINHAS_INFERIORES.items():
        if cel(lay.linha(item), 4) != rotulo:
            raise ValueError(
                f"linha {lay.linha(item)}: esperado rótulo {rotulo!r}, encontrado {cel(lay.linha(item), 4)!r}"
            )
    titulo = (cel(1, 1), cel(2, 1))

    return MIP(
        setores=setores,
        Z=Z,
        F=F,
        total_linha=total_linha,
        X=X,
        V=V,
        V_df=V_df,
        X_df=X_df,
        S=S,
        totais_ev=totais_ev,
        coef_latentes=coef_latentes,
        titulo=titulo,
    )


def ler_mip_original() -> MIP:
    """Lê a planilha original de `dados/original/` e confere que tem 136 setores-região."""
    mip = ler_mip(CAMINHO_ORIGINAL)
    if mip.n != N_ORIGINAL:
        raise ValueError(f"a original deveria ter {N_ORIGINAL} setores-região; tem {mip.n}")
    assert LINHA_PRIMEIRO_SETOR == 7 and mip.layout.col_total == 152  # EV
    return mip


def formato_longo(mip: MIP) -> pd.DataFrame:
    """MIP em formato longo: uma linha por célula (Z, demanda final, insumos primários, VBP, satélite).

    Colunas: bloco, regiao_origem, setor_origem, regiao_destino, setor_ou_coluna_destino, valor, unidade.
    Para insumos primários e contas satélite, `setor_origem` é o item (ex.: importacoes, co2e)
    e `regiao_origem` fica vazia.
    """
    reg = mip.setores["regiao"].to_numpy()
    cod = mip.setores["codigo"].to_numpy()
    n = mip.n
    partes = []
    # Z
    partes.append(
        pd.DataFrame(
            {
                "bloco": "Z",
                "regiao_origem": np.repeat(reg, n),
                "setor_origem": np.repeat(cod, n),
                "regiao_destino": np.tile(reg, n),
                "setor_ou_coluna_destino": np.tile(cod, n),
                "valor": mip.Z.ravel(),
                "unidade": "R$ milhões",
            }
        )
    )
    # demanda final
    ndf = len(COLUNAS_DF)
    partes.append(
        pd.DataFrame(
            {
                "bloco": "demanda_final",
                "regiao_origem": np.repeat(reg, ndf),
                "setor_origem": np.repeat(cod, ndf),
                "regiao_destino": np.tile(COLUNAS_DF["regiao"].to_numpy(), n),
                "setor_ou_coluna_destino": np.tile(COLUNAS_DF["codigo"].to_numpy(), n),
                "valor": mip.F.ravel(),
                "unidade": "R$ milhões",
            }
        )
    )
    # insumos primários e VBP nas colunas dos setores
    for nome_item, vetor in list(zip(ITENS_PRIMARIOS, mip.V)) + [("vbp", mip.X)]:
        partes.append(
            pd.DataFrame(
                {
                    "bloco": "vbp" if nome_item == "vbp" else "insumos_primarios",
                    "regiao_origem": "",
                    "setor_origem": nome_item,
                    "regiao_destino": reg,
                    "setor_ou_coluna_destino": cod,
                    "valor": vetor,
                    "unidade": "R$ milhões",
                }
            )
        )
    # insumos primários e total nas colunas de demanda final
    for nome_item, vetor in list(zip(ITENS_PRIMARIOS_DF, mip.V_df)) + [("vbp", mip.X_df)]:
        partes.append(
            pd.DataFrame(
                {
                    "bloco": "total_demanda_final" if nome_item == "vbp" else "insumos_primarios_df",
                    "regiao_origem": "",
                    "setor_origem": "total" if nome_item == "vbp" else nome_item,
                    "regiao_destino": COLUNAS_DF["regiao"].to_numpy(),
                    "setor_ou_coluna_destino": COLUNAS_DF["codigo"].to_numpy(),
                    "valor": vetor,
                    "unidade": "R$ milhões",
                }
            )
        )
    # contas satélite
    for nome_item, vetor in zip(ITENS_SATELITE, mip.S):
        partes.append(
            pd.DataFrame(
                {
                    "bloco": "satelite",
                    "regiao_origem": "",
                    "setor_origem": nome_item,
                    "regiao_destino": reg,
                    "setor_ou_coluna_destino": cod,
                    "valor": vetor,
                    "unidade": UNIDADES[nome_item],
                }
            )
        )
    return pd.concat(partes, ignore_index=True)


def caminho_coef_latentes_ao_lado(caminho_longo: str | Path) -> Path:
    """Arquivo dos coeficientes latentes que acompanha um CSV longo: x_longo.csv → x_coef_latentes.csv
    (outro nome: x.csv → x_coef_latentes.csv). Usado por quem grava e por quem lê."""
    p = Path(caminho_longo)
    base = p.name[: -len("_longo.csv")] if p.name.endswith("_longo.csv") else p.stem
    return p.with_name(base + "_coef_latentes.csv")


def ler_mip_longo(caminho: str | Path, caminho_coef_latentes: str | Path | None = None) -> MIP:
    """Reconstrói uma MIP a partir do CSV em formato longo (ver `formato_longo`).

    Os coeficientes dos subsetores latentes não estão no formato longo: se houver um arquivo
    `<nome>_coef_latentes.csv` ao lado (ou `caminho_coef_latentes`), ele é lido. Os nomes dos setores
    não estão no formato longo (ficam iguais aos códigos); `total_linha` = Σ Z + Σ F e `totais_ev` = somas.
    """
    caminho = Path(caminho)
    df = pd.read_csv(caminho, dtype={"regiao_origem": str, "regiao_destino": str}, keep_default_na=False)
    z = df[df["bloco"] == "Z"]
    ordem = list(dict.fromkeys(zip(z["regiao_origem"], z["setor_origem"])))
    n = len(ordem)
    pos = {k: i for i, k in enumerate(ordem)}
    setores = pd.DataFrame({"regiao": [a for a, _ in ordem], "codigo": [b for _, b in ordem],
                            "nome": [b for _, b in ordem]})
    Z = np.zeros((n, n))
    Z[[pos[k] for k in zip(z["regiao_origem"], z["setor_origem"])],
      [pos[k] for k in zip(z["regiao_destino"], z["setor_ou_coluna_destino"])]] = z["valor"].to_numpy(float)
    dfc = {c: i for i, c in enumerate(COLUNAS_DF["codigo"])}
    fd = df[df["bloco"] == "demanda_final"]
    F = np.zeros((n, len(COLUNAS_DF)))
    F[[pos[k] for k in zip(fd["regiao_origem"], fd["setor_origem"])], [dfc[c] for c in fd["setor_ou_coluna_destino"]]] = \
        fd["valor"].to_numpy(float)

    def vetor(bloco, item):
        v = np.zeros(n)
        b = df[(df["bloco"] == bloco) & (df["setor_origem"] == item)]
        v[[pos[k] for k in zip(b["regiao_destino"], b["setor_ou_coluna_destino"])]] = b["valor"].to_numpy(float)
        return v

    def vetor_df(bloco, item):
        v = np.zeros(len(COLUNAS_DF))
        b = df[(df["bloco"] == bloco) & (df["setor_origem"] == item)]
        v[[dfc[c] for c in b["setor_ou_coluna_destino"]]] = b["valor"].to_numpy(float)
        return v

    V = np.vstack([vetor("insumos_primarios", it) for it in ITENS_PRIMARIOS])
    X = vetor("vbp", "vbp")
    S = np.vstack([vetor("satelite", it) for it in ITENS_SATELITE])
    V_df = np.vstack([vetor_df("insumos_primarios_df", it) for it in ITENS_PRIMARIOS_DF])
    X_df = vetor_df("total_demanda_final", "total")
    coef = None
    if caminho_coef_latentes and not Path(caminho_coef_latentes).exists():
        raise FileNotFoundError(f"coeficientes latentes não encontrados: {caminho_coef_latentes}")
    cand = Path(caminho_coef_latentes) if caminho_coef_latentes else caminho_coef_latentes_ao_lado(caminho)
    if cand.exists():
        coef = pd.read_csv(cand)
    totais = {it: float(V[k].sum()) for k, it in enumerate(ITENS_PRIMARIOS)} | {"vbp": float(X.sum())}
    return MIP(setores=setores, Z=Z, F=F, total_linha=Z.sum(axis=1) + F.sum(axis=1), X=X, V=V, V_df=V_df, X_df=X_df,
               S=S, totais_ev=totais, coef_latentes=coef)
