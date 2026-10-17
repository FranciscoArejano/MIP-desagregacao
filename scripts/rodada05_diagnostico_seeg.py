"""Rodada 05 — diagnóstico da conta-satélite de emissões do S38 (NEREUS × SEEG 13.0) e componentes da E4.

Uso (na raiz do repositório):
    python scripts/rodada05_diagnostico_seeg.py

Instruções: instrucoes/RODADA_05_instrucoes_CODE.md, Seções 3, 3.1 e 3.2. O script NÃO muda a MIP: só lê o SEEG e a
base NEREUS (original, intocável) e grava o diagnóstico.

Fonte: SEEG 13.0, "Dados nacionais", aba 'Dados' (xlsx de ~185 MB, fora do git; registro do download e SHA-256 em
dados/externos/seeg/DOWNLOAD.txt): https://seeg.eco.br/wp-content/uploads/2025/12/Dados-nacionais-13.0.xlsx

Cache: ler o xlsx com openpyxl leva ~2 min. Na primeira execução o script grava todas as linhas da aba 'Dados' com o
valor de 2019 (todas as métricas de gás, todos os tipos de 'Emissão/Remoção/Bunker'; colunas de texto + valor_2019_t)
em dados/externos/brutos/seeg/seeg13_2019_todas_linhas.csv.gz, com um seeg13_2019_todas_linhas.meta.json ao lado
(SHA-256 do bruto, ano, número de linhas). A pasta dados/externos/brutos/ é ignorada pelo git. Nas execuções
seguintes o cache é reaproveitado só se o SHA-256 do bruto (recalculado a cada execução) conferir com o do .meta.json
e com o registrado em DOWNLOAD.txt; senão o xlsx é lido de novo. Apagar o cache força a releitura. O cache não é
cópia literal do bruto: os rótulos de texto têm os espaços das pontas tirados (str.strip; no bruto, 'Efluentes líquidos
industriais ', 'Disposição final ' e 'Alimentos ' têm espaço no fim, 3544 linhas, nenhuma em Geração de eletricidade
ou Produção de combustíveis); nenhuma soma muda.

Critérios (Seção 3): ano 2019; Gás = 'CO2e (t) GWP-AR5'; 'Emissão/Remoção/Bunker' ∈ {'Emissão', 'Emissão NCI'}
(emissão bruta); região do projeto: Estado = 'Rio Grande do Sul' → RS; 'Não Alocado' → 'Não alocado'; demais
estados → RB. "S38 pelo critério do revisor" = 'Atividade geral' ∈ {'Geração de eletricidade', 'Produção de
combustíveis'}.

Métricas (Seção 3.1, item 2): GWP-AR2, AR4, AR5, AR6 e GTP-AR2, AR4, AR5, AR6. No arquivo, parte das linhas do
GTP-AR4 (736 linhas, todas do setor Resíduos) vem com o rótulo 'CO2e (t) GTP AR4' (sem hífen); não há sobreposição
com as linhas 'CO2e (t) GTP-AR4' e as duas grafias são tratadas como a mesma métrica. Bases (principais = convenção
do próprio SEEG, a da pivot 'Consulta Série Histórica' do arquivo): bruta (SEEG) = 'Emissão'; líquida (SEEG) =
'Emissão' + 'Remoção' (remoções vêm negativas). Variantes: bruta com NCI = 'Emissão' + 'Emissão NCI' (o critério da
Seção 3; NCI = não contabilizadas no inventário: solos manejados e queimadas não associadas a desmatamento; no S38
não há NCI, então as duas brutas coincidem); líquida com NCI = as quatro; e, só como sensibilidade, líquida (SEEG) +
'Bunker'. Fora disso, 'Bunker' (combustível de navegação e aviação internacionais, todo em Estado = 'Não Alocado')
fica fora. Total Brasil = 26 UF + RS + 'Não Alocado' (a NEREUS, RS + RB, cobre o país todo); a versão só com as UF
também é mostrada. Empates exatos entre métricas/bases são registrados (GTP-AR2 = GTP-AR4).

Casamentos (Seção 3.1, item 3): setores S01, S02, S03, S19, S26, S39, S43 do RS (e S38 como controle) contra somas
de 1, 2 ou 3 unidades do SEEG do RS (GWP-AR5, bruta), dentro de cada tipo de unidade: 'Atividade geral'; 'Setor de
emissão › Categoria emissora'; 'Atividade geral › Categoria emissora › Sub-categoria emissora' (só unidades com valor
≠ 0). No arquivo, GTP-AR2 e GTP-AR4 têm valores idênticos linha a linha (conferido nesta execução e
registrado no LEIAME). O sentido econômico de cada unidade para cada setor (sim / discutível / não) é uma classificação
deste diagnóstico (HIPÓTESE, documentada na aba unidades_RS); o da combinação é o pior nível entre as suas unidades.
O critério é o do SCN para todos os setores: 'Produção de combustíveis' (refino, carvão mineral, exploração) é 'não'
para o S38, sem exceção pela hipótese em teste; a combinação do revisor (S38 = Geração de eletricidade + Produção de
combustíveis) é marcada à parte como 'hipótese formulada antes da busca'. Mudança de uso da terra (MUT) é 'discutível'
para S01, S02 e S03 em U2 e em U3. 'Esperado por acaso' (nº em ±5% ÷ 50) é só ordem de grandeza (densidade uniforme);
unidades desprezíveis (|valor| < 0.1% do alvo) geram combinações quase repetidas, por isso também se contam os
'núcleos' distintos (a combinação sem as unidades desprezíveis).

Saídas (só arquivos novos):
- dados/externos/seeg/seeg13_2019_recorte_energia_v04.csv — o recorte da tabela da Seção 3 (versionado; as linhas
  com valor 2019 = 0 ficam fora do CSV para mantê-lo pequeno, o que não muda nenhuma soma);
- dados/externos/seeg/seeg13_2019_componentes_e4_v04.csv — as 5 linhas da regra da E4 (Seção 3.2);
- resultados/rodada05/diagnostico_seeg_nereus.xlsx — abas LEIAME, componentes_S38, componentes_S38_detalhe,
  metricas, casamentos_RS, unidades_RS, resumo.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import itertools
import json
import re
import sys
import time
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402

from mip.caminhos import PASTA_EXTERNOS, PASTA_RESULTADOS  # noqa: E402
from mip.estrutura import ITENS_SATELITE  # noqa: E402
from mip.leitura import ler_mip_original  # noqa: E402

# --------------------------------------------------------------------------------------------------------------------
# Constantes
# --------------------------------------------------------------------------------------------------------------------
ANO = 2019
URL_SEEG = "https://seeg.eco.br/wp-content/uploads/2025/12/Dados-nacionais-13.0.xlsx"
SHA_SEEG = "21fc49654828cdb2eb63c104f06ba19d08b7dd81229fd9bf8972d80834aff350"
BRUTO = PASTA_EXTERNOS / "brutos" / "seeg" / "Dados-nacionais-13.0.xlsx"
CACHE = PASTA_EXTERNOS / "brutos" / "seeg" / f"seeg13_{ANO}_todas_linhas.csv.gz"
CACHE_META = PASTA_EXTERNOS / "brutos" / "seeg" / f"seeg13_{ANO}_todas_linhas.meta.json"
DOWNLOAD_TXT = PASTA_EXTERNOS / "seeg" / "DOWNLOAD.txt"

SAIDA_RECORTE = PASTA_EXTERNOS / "seeg" / "seeg13_2019_recorte_energia_v04.csv"
SAIDA_E4 = PASTA_EXTERNOS / "seeg" / "seeg13_2019_componentes_e4_v04.csv"
SAIDA_XLSX = PASTA_RESULTADOS / "rodada05" / "diagnostico_seeg_nereus.xlsx"

COLUNAS_TEXTO = [
    "Emissão/Remoção/Bunker", "Gás", "Setor de emissão", "Categoria emissora", "Sub-categoria emissora",
    "Produto ou sistema", "Detalhamento", "Recorte", "Atividade geral", "Estado", "Bioma",
]
ERB, GAS, SETOR, CAT, SUB, ATIV, ESTADO = (
    "Emissão/Remoção/Bunker", "Gás", "Setor de emissão", "Categoria emissora", "Sub-categoria emissora",
    "Atividade geral", "Estado",
)
COL_VALOR = "valor_2019_t"

GAS_AR5 = "CO2e (t) GWP-AR5"
METRICAS = [
    "CO2e (t) GWP-AR2", "CO2e (t) GWP-AR4", "CO2e (t) GWP-AR5", "CO2e (t) GWP-AR6",
    "CO2e (t) GTP-AR2", "CO2e (t) GTP-AR4", "CO2e (t) GTP-AR5", "CO2e (t) GTP-AR6",
]
GRAFIAS_METRICA = {"CO2e (t) GTP AR4": "CO2e (t) GTP-AR4"}  # 736 linhas de Resíduos com o rótulo sem hífen

EMISSAO = ("Emissão", "Emissão NCI")
REMOCAO = ("Remoção", "Remoção NCI")
BUNKER = "Bunker"
BRUTA, LIQUIDA = "bruta (SEEG)", "líquida (SEEG)"
BRUTA_NCI, LIQUIDA_NCI, LIQUIDA_BUNKER = "bruta com NCI", "líquida com NCI", "líquida (SEEG) + Bunker"
BASES = {  # base → valores de 'Emissão/Remoção/Bunker' somados
    BRUTA: ("Emissão",),  # convenção do SEEG (pivot 'Consulta Série Histórica' do próprio arquivo)
    LIQUIDA: ("Emissão", "Remoção"),  # convenção do SEEG
    BRUTA_NCI: ("Emissão", "Emissão NCI"),  # critério da Seção 3 (no S38 = bruta (SEEG): não há NCI)
    LIQUIDA_NCI: ("Emissão", "Emissão NCI", "Remoção", "Remoção NCI"),
    LIQUIDA_BUNKER: ("Emissão", "Remoção", "Bunker"),  # só sensibilidade
}
BASES_PRINCIPAIS = (BRUTA, LIQUIDA)
GRUPOS_BASES = {  # conjunto de bases → bases consideradas na escolha da 'métrica mais perto'
    "convenção SEEG (bruta e líquida)": BASES_PRINCIPAIS,
    "convenção SEEG + variantes com NCI": (BRUTA, LIQUIDA, BRUTA_NCI, LIQUIDA_NCI),
    "todas (inclui líquida + Bunker, sensibilidade)": tuple(BASES),
}

ATIV_GE, ATIV_PC = "Geração de eletricidade", "Produção de combustíveis"
ATIVIDADES_S38 = (ATIV_GE, ATIV_PC)
CAT_PUBLICO = "Geração de eletricidade (serviço público)"
SUB_REFINO, SUB_EXPLOR, SUB_ALCOOL, SUB_CARVAO = (
    "Refino de petróleo", "Exploração de petróleo e gás natural", "Produção de álcool",
    "Produção de carvão mineral e outros",
)

RS, RB, NA = "RS", "RB", "Não alocado"
TOT_BR = "Total Brasil (26 UF + RS + Não Alocado)"
TOT_BR_UF = "Total Brasil (só UF: 26 UF + RS)"
UF_RS, ESTADO_NA = "Rio Grande do Sul", "Não Alocado"

# Componentes da tabela da Seção 3 e valores de referência das instruções (mil t CO2e; RS, RB (estados), Não alocado).
# Usados só para comparação, nunca como dado.
COMPONENTES = [
    ("GE_publico", "Geração de eletricidade › Geração de eletricidade (serviço público)", (3842.6, 34181.1, 344.0)),
    ("GE_demais", "Geração de eletricidade › demais categorias (autoprodução: indústria, comércio, público, "
                  "agropecuária)", (96.5, 2094.4, 4833.4)),
    ("GE_PC_refino", "Geração de eletricidade › Produção de combustíveis › Refino de petróleo", (76.1, 759.7, 496.3)),
    ("GE_PC_exploracao", "Geração de eletricidade › Produção de combustíveis › Exploração de petróleo e gás natural",
     (0.0, 5778.2, 0.0)),
    ("GE_PC_alcool", "Geração de eletricidade › Produção de combustíveis › Produção de álcool", (0.0, 0.0, 412.4)),
    ("GE_PC_outros", "Geração de eletricidade › Produção de combustíveis › outras sub-categorias (não está na tabela "
                     "das instruções)", (None, None, None)),
    ("PC_refino", "Produção de combustíveis › Refino de petróleo", (1736.0, 17329.1, 7053.5)),
    ("PC_exploracao", "Produção de combustíveis › Exploração de petróleo e gás natural", (0.0, 17721.0, 53.4)),
    ("PC_carvao_mineral", "Produção de combustíveis › Produção de carvão mineral e outros", (88.1, 1155.4, 355.6)),
    ("PC_demais", "Produção de combustíveis › demais (carvão vegetal, álcool, transporte de petróleo e gás)",
     (0.0, 0.0, 5151.2)),
]
REF_SOMA = (5839.2, 79018.9, 18700.0)
REF_NEREUS = (5838.0, 84518.3, None)

# Regra da E4 (Seção 3.2): (região, rótulo, componentes, setor de destino, referência em mil t)
REGRA_E4 = [
    (RS, "Refino de petróleo (Produção de combustíveis + Geração de eletricidade › Produção de combustíveis)",
     ("PC_refino", "GE_PC_refino"), "S19", 1812.1),
    (RS, "Produção de carvão mineral e outros", ("PC_carvao_mineral",), "S04", 88.1),
    (RB, "Refino de petróleo (Produção de combustíveis + Geração de eletricidade › Produção de combustíveis)",
     ("PC_refino", "GE_PC_refino"), "S19", 18088.8),
    (RB, "Exploração de petróleo e gás natural (Produção de combustíveis + Geração de eletricidade › Produção de "
         "combustíveis)", ("PC_exploracao", "GE_PC_exploracao"), "S05", 23499.2),
    (RB, "Produção de carvão mineral e outros", ("PC_carvao_mineral",), "S04", 1155.4),
]

ACASO = "esperado por acaso em ±0.1% (ordem de grandeza: nº em ±5% ÷ 50)"
SETORES_CASAMENTO = ["S01", "S02", "S03", "S19", "S26", "S39", "S43", "S38"]  # S38 = controle
TOL_CASAMENTO = 0.1  # %
TAMANHOS = (1, 2, 3)
TIPOS_UNIDADE = {
    "U1": ("Atividade geral", [ATIV]),
    "U2": ("Setor de emissão › Categoria emissora", [SETOR, CAT]),
    "U3": ("Atividade geral › Categoria emissora › Sub-categoria emissora", [ATIV, CAT, SUB]),
}
NIVEIS = {"sim": 0, "discutível": 1, "não": 2}
# Teste direto do padrão "atividade do SEEG → setor da MIP": a(s) 'Atividade geral' que corresponde(m) a cada setor.
ATIV_NATURAL = {
    "S01": ["Agricultura"], "S02": ["Pecuária"], "S03": ["Silvicultura"], "S19": ["Produção de combustíveis"],
    "S26": ["Cimento"], "S39": ["Saneamento Básico"], "S43": ["Transporte de carga", "Transporte de passageiros"],
    "S38": ["Geração de eletricidade", "Produção de combustíveis"],
}


# --------------------------------------------------------------------------------------------------------------------
# Leitura
# --------------------------------------------------------------------------------------------------------------------
def sha256(caminho: Path) -> str:
    h = hashlib.sha256()
    with open(caminho, "rb") as f:
        for bloco in iter(lambda: f.read(1 << 22), b""):
            h.update(bloco)
    return h.hexdigest()


NOTA_ROTULOS = ("rótulos de texto com os espaços das pontas tirados (str.strip); no bruto, 'Efluentes líquidos "
                "industriais ', 'Disposição final ' e 'Alimentos ' têm espaço no fim; nenhuma soma muda")


def ler_seeg_2019() -> tuple[pd.DataFrame, str]:
    """Todas as linhas da aba 'Dados' com o valor de 2019 (todas as métricas de gás), com cache conferido por SHA-256."""
    if not BRUTO.exists():
        raise FileNotFoundError(f"bruto do SEEG ausente: {BRUTO} (ver {DOWNLOAD_TXT.relative_to(RAIZ)})")
    if SHA_SEEG not in DOWNLOAD_TXT.read_text(encoding="utf-8"):
        raise ValueError(f"o SHA-256 esperado não está em {DOWNLOAD_TXT}")
    sha = sha256(BRUTO)
    if sha != SHA_SEEG:
        raise ValueError(f"SHA-256 do bruto não confere: {sha} (esperado {SHA_SEEG})")
    if CACHE.exists() and CACHE_META.exists():
        meta = json.loads(CACHE_META.read_text(encoding="utf-8"))
        if meta.get("sha256_bruto") == sha and meta.get("ano") == ANO:
            df = pd.read_csv(CACHE, dtype={c: str for c in COLUNAS_TEXTO}, keep_default_na=False)
            if len(df) == meta.get("linhas"):
                if "rotulos" not in meta:  # caches gravados antes desta nota
                    meta["rotulos"] = NOTA_ROTULOS
                    CACHE_META.write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
                return df, f"cache {CACHE.relative_to(RAIZ)} (SHA-256 do bruto conferido)"
    import openpyxl

    t0 = time.time()
    wb = openpyxl.load_workbook(BRUTO, read_only=True, data_only=True)
    ws = wb["Dados"]
    linhas = ws.iter_rows(values_only=True)
    cab = list(next(linhas))
    if cab[: len(COLUNAS_TEXTO)] != COLUNAS_TEXTO:
        raise ValueError(f"cabeçalho inesperado: {cab[:len(COLUNAS_TEXTO)]}")
    j = cab.index(ANO)
    n_txt = len(COLUNAS_TEXTO)
    reg = []
    for r in linhas:
        if r is None or all(v is None for v in r[:n_txt]):
            continue
        v = r[j]
        reg.append(["" if x is None else str(x).strip() for x in r[:n_txt]]
                   + [0.0 if v is None or v == "" else float(v)])
    wb.close()
    df = pd.DataFrame(reg, columns=COLUNAS_TEXTO + [COL_VALOR])
    df.to_csv(CACHE, index=False, compression="gzip")
    CACHE_META.write_text(json.dumps({"sha256_bruto": sha, "ano": ANO, "linhas": len(df),
                                      "gerado_por": "scripts/rodada05_diagnostico_seeg.py",
                                      "segundos_leitura": round(time.time() - t0, 1),
                                      "rotulos": NOTA_ROTULOS}, ensure_ascii=False), encoding="utf-8")
    return df, f"xlsx bruto lido em {time.time() - t0:.0f} s; cache gravado em {CACHE.relative_to(RAIZ)}"


def regiao_do_estado(estado: str) -> str:
    if estado == UF_RS:
        return RS
    if estado == ESTADO_NA:
        return NA
    return RB


def preparar(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["metrica"] = df[GAS].replace(GRAFIAS_METRICA)
    df["regiao"] = df[ESTADO].map(regiao_do_estado)
    return df


def ler_nereus() -> pd.DataFrame:
    m = ler_mip_original()
    s = m.setores[["regiao", "codigo", "nome"]].copy()
    s["co2e_t"] = np.asarray(m.S[ITENS_SATELITE.index("co2e")], dtype=float)
    return s


# --------------------------------------------------------------------------------------------------------------------
# Seção 3: recorte, tabela de componentes e E4
# --------------------------------------------------------------------------------------------------------------------
def componente(ativ: str, cat: str, sub: str) -> str:
    if ativ == ATIV_GE:
        if cat == CAT_PUBLICO:
            return "GE_publico"
        if cat == ATIV_PC:
            return {SUB_REFINO: "GE_PC_refino", SUB_EXPLOR: "GE_PC_exploracao",
                    SUB_ALCOOL: "GE_PC_alcool"}.get(sub, "GE_PC_outros")
        return "GE_demais"
    if ativ == ATIV_PC:
        return {SUB_REFINO: "PC_refino", SUB_EXPLOR: "PC_exploracao",
                SUB_CARVAO: "PC_carvao_mineral"}.get(sub, "PC_demais")
    raise ValueError(ativ)


def recorte_energia(df: pd.DataFrame) -> pd.DataFrame:
    sel = df[(df[GAS] == GAS_AR5) & df[ERB].isin(EMISSAO) & df[ATIV].isin(ATIVIDADES_S38)].copy()
    sel["componente"] = [componente(a, c, s) for a, c, s in zip(sel[ATIV], sel[CAT], sel[SUB])]
    return sel[COLUNAS_TEXTO + [COL_VALOR, "regiao", "componente"]].reset_index(drop=True)


def _pct(a: float, b: float) -> float:
    return np.nan if not b else (a - b) / b * 100.0


def tabela_componentes(rec: pd.DataFrame, nereus: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    piv = rec.pivot_table(index="componente", columns="regiao", values=COL_VALOR, aggfunc="sum", fill_value=0.0)
    piv = piv.reindex(columns=[RS, RB, NA], fill_value=0.0)
    linhas = []
    for cod, rotulo, ref in COMPONENTES:
        v = piv.loc[cod] if cod in piv.index else pd.Series(0.0, index=[RS, RB, NA])
        if cod == "GE_PC_outros" and float(v.abs().sum()) == 0.0:
            continue  # não existe no arquivo: fica fora da tabela
        cats = sorted(set(rec.loc[rec.componente == cod, CAT] + " › " + rec.loc[rec.componente == cod, SUB]))
        linhas.append((cod, rotulo, "; ".join(cats), v[RS], v[RB], v[NA], ref))
    soma = piv.sum()
    linhas.append(("soma", "Soma das duas atividades", "", soma[RS], soma[RB], soma[NA], REF_SOMA))
    ner = nereus[nereus.codigo == "S38"].set_index("regiao").co2e_t
    linhas.append(("nereus", "S38 na base NEREUS", "MIP original, linha CO2e (t)", ner[RS], ner[RB], np.nan,
                   REF_NEREUS))
    out = []
    for cod, rotulo, detalhe, vrs, vrb, vna, ref in linhas:
        reg = {"codigo": cod, "componente (Atividade geral › Categoria emissora › Sub-categoria emissora)": rotulo,
               "categorias › sub-categorias incluídas": detalhe,
               "RS (t)": vrs, "RB estados (t)": vrb, "Não alocado (t)": vna,
               "RS (mil t)": vrs / 1e3, "RB estados (mil t)": vrb / 1e3, "Não alocado (mil t)": vna / 1e3}
        for nome, v, r in zip(("RS", "RB estados", "Não alocado"), (vrs, vrb, vna), ref):
            reg[f"{nome}: referência das instruções (mil t)"] = np.nan if r is None else r
            reg[f"{nome}: diferença (mil t)"] = np.nan if r is None or pd.isna(v) else v / 1e3 - r
        out.append(reg)
    tab = pd.DataFrame(out)
    # Linha extra: soma menos NEREUS (diferença % da Seção 3)
    for nome in ("RS", "RB estados"):
        a = tab.loc[tab.codigo == "soma", f"{nome} (t)"].iloc[0]
        b = tab.loc[tab.codigo == "nereus", f"{nome} (t)"].iloc[0]
        tab.loc[tab.codigo == "nereus", f"{nome}: soma SEEG − NEREUS (%)"] = _pct(a, b)
    detalhe = (rec.groupby(["componente", ATIV, CAT, SUB, "regiao"])[COL_VALOR].sum().unstack("regiao", fill_value=0.0)
               .reindex(columns=[RS, RB, NA], fill_value=0.0).reset_index())
    detalhe = detalhe.rename(columns={RS: "RS (t)", RB: "RB estados (t)", NA: "Não alocado (t)"})
    for c in ("RS", "RB estados", "Não alocado"):
        detalhe[f"{c} (mil t)"] = detalhe[f"{c} (t)"] / 1e3
    return tab, detalhe


def componentes_e4(rec: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    linhas, conf = [], []
    for reg, rotulo, comps, destino, ref in REGRA_E4:
        sub = rec[(rec.regiao == reg) & rec.componente.isin(comps)]
        parcelas = [float(sub.loc[sub.componente == c, COL_VALOR].sum()) / 1e3 for c in comps]
        exato = sum(parcelas)
        soma_arred = round(sum(round(x, 1) for x in parcelas), 1)
        if abs(exato - ref) <= 0.05:
            nota_ref = "confere (arredondamento a 0.1 mil t)"
        elif abs(soma_arred - ref) <= 0.05:
            nota_ref = (f"a referência {ref} das instruções é a soma das parcelas já arredondadas ("
                        + " + ".join(f"{round(x, 1)}" for x in parcelas) + f" = {soma_arred}); o valor exato é "
                        f"{exato:.2f} mil t ({exato:.1f} arredondado), e é o que vai no valor_t")
        else:
            nota_ref = f"DIVERGE da referência {ref} mil t (exato {exato:.2f}) — conferir"
        conf.append({"regiao": reg, "setor_destino": destino, "componente": rotulo,
                     "parcelas (mil t, exatas)": "; ".join(f"{c} = {x:.4f}" for c, x in zip(comps, parcelas)),
                     "valor exato (mil t)": exato, "referência das instruções (mil t)": ref,
                     "diferença (mil t)": exato - ref, "soma das parcelas arredondadas (mil t)": soma_arred,
                     "nota": nota_ref})
        filtros = sorted(set(sub[ATIV] + " › " + sub[CAT] + " › " + sub[SUB]))
        filtro = (f"ano {ANO}; Gás = '{GAS_AR5}'; Emissão/Remoção/Bunker ∈ {{Emissão, Emissão NCI}}; "
                  + ("Estado = 'Rio Grande do Sul'" if reg == RS else
                     "Estado ∉ {'Rio Grande do Sul', 'Não Alocado'} (26 UF)")
                  + "; Atividade geral › Categoria emissora › Sub-categoria emissora ∈ {" + "; ".join(filtros) + "}")
        valor = float(sub[COL_VALOR].sum())
        linhas.append({
            "regiao": reg, "componente": rotulo, "setor_destino": destino, "valor_t": valor,
            "filtro_seeg": filtro,
            "fonte": (f"SEEG 13.0, Dados nacionais, aba Dados ({URL_SEEG}; SHA-256 {SHA_SEEG}); "
                      f"recorte em dados/externos/seeg/{SAIDA_RECORTE.name}; scripts/rodada05_diagnostico_seeg.py; "
                      f"regra da E4 = instruções da Rodada 05, Seção 3.2 (referência {ref} mil t: {nota_ref}); "
                      "'Não alocado' não entra"),
        })
    return pd.DataFrame(linhas), pd.DataFrame(conf)


# --------------------------------------------------------------------------------------------------------------------
# Seção 3.1 item 2: métricas
# --------------------------------------------------------------------------------------------------------------------
def tabela_metricas(df: pd.DataFrame, nereus: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    ner = nereus.set_index(["regiao", "codigo"]).co2e_t
    alvos = {
        "S38-RS": ner[(RS, "S38")],
        "S38-RB (estados)": ner[(RB, "S38")],
        "S38-RB (estados + Não alocado)": ner[(RB, "S38")],
        "Total RS": nereus.loc[nereus.regiao == RS, "co2e_t"].sum(),
        TOT_BR: nereus.co2e_t.sum(),
        TOT_BR_UF: nereus.co2e_t.sum(),
    }
    linhas = []
    for met in METRICAS:
        d = df[df.metrica == met]
        if d.empty:
            raise ValueError(f"métrica ausente no arquivo: {met}")
        for base, erbs in BASES.items():
            b = d[d[ERB].isin(erbs)]
            s38 = b[b[ATIV].isin(ATIVIDADES_S38)]
            seeg = {
                "S38-RS": s38.loc[s38.regiao == RS, COL_VALOR].sum(),
                "S38-RB (estados)": s38.loc[s38.regiao == RB, COL_VALOR].sum(),
                "S38-RB (estados + Não alocado)": s38.loc[s38.regiao.isin([RB, NA]), COL_VALOR].sum(),
                "Total RS": b.loc[b.regiao == RS, COL_VALOR].sum(),
                TOT_BR: b[COL_VALOR].sum(),  # todas as linhas: 26 UF + RS + 'Não Alocado'
                TOT_BR_UF: b.loc[b.regiao != NA, COL_VALOR].sum(),
            }
            for agr, v in seeg.items():
                linhas.append({"metrica": met, "base": base, "principal": base in BASES_PRINCIPAIS,
                               "agregado": agr, "SEEG (t)": v, "NEREUS (t)": alvos[agr],
                               "SEEG (mil t)": v / 1e3, "NEREUS (mil t)": alvos[agr] / 1e3,
                               "diferença SEEG − NEREUS (%)": _pct(v, alvos[agr])})
    longo = pd.DataFrame(linhas)
    longo["|diferença| (%)"] = longo["diferença SEEG − NEREUS (%)"].abs()
    mais_perto = []
    for agr in alvos:
        for grupo, bases in GRUPOS_BASES.items():
            sub = longo[(longo.agregado == agr) & longo.base.isin(bases)]
            r = sub.loc[sub["|diferença| (%)"].idxmin()]
            # empates: outras métrica/base do mesmo conjunto com o mesmo valor (ex.: GTP-AR2 = GTP-AR4)
            emp = sub[np.isclose(sub["SEEG (t)"], r["SEEG (t)"], rtol=1e-12, atol=1e-3) & (sub.index != r.name)]
            ar5 = {b: longo[(longo.agregado == agr) & (longo.metrica == GAS_AR5) & (longo.base == b)].iloc[0]
                   for b in (BRUTA, LIQUIDA)}
            mais_perto.append({
                "agregado": agr, "bases consideradas": grupo,
                "métrica mais perto": r.metrica, "base": r.base,
                "empatadas (mesmo valor)": "; ".join(f"{x.metrica} {x.base}" for x in emp.itertuples()) or "—",
                "SEEG (mil t)": r["SEEG (mil t)"],
                "NEREUS (mil t)": r["NEREUS (mil t)"], "diferença (%)": r["diferença SEEG − NEREUS (%)"],
                f"GWP-AR5 {BRUTA}: diferença (%)": ar5[BRUTA]["diferença SEEG − NEREUS (%)"],
                f"GWP-AR5 {LIQUIDA}: diferença (%)": ar5[LIQUIDA]["diferença SEEG − NEREUS (%)"],
            })
    mais_perto = pd.DataFrame(mais_perto)
    erb = (df[df.metrica == GAS_AR5].groupby(ERB)[COL_VALOR].agg(["count", "sum"]).reset_index()
           .rename(columns={"count": "linhas (GWP-AR5)", "sum": "soma 2019 Brasil, GWP-AR5 (t)"}))
    entra = {"Emissão": "todas as bases (bruta e líquida do SEEG e variantes)",
             "Emissão NCI": "só nas variantes 'com NCI' (bruta com NCI = critério da Seção 3; não há NCI no S38); "
                            "NCI = não contabilizadas no inventário: solos manejados e queimadas não associadas a "
                            "desmatamento",
             "Remoção": "líquida (SEEG), líquida com NCI e líquida + Bunker",
             "Remoção NCI": "só em 'líquida com NCI' (aumento do estoque de C no solo)",
             "Bunker": "fora das bases principais (combustível de navegação e aviação internacionais; não é emissão "
                       "do território; todo em Estado = 'Não Alocado'); entra só na sensibilidade 'líquida (SEEG) + "
                       "Bunker'"}
    erb["onde entra"] = erb[ERB].map(entra).fillna("NÃO PREVISTO — conferir")
    erb["soma 2019 Brasil, GWP-AR5 (mil t)"] = erb["soma 2019 Brasil, GWP-AR5 (t)"] / 1e3
    return longo, mais_perto, erb


# --------------------------------------------------------------------------------------------------------------------
# Seção 3.1 item 3: casamentos no RS
# --------------------------------------------------------------------------------------------------------------------
MUT_CATS = {"Alterações de uso da terra", "Carbono orgânico no solo", "Resíduos florestais",
            "Queimadas não associadas a desmatamento"}
AUTOPRODUCAO = "autoprodução de eletricidade, que na MIP tende a ficar no setor que a produz"
PC_NAO_S38 = ("não", "pelo SCN, produção de combustíveis é S19/S04/S05; só entraria no S38 pela hipótese em teste, "
                     "que não serve de critério")
MUT_DISC = "MUT: a conta-satélite pode não seguir a atribuição do SEEG a uma atividade (mesmo critério em U2 e U3)"


def sentido_unidade(tipo: str, chave: tuple[str, ...]) -> dict[str, tuple[str, str]]:
    """{setor: (nível, motivo)} para os setores com nível 'sim' ou 'discutível'; os demais são 'não'.

    HIPÓTESE deste diagnóstico: classificação feita pela descrição do SEEG e pela definição dos setores do SCN.
    """
    agro = ("S01", "S02", "S03")
    r: dict[str, tuple[str, str]] = {}
    if tipo == "U1":
        (a,) = chave
        regra = {
            "Agricultura": {"S01": ("sim", "atividade geral do SEEG = agricultura (inclui a MUT atribuída a ela)")},
            "Pecuária": {"S02": ("sim", "atividade geral do SEEG = pecuária (inclui a MUT atribuída a ela)")},
            "Silvicultura": {"S03": ("sim", "silvicultura = produção florestal")},
            "Vegetação nativa": {"S03": ("discutível", "MUT em vegetação nativa, sem atividade produtiva")},
            "Outros (geral)": {"S03": ("discutível", "MUT sem atividade identificada")},
            "Agropecuária (finalidade não identificada)": {
                s: ("discutível", "combustível da agropecuária, sem separar lavoura, pecuária e floresta") for s in agro},
            "Cimento": {"S26": ("sim", "cimento é minerais não metálicos")},
            "Outras matérias primas e indústrias": {
                "S26": ("discutível", "inclui cerâmica, mas também alimentos, papel, têxtil e mineração")},
            "Geração de eletricidade": {"S38": ("sim", "geração de eletricidade")},
            "Produção de combustíveis": {
                "S19": ("discutível", "no RS é quase só refino, mas inclui carvão mineral (S04)"),
                "S38": PC_NAO_S38},
            "Saneamento Básico": {"S39": ("sim", "saneamento (resíduos sólidos e efluentes domésticos)")},
            "Produção Industrial": {"S39": ("discutível", "efluentes industriais: o tratamento pode ser da indústria")},
            "Transporte de carga": {"S43": ("sim", "transporte de carga, quase todo rodoviário e ferroviário")},
            "Transporte de passageiros": {"S43": ("discutível", "inclui aéreo (S45) e automóveis das famílias")},
        }
        r = regra.get(a, {})
    elif tipo == "U2":
        s, c = chave
        if s == "Agropecuária":
            r = {"Cultivo de arroz": {"S01": ("sim", "lavoura de arroz")},
                 "Queima de resíduos agrícolas": {"S01": ("sim", "resíduos de lavoura")},
                 "Fermentação entérica": {"S02": ("sim", "rebanho")},
                 "Manejo de dejetos animais": {"S02": ("sim", "rebanho")},
                 "Solos manejados": {"S01": ("discutível", "mistura lavoura e pastagem"),
                                     "S02": ("discutível", "mistura lavoura e pastagem")}}.get(c, {})
        elif s == "Energia":
            r = {CAT_PUBLICO: {"S38": ("sim", "geração de serviço público")},
                 "Produção de combustíveis": {
                     "S19": ("discutível", "no RS é quase só refino, mas inclui carvão mineral (S04)"),
                     "S38": PC_NAO_S38},
                 "Transportes": {"S43": ("discutível", "inclui aéreo, hidroviário e automóveis das famílias")},
                 "Agropecuária": {x: ("discutível", "combustível da agropecuária, sem separar a atividade")
                                  for x in agro}}.get(c, {})
        elif s == "Mudança de Uso da Terra e Floresta":
            r = {x: ("discutível", "MUT de várias atividades, sem separar; " + MUT_DISC) for x in agro}
        elif s == "Processos Industriais":
            r = {"Produtos minerais": {"S26": ("sim", "cimento e cerâmica (processo)")}}.get(c, {})
        elif s == "Resíduos":
            if c == "Efluentes líquidos industriais":
                r = {"S39": ("discutível", "efluentes industriais: o tratamento pode ser da indústria")}
            else:
                r = {"S39": ("sim", "resíduos sólidos e efluentes domésticos")}
    elif tipo == "U3":
        a, c, sb = chave
        if c in MUT_CATS:  # mesmo critério do U2: MUT é 'discutível' para os três setores agropecuários
            r = {x: ("discutível", f"MUT atribuída pelo SEEG a '{a}'; " + MUT_DISC) for x in agro}
        elif a in ("Agricultura", "Pecuária"):
            alvo = "S01" if a == "Agricultura" else "S02"
            r = {alvo: ("sim", a.lower())}
        elif a == "Silvicultura":
            r = {"S03": ("sim", "silvicultura")}
        elif a in ("Vegetação nativa", "Outros (geral)"):
            r = {"S03": ("discutível", "MUT sem atividade produtiva identificada")}
        elif a == "Agropecuária (finalidade não identificada)":
            r = {x: ("discutível", "combustível da agropecuária, sem separar a atividade") for x in agro}
        elif a == "Cimento":
            r = {"S26": ("sim", "cimento")}
        elif a == "Outras matérias primas e indústrias" and sb in ("Cerâmica", "Produção de cerâmica"):
            r = {"S26": ("sim", "cerâmica")}
        elif a == ATIV_GE:
            if c == CAT_PUBLICO:
                r = {"S38": ("sim", "geração de serviço público")}
            elif c == ATIV_PC:
                r = {"S38": ("discutível", AUTOPRODUCAO + " (aqui, produção de combustíveis)")}
                if sb == SUB_REFINO:
                    r["S19"] = ("sim", "eletricidade gerada na refinaria")
            else:
                r = {"S38": ("discutível", AUTOPRODUCAO)}
                if sb in ("Cimento", "Cerâmica"):
                    r["S26"] = ("discutível", "autoprodução de eletricidade na indústria de " + sb.lower())
                if c == "Agropecuária":
                    r.update({x: ("discutível", "autoprodução de eletricidade na agropecuária") for x in agro})
        elif a == ATIV_PC:
            r = {"S38": PC_NAO_S38}
            if sb == SUB_REFINO:
                r["S19"] = ("sim", "refino de petróleo")
        elif a == "Saneamento Básico":
            r = {"S39": ("sim", "saneamento")}
        elif a == "Produção Industrial":
            r = {"S39": ("discutível", "efluentes industriais: o tratamento pode ser da indústria")}
        elif a == "Transporte de carga" and sb in ("Rodoviário", "Ferroviário"):
            r = {"S43": ("sim", "transporte terrestre de carga")}
        elif a == "Transporte de passageiros" and sb == "Rodoviário":
            r = {"S43": ("discutível", "rodoviário de passageiros inclui automóveis e motos das famílias")}
    return r


def unidades_rs(df: pd.DataFrame) -> pd.DataFrame:
    base = df[(df.metrica == GAS_AR5) & df[ERB].isin(EMISSAO) & (df.regiao == RS)]
    linhas = []
    for tipo, (desc, cols) in TIPOS_UNIDADE.items():
        g = base.groupby(cols)[COL_VALOR].sum()
        for chave, v in g.items():
            chave = chave if isinstance(chave, tuple) else (chave,)
            sentido = sentido_unidade(tipo, chave)
            reg = {"tipo": tipo, "descrição do tipo": desc, "unidade": " › ".join(chave), "valor_t": float(v),
                   "valor (mil t)": float(v) / 1e3, "usada na busca": float(v) != 0.0}
            for s in SETORES_CASAMENTO:
                nivel, motivo = sentido.get(s, ("não", ""))
                reg[f"sentido {s}"] = nivel if not motivo else f"{nivel}: {motivo}"
            linhas.append(reg)
    return pd.DataFrame(linhas)


def atividade_natural(unid: pd.DataFrame, nereus: pd.DataFrame) -> pd.DataFrame:
    ner = nereus[nereus.regiao == RS].set_index("codigo")
    u1 = unid[unid.tipo == "U1"].set_index("unidade").valor_t
    linhas = []
    for s, ativs in ATIV_NATURAL.items():
        v = float(sum(u1.get(a, 0.0) for a in ativs))
        alvo = float(ner.loc[s, "co2e_t"])
        linhas.append({"setor": s, "nome": ner.loc[s, "nome"], "Atividade geral do SEEG": " + ".join(ativs),
                       "SEEG RS (mil t)": v / 1e3, "NEREUS RS (mil t)": alvo / 1e3, "diferença (%)": _pct(v, alvo),
                       "nota": "sem emissão bruta no RS em 2019" if v == 0.0 else
                       ("hipótese do revisor (controle)" if s == "S38" else "")})
    return pd.DataFrame(linhas)


def _nivel(texto: str) -> str:
    return texto.split(":")[0]


def casamentos(unid: pd.DataFrame, nereus: pd.DataFrame) -> dict[str, pd.DataFrame]:
    ner = nereus[nereus.regiao == RS].set_index("codigo")
    usadas = unid[unid["usada na busca"]]
    combos = {}  # tipo → (matriz de índices por tamanho, valores)
    espaco = []
    for tipo, (desc, _) in TIPOS_UNIDADE.items():
        u = usadas[usadas.tipo == tipo].reset_index(drop=True)
        n = len(u)
        por_k = {}
        for k in TAMANHOS:
            idx = np.array(list(itertools.combinations(range(n), k)), dtype=np.int32).reshape(-1, k)
            por_k[k] = idx
        combos[tipo] = (u, por_k)
        espaco.append({"tipo": tipo, "unidade": desc, "unidades com valor ≠ 0": n,
                       "unidades com valor = 0 (fora)": int(((unid.tipo == tipo) & ~unid["usada na busca"]).sum()),
                       **{f"combinações de {k}": len(por_k[k]) for k in TAMANHOS},
                       "total de combinações": sum(len(v) for v in por_k.values())})
    espaco = pd.DataFrame(espaco)
    n_por_setor = int(espaco["total de combinações"].sum())

    achados, melhores, contagens = [], [], []
    for s in SETORES_CASAMENTO:
        alvo = float(ner.loc[s, "co2e_t"])
        cont = {"setor": s, "nome": ner.loc[s, "nome"], "NEREUS RS (mil t)": alvo / 1e3,
                "combinações testadas": n_por_setor}
        n01 = n1 = n5 = 0
        n01_por_nivel = {k: 0 for k in NIVEIS}
        melhor_geral, melhor_sim, melhor_nao_nao = None, None, None
        for tipo, (u, por_k) in combos.items():
            vals = u.valor_t.to_numpy()
            niveis_u = np.array([NIVEIS[_nivel(x)] for x in u[f"sentido {s}"]])
            for k, idx in por_k.items():
                soma = vals[idx].sum(axis=1)
                dif = (soma - alvo) / alvo * 100.0
                adif = np.abs(dif)
                nivel_c = niveis_u[idx].max(axis=1)
                n1 += int((adif < 1.0).sum())
                n5 += int((adif < 5.0).sum())
                hit = np.flatnonzero(adif < TOL_CASAMENTO)
                n01 += len(hit)
                for h in hit:
                    achados.append(_registro(s, ner, alvo, tipo, k, u, idx[h], soma[h], dif[h], nivel_c[h]))
                    n01_por_nivel[_nome_nivel(nivel_c[h])] += 1
                for crit, atual in (("geral", melhor_geral), ("sim", melhor_sim), ("nao_nao", melhor_nao_nao)):
                    mask = np.ones(len(adif), bool) if crit == "geral" else (
                        nivel_c == 0 if crit == "sim" else nivel_c <= 1)
                    if not mask.any():
                        continue
                    j = np.flatnonzero(mask)[np.argmin(adif[mask])]
                    if atual is None or adif[j] < abs(atual["diferença (%)"]):
                        novo = _registro(s, ner, alvo, tipo, k, u, idx[j], soma[j], dif[j], nivel_c[j])
                        if crit == "geral":
                            melhor_geral = novo
                        elif crit == "sim":
                            melhor_sim = novo
                        else:
                            melhor_nao_nao = novo
        for crit, reg in (("melhor de todas", melhor_geral), ("melhor com sentido 'sim'", melhor_sim),
                          ("melhor com sentido 'sim' ou 'discutível'", melhor_nao_nao)):
            if reg is not None:
                melhores.append({"critério": crit, **reg})
        nucleos = {(a["tipo de unidade"], a["núcleo (sem as desprezíveis)"]) for a in achados if a["setor"] == s}
        cont.update({"com |dif| < 0.1%": n01, "dos quais 'sim'": n01_por_nivel["sim"],
                     "dos quais 'discutível'": n01_por_nivel["discutível"], "dos quais 'não'": n01_por_nivel["não"],
                     "núcleos distintos em |dif| < 0.1%": len(nucleos),
                     "com |dif| < 1%": n1, "com |dif| < 5%": n5,
                     ACASO: n5 / 50.0})
        contagens.append(cont)
    achados = pd.DataFrame(achados)
    if not achados.empty:
        achados = achados.sort_values(["setor", "sentido econômico", "|diferença| (%)"],
                                      key=lambda c: c.map(NIVEIS) if c.name == "sentido econômico" else c)
    contagens = pd.DataFrame(contagens)
    tot = contagens.drop(columns=["setor", "nome", "NEREUS RS (mil t)"]).sum()
    contagens = pd.concat([contagens, pd.DataFrame([{"setor": "Total", "nome": "8 setores", **tot.to_dict()}])],
                          ignore_index=True)
    return {"espaco": espaco, "contagens": contagens, "achados": achados,
            "melhores": pd.DataFrame(melhores), "n_por_setor": n_por_setor}


def _nome_nivel(i: int) -> str:
    return {v: k for k, v in NIVEIS.items()}[int(i)]


HIPOTESE_REVISOR = ("S38", "U1", f"{ATIV_GE} + {ATIV_PC}")  # formulada antes da busca (Seção 3)
TOL_DESPREZIVEL = 0.1  # % do alvo: unidade menor que isso não muda o casamento a ±0.1%


def _registro(setor, ner, alvo, tipo, k, u, idx, soma, dif, nivel) -> dict:
    nomes = [u.unidade[i] for i in idx]
    sentidos = [u[f"sentido {setor}"][i] for i in idx]
    desprez = [abs(float(u.valor_t[i])) < TOL_DESPREZIVEL / 100.0 * alvo for i in idx]
    nucleo = " + ".join(n for n, d in zip(nomes, desprez) if not d) or "(só unidades desprezíveis)"
    nivel = _nome_nivel(nivel)
    nome_setor = ner.loc[setor, "nome"]
    hipotese = (setor, tipo, " + ".join(nomes)) == HIPOTESE_REVISOR
    if hipotese:
        just = ("Não pelo critério SCN (refino, carvão mineral e exploração são S19/S04/S05), mas é a hipótese do "
                "revisor, formulada antes da busca: o peso dela vem disso, não do rótulo de sentido econômico.")
    elif nivel == "sim":
        just = f"Sim: todas as unidades são atividades do {setor} ({nome_setor}) — " + "; ".join(
            f"{n}: {t.split(': ', 1)[1]}" for n, t in zip(nomes, sentidos)) + "."
    elif nivel == "discutível":
        just = f"Discutível: nenhuma unidade é alheia ao {setor}, mas " + "; ".join(
            f"'{n}' ({t.split(': ', 1)[1]})" for n, t in zip(nomes, sentidos) if _nivel(t) == "discutível") + "."
    else:
        just = ("Não: '" + "', '".join(n for n, t in zip(nomes, sentidos) if _nivel(t) == "não")
                + f"' não pertence(m) ao {setor} ({nome_setor}); coincidência numérica.")
    return {"setor": setor, "nome do setor": nome_setor, "NEREUS RS (t)": alvo, "tipo de unidade": tipo,
            "nº de unidades": k, "combinação": " + ".join(nomes), "SEEG RS (t)": float(soma),
            "diferença (%)": float(dif), "|diferença| (%)": abs(float(dif)), "sentido econômico": nivel,
            "hipótese formulada antes da busca": "sim" if hipotese else "não",
            f"unidades desprezíveis (|valor| < {TOL_DESPREZIVEL}% do alvo)": int(sum(desprez)),
            "núcleo (sem as desprezíveis)": nucleo, "justificativa": just}


# --------------------------------------------------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------------------------------------------------
NEGRITO = Font(bold=True)


def escrever_blocos(xw: pd.ExcelWriter, aba: str, blocos: list[tuple[str, pd.DataFrame | list[str]]]) -> None:
    """Escreve vários blocos (título + tabela ou lista de linhas de texto) numa aba, um abaixo do outro."""
    linha = 0
    larguras: dict[int, int] = {}
    for titulo, conteudo in blocos:
        if isinstance(conteudo, list):
            conteudo = pd.DataFrame({titulo: conteudo})
            conteudo.to_excel(xw, sheet_name=aba, startrow=linha, index=False)
            ws = xw.sheets[aba]
            ws.cell(row=linha + 1, column=1).font = NEGRITO
            larguras[1] = max(larguras.get(1, 0), 140)
            linha += len(conteudo) + 2
            continue
        pd.DataFrame({titulo: []}).to_excel(xw, sheet_name=aba, startrow=linha, index=False)
        conteudo.to_excel(xw, sheet_name=aba, startrow=linha + 1, index=False)
        ws = xw.sheets[aba]
        ws.cell(row=linha + 1, column=1).font = Font(bold=True, size=12)
        for j, col in enumerate(conteudo.columns, start=1):
            c = ws.cell(row=linha + 2, column=j)
            c.font = NEGRITO
            c.alignment = Alignment(wrap_text=True, vertical="top")
            serie = conteudo[col]
            larg = min(60, max(10, len(str(col)) // 2 + 2,
                               *(len(str(x)) + 2 for x in serie.head(200)) if len(serie) else [10]))
            larguras[j] = max(larguras.get(j, 0), larg)
            if pd.api.types.is_float_dtype(serie):
                fmt = "0.000" if "%" in str(col) else "0.0"
                for i in range(len(serie)):
                    ws.cell(row=linha + 3 + i, column=j).number_format = fmt
        linha += len(conteudo) + 4
    ws = xw.sheets[aba]
    from openpyxl.utils import get_column_letter
    for j, w in larguras.items():
        ws.column_dimensions[get_column_letter(j)].width = w


# --------------------------------------------------------------------------------------------------------------------
# Principal
# --------------------------------------------------------------------------------------------------------------------
def main() -> None:
    t0 = time.time()
    bruto, origem = ler_seeg_2019()
    df = preparar(bruto)
    print(f"SEEG 13.0: {len(df)} linhas de {ANO} ({origem})")
    erbs_arquivo = sorted(df[ERB].unique())
    previstos = set().union(*BASES.values())
    if set(erbs_arquivo) - previstos:
        raise ValueError(f"valores não previstos em '{ERB}': {set(erbs_arquivo) - previstos}")
    nereus = ler_nereus()

    # --- Seção 3: recorte, tabela, E4 ---
    rec = recorte_energia(df)
    SAIDA_RECORTE.parent.mkdir(parents=True, exist_ok=True)
    rec[rec[COL_VALOR] != 0.0].to_csv(SAIDA_RECORTE, index=False)  # linhas com valor 0 não mudam nenhuma soma
    tab, detalhe = tabela_componentes(rec, nereus)
    e4, e4_conf = componentes_e4(rec)
    e4.to_csv(SAIDA_E4, index=False)

    # --- 3.1 item 2: métricas ---
    longo, mais_perto, erb = tabela_metricas(df, nereus)

    # --- 3.1 item 3: casamentos ---
    unid = unidades_rs(df)
    cas = casamentos(unid, nereus)
    n_total = cas["n_por_setor"] * len(SETORES_CASAMENTO)
    natural = atividade_natural(unid, nereus)

    # --- números-chave ---
    t = tab.set_index("codigo")
    soma_rs, soma_rb = t.loc["soma", "RS (t)"], t.loc["soma", "RB estados (t)"]
    ner_rs, ner_rb = t.loc["nereus", "RS (t)"], t.loc["nereus", "RB estados (t)"]
    difs = tab.filter(like="diferença (mil t)")
    max_dif = float(np.nanmax(np.abs(difs.to_numpy(dtype=float))))
    ach = cas["achados"]
    cont = cas["contagens"].set_index("setor")
    vazio = ach.iloc[0:0] if not ach.empty else pd.DataFrame(columns=["setor", "sentido econômico", "combinação",
                                                                     "tipo de unidade", "diferença (%)"])
    ach = ach if not ach.empty else vazio
    ach_sim = ach[ach["sentido econômico"] == "sim"]
    ach_disc = ach[ach["sentido econômico"] == "discutível"]
    mel = cas["melhores"]
    ner_tot_rs = float(nereus.loc[nereus.regiao == RS, "co2e_t"].sum())
    ner_tot_br = float(nereus.co2e_t.sum())

    def _f(x, d=1):
        return f"{x:.{d}f}"

    def _dif(agr, met, base):
        r = longo[(longo.agregado == agr) & (longo.metrica == met) & (longo.base == base)]
        return float(r["diferença SEEG − NEREUS (%)"].iloc[0])

    def _seeg(agr, met, base):
        r = longo[(longo.agregado == agr) & (longo.metrica == met) & (longo.base == base)]
        return float(r["SEEG (t)"].iloc[0])

    s38rs = "; ".join(f"{m.replace('CO2e (t) ', '')} {_f(_dif('S38-RS', m, BRUTA), 2)}%" for m in METRICAS)
    s38rb = "; ".join(f"{m.replace('CO2e (t) ', '')} {_f(_dif('S38-RB (estados)', m, BRUTA), 2)}%"
                      for m in METRICAS)
    s38_nci_igual = all(np.isclose(_seeg(a, m, BRUTA), _seeg(a, m, BRUTA_NCI), rtol=1e-12, atol=1e-3)
                        for a in ("S38-RS", "S38-RB (estados)") for m in METRICAS)

    def _mais_perto(agr, grupo):
        r = mais_perto[(mais_perto.agregado == agr) & (mais_perto["bases consideradas"] == grupo)].iloc[0]
        emp = r["empatadas (mesmo valor)"]
        emp = "" if emp == "—" else f"; empata com {emp.replace('CO2e (t) ', '')}"
        return f"{r['métrica mais perto'].replace('CO2e (t) ', '')} {r.base} ({_f(r['diferença (%)'], 2)}%{emp})"

    def _melhor(setor, crit):
        r = mel[(mel.setor == setor) & (mel["critério"] == crit)]
        if r.empty:
            return "nenhuma unidade com esse sentido"
        r = r.iloc[0]
        return f"{r['combinação']} ({_f(r['diferença (%)'], 2)}%)"

    gtp_iguais = bool(np.allclose(
        df[df.metrica == "CO2e (t) GTP-AR2"].set_index(COLUNAS_TEXTO[:1] + COLUNAS_TEXTO[2:])[COL_VALOR].sort_index(),
        df[df.metrica == "CO2e (t) GTP-AR4"].set_index(COLUNAS_TEXTO[:1] + COLUNAS_TEXTO[2:])[COL_VALOR].sort_index(),
        rtol=0, atol=1e-9))
    acaso_total = float(cont.loc["Total", ACASO])
    nucleos_total = int(cont.loc["Total", "núcleos distintos em |dif| < 0.1%"])
    outros = [s for s in SETORES_CASAMENTO if s != "S38"]
    melhores_sim = "; ".join(f"{s}: {_melhor(s, chr(39).join(['melhor com sentido ', 'sim', '']))}" for s in outros)
    n_agri = int(((unid.tipo == "U3") & unid["usada na busca"] & unid.unidade.str.startswith("Agricultura › ")).sum())
    m01 = mel[(mel.setor == "S01") & (mel["critério"] == "melhor com sentido 'sim'")]
    n_s01 = int(m01["nº de unidades"].iloc[0]) if not m01.empty else 0
    hip = ach[(ach.setor == HIPOTESE_REVISOR[0]) & (ach["tipo de unidade"] == HIPOTESE_REVISOR[1])
              & (ach["combinação"] == HIPOTESE_REVISOR[2])]
    hip_dif = float(hip["diferença (%)"].iloc[0]) if not hip.empty else _pct(soma_rs, ner_rs)

    # Casamentos concorrentes: outros setores com |dif| < 0.1% que usam o refino de 'Produção de combustíveis'
    u_refino = f"{ATIV_PC} › {ATIV_PC} › {SUB_REFINO}"
    u3 = unid[unid.tipo == "U3"].set_index("unidade").valor_t
    refino_rs = float(sum(v for n, v in u3.items() if SUB_REFINO in n))  # PC › refino + GE › PC › refino
    unid_refino = [n for n in u3.index if SUB_REFINO in n or re.search(r"\bcoque", n, re.IGNORECASE)]
    conc = ach[(ach.setor != "S38") & ach["combinação"].str.contains(u_refino, regex=False)]
    conc_s19 = conc[conc.setor == "S19"].sort_values("|diferença| (%)") if not conc.empty else conc
    ner_s19 = float(nereus.loc[(nereus.regiao == RS) & (nereus.codigo == "S19"), "co2e_t"].iloc[0])
    e4_rs_s19 = float(e4.loc[(e4.regiao == RS) & (e4.setor_destino == "S19"), "valor_t"].sum())
    if not conc_s19.empty:
        c0 = conc_s19.iloc[0]
        txt_conc = (f"A busca também encontrou S19-RS ≈ {c0['combinação']} ({_f(float(c0['diferença (%)']), 3)}%, "
                    f"sentido '{c0['sentido econômico']}'), que usa as mesmas emissões do refino; ao todo, "
                    f"{len(conc)} combinações abaixo de 0.1% em outros setores usam '{u_refino}' "
                    f"({len(conc_s19)} no S19). ")
    else:
        txt_conc = f"Nenhuma combinação abaixo de 0.1% de outro setor usa '{u_refino}'. "
    txt_conc += (
        f"As duas leituras se excluem: as {_f(u3.get(u_refino, 0.0) / 1e3)} mil t do refino não podem estar ao "
        f"mesmo tempo no S38 e no S19. E, se o refino estiver no S38, as {_f(ner_s19 / 1e3)} mil t do S19-RS "
        f"(Refino de petróleo e coquerias) da NEREUS não têm contrapartida no SEEG do RS, cujas únicas unidades de "
        f"refino ou coqueria são {'; '.join(unid_refino)} ({_f(refino_rs / 1e3)} mil t, ambas usadas na hipótese "
        f"do S38). Isso enfraquece o indício e é um risco para a E4: se o S19 da NEREUS já contiver o refino, a "
        f"E4 o conta duas vezes no S19-RS ({_f(ner_s19 / 1e3)} + {_f(e4_rs_s19 / 1e3)} = "
        f"{_f((ner_s19 + e4_rs_s19) / 1e3)} mil t); vale tratá-lo como sensibilidade.")

    # Os 'discutível' abaixo de 0.1%: por setor e a unidade comum a todos (se houver)
    txt_disc = ""
    for s_, g in ach_disc.groupby("setor"):
        comuns = set.intersection(*(set(c.split(" + ")) for c in g["combinação"]))
        txt_disc += f" {len(g)} 'discutível' no {s_}"
        for cu in sorted(comuns):
            v = float(unid.loc[unid.unidade == cu, "valor_t"].iloc[0])
            alvo = float(nereus.loc[(nereus.regiao == RS) & (nereus.codigo == s_), "co2e_t"].iloc[0])
            txt_disc += (f", todas com '{cu}' ({_f(v / 1e3)} mil t, {_f(_pct(v, alvo), 1)}% do {s_} sozinha) "
                         "mais unidades pequenas: só MUT e fragmentos, não a atividade")
        txt_disc += "."

    e4_div = e4_conf[e4_conf.nota.str.startswith("a referência")]
    txt_e4 = (" Regra da E4: " + "; ".join(f"{r.regiao} → {r.setor_destino}: exato {_f(r._5, 2)} mil t contra "
                                          f"{r._6} das instruções ({r.nota})" for r in e4_div.itertuples())
              + "." if not e4_div.empty else " Regra da E4: as 5 linhas conferem com as referências das instruções.")

    tot = {(a, b): _dif(a, GAS_AR5, b) for a in ("Total RS", TOT_BR, TOT_BR_UF) for b in BASES}
    conclusao = [
        f"1. Reprodução da tabela da Seção 3: todos os componentes conferem com os das instruções (maior diferença: "
        f"{_f(max_dif, 2)} mil t, arredondamento a 0.1 mil t). Soma 'Geração de eletricidade' + 'Produção de "
        f"combustíveis' (GWP-AR5, Emissão + Emissão NCI, que no S38 é igual à bruta do SEEG): RS {_f(soma_rs / 1e3)} "
        f"mil t contra {_f(ner_rs / 1e3)} no S38-RS da NEREUS ({_f(_pct(soma_rs, ner_rs), 3)}%); RB estados "
        f"{_f(soma_rb / 1e3)} mil t contra {_f(ner_rb / 1e3)} ({_f(_pct(soma_rb, ner_rb), 2)}%). Não alocado: "
        f"{_f(t.loc['soma', 'Não alocado (t)'] / 1e3)} mil t." + txt_e4,
        f"2. Métricas no S38 (bruta; {'igual com ou sem NCI e ' if s38_nci_igual else ''}igual à líquida, sem "
        f"remoções nessas atividades). S38-RS: {s38rs}. A GWP-AR5 é a mais perto no S38-RS; nenhuma métrica aproxima "
        f"o S38-RB (estados): {s38rb}."
        + (" GTP-AR2 e GTP-AR4 têm valores idênticos linha a linha no arquivo: empatam em todo agregado." if gtp_iguais
           else ""),
        f"3. Totais. Bases principais = convenção do próprio SEEG (a da pivot 'Consulta Série Histórica' do arquivo): "
        f"bruta = 'Emissão'; líquida = 'Emissão' + 'Remoção'. Com GWP-AR5: total RS bruta "
        f"{_f(tot[('Total RS', BRUTA)], 2)}%, líquida {_f(tot[('Total RS', LIQUIDA)], 2)}%; total Brasil (26 UF + RS "
        f"+ Não Alocado) bruta {_f(tot[(TOT_BR, BRUTA)], 2)}%, líquida {_f(tot[(TOT_BR, LIQUIDA)], 2)}%. As variantes "
        f"com NCI (emissões e remoções não contabilizadas no inventário: solos manejados e queimadas não associadas a "
        f"desmatamento) dão: RS bruta com NCI {_f(tot[('Total RS', BRUTA_NCI)], 2)}%, líquida com NCI "
        f"{_f(tot[('Total RS', LIQUIDA_NCI)], 2)}%; Brasil bruta com NCI {_f(tot[(TOT_BR, BRUTA_NCI)], 2)}%, "
        f"líquida com NCI {_f(tot[(TOT_BR, LIQUIDA_NCI)], 2)}%. Na convenção do SEEG, a métrica mais perto é: total RS "
        f"{_mais_perto('Total RS', 'convenção SEEG (bruta e líquida)')}; total Brasil "
        f"{_mais_perto(TOT_BR, 'convenção SEEG (bruta e líquida)')}. O total Brasil depende de incluir o 'Não "
        f"Alocado' (inclusão é a escolha defensável, porque RS + RB cobre o país): só com as UF, GWP-AR5 líquida "
        f"{_f(tot[(TOT_BR_UF, LIQUIDA)], 2)}% e a mais perto passa a "
        f"{_mais_perto(TOT_BR_UF, 'convenção SEEG (bruta e líquida)')}. E a 'mais perto' depende do conjunto de "
        f"bases testado: somando o Bunker (todo no 'Não Alocado') à líquida, o total Brasil vai a "
        f"{_mais_perto(TOT_BR, 'todas (inclui líquida + Bunker, sensibilidade)')}, sem mudar o RS "
        f"({_mais_perto('Total RS', 'todas (inclui líquida + Bunker, sensibilidade)')}). Com ~8 métricas × 5 bases × "
        "2 escopos, aproximações de totais abaixo de 1% aparecem com facilidade: são coincidências possíveis, não "
        "evidência de qual base ou versão do SEEG a conta-satélite usou; a distância (0.1% a 1%) é muito maior que a "
        f"do S38-RS ({_f(abs(_pct(soma_rs, ner_rs)), 2)}%).",
        f"4. Casamentos no RS: {cas['n_por_setor']} combinações por setor (1, 2 ou 3 unidades de um mesmo tipo) × "
        f"{len(SETORES_CASAMENTO)} setores (7 + o controle S38) = {n_total} testes. Com |dif| < 0.1%: {len(ach)} "
        f"combinações ({nucleos_total} núcleos distintos depois de tirar as unidades desprezíveis, |valor| < 0.1% do "
        f"alvo, que geram quase repetições). Por acaso seriam esperadas da ordem de {_f(acaso_total, 0)} (heurística de "
        "densidade uniforme: nº em ±5% ÷ 50; por setor erra até ~3 vezes, ver contagens). Pelo critério SCN, aplicado "
        f"igual a todos os setores, {len(ach_sim)} têm sentido econômico 'sim' e {len(ach_disc)} 'discutível'; as "
        "demais não têm sentido econômico." + txt_disc,
        f"5. No controle S38, a combinação mais perto de todas é {_melhor('S38', 'melhor de todas')}, sem sentido "
        "econômico: com esse número de testes, casamentos melhores que 0.1% aparecem por acaso. O casamento do S38-RS "
        f"com 'Geração de eletricidade + Produção de combustíveis' ({_f(hip_dif, 3)}%) também é 'não' pelo critério "
        "SCN (refino e carvão mineral são S19/S04): o peso dele vem só de ter sido formulado antes da busca (hipótese "
        "do revisor) e de o SEEG pôr refino e mineração de carvão em 'Produção de combustíveis', não de ser o mais "
        "exato nem do rótulo de sentido econômico (classificá-lo como 'discutível' seria circular).",
        "6. Casamento concorrente. " + txt_conc,
        "7. O padrão 'atividade do SEEG → setor da MIP' não se repete nos outros setores do RS. Atividade geral "
        "correspondente × NEREUS: " + "; ".join(f"{r.setor} = {r._3} ({_f(r._6, 1)}%)" for r in natural.itertuples()
                                               if r.setor != "S38") + ". Melhor combinação com sentido 'sim' por "
        f"setor (busca completa): {melhores_sim}. A do S01 junta {n_s01} das {n_agri} unidades U3 da agricultura, "
        "sem critério que as separe das demais: é um fragmento, não a atividade.",
        "8. Conclusão: o S38-RS da base NEREUS é compatível com um S38 que abranja as emissões que o SEEG 13.0 "
        "(GWP-AR5, bruta) atribui a 'Geração de eletricidade' e 'Produção de combustíveis' (inclusive refino e carvão "
        "mineral). É um indício fraco, compatível com a hipótese do revisor, não uma prova: não temos a documentação "
        f"da conta-satélite; o RB não fecha com o mesmo critério ({_f(_pct(soma_rb, ner_rb), 2)}%); os totais do RS e "
        "do Brasil só se aproximam conforme a base e o escopo escolhidos; os outros setores do RS não mostram o mesmo "
        f"padrão; com {n_total} combinações testadas, coincidências abaixo de 0.1% são esperadas; e há um casamento "
        "concorrente (S19 com o mesmo refino) incompatível com ele, além do S19-RS sem contrapartida no SEEG se o "
        "refino estiver no S38.",
    ]

    resumo = pd.DataFrame([
        ("S38-RS: SEEG GE + PC (GWP-AR5, bruta)", soma_rs / 1e3, "mil t"),
        ("S38-RS: NEREUS", ner_rs / 1e3, "mil t"),
        ("S38-RS: diferença", _pct(soma_rs, ner_rs), "%"),
        ("S38-RB: SEEG GE + PC, estados (GWP-AR5, bruta)", soma_rb / 1e3, "mil t"),
        ("S38-RB: NEREUS", ner_rb / 1e3, "mil t"),
        ("S38-RB: diferença", _pct(soma_rb, ner_rb), "%"),
        ("Não alocado (GE + PC)", t.loc["soma", "Não alocado (t)"] / 1e3, "mil t"),
        ("Total RS: NEREUS", ner_tot_rs / 1e3, "mil t"),
        ("Total RS: SEEG GWP-AR5 bruta (SEEG: Emissão) − NEREUS", tot[("Total RS", BRUTA)], "%"),
        ("Total RS: SEEG GWP-AR5 líquida (SEEG: Emissão + Remoção) − NEREUS", tot[("Total RS", LIQUIDA)], "%"),
        ("Total RS: SEEG GWP-AR5 bruta com NCI − NEREUS (variante)", tot[("Total RS", BRUTA_NCI)], "%"),
        ("Total RS: SEEG GWP-AR5 líquida com NCI − NEREUS (variante)", tot[("Total RS", LIQUIDA_NCI)], "%"),
        ("Total Brasil: NEREUS", ner_tot_br / 1e3, "mil t"),
        ("Total Brasil: SEEG GWP-AR5 bruta (SEEG: Emissão; 26 UF + RS + Não Alocado) − NEREUS", tot[(TOT_BR, BRUTA)],
         "%"),
        ("Total Brasil: SEEG GWP-AR5 líquida (SEEG: Emissão + Remoção; 26 UF + RS + Não Alocado) − NEREUS",
         tot[(TOT_BR, LIQUIDA)], "%"),
        ("Total Brasil: SEEG GWP-AR5 bruta com NCI (26 UF + RS + Não Alocado) − NEREUS (variante)",
         tot[(TOT_BR, BRUTA_NCI)], "%"),
        ("Total Brasil: SEEG GWP-AR5 líquida com NCI (26 UF + RS + Não Alocado) − NEREUS (variante)",
         tot[(TOT_BR, LIQUIDA_NCI)], "%"),
        ("Total Brasil: SEEG GWP-AR5 líquida + Bunker − NEREUS (sensibilidade)", tot[(TOT_BR, LIQUIDA_BUNKER)], "%"),
        ("Total Brasil só UF (sem Não Alocado): SEEG GWP-AR5 bruta (SEEG) − NEREUS", tot[(TOT_BR_UF, BRUTA)], "%"),
        ("Total Brasil só UF (sem Não Alocado): SEEG GWP-AR5 líquida (SEEG) − NEREUS", tot[(TOT_BR_UF, LIQUIDA)], "%"),
        *[(f"E4 {r.regiao} → {r.setor_destino}: {r.componente}", r.valor_t / 1e3, "mil t") for r in e4.itertuples()],
        ("Casamentos: combinações por setor", cas["n_por_setor"], "nº"),
        ("Casamentos: setores testados (com o controle S38)", len(SETORES_CASAMENTO), "nº"),
        ("Casamentos: total de testes", n_total, "nº"),
        ("Casamentos: |dif| < 0.1% (todos)", len(ach), "nº"),
        ("Casamentos: |dif| < 0.1%, núcleos distintos (sem unidades desprezíveis)", nucleos_total, "nº"),
        ("Casamentos: |dif| < 0.1% esperados por acaso (ordem de grandeza)", acaso_total, "nº"),
        ("Casamentos: |dif| < 0.1% com sentido 'sim' (critério SCN)", len(ach_sim), "nº"),
        ("Casamentos: |dif| < 0.1% 'discutível' (critério SCN)", len(ach_disc), "nº"),
        ("Casamentos: S38 = GE + PC (hipótese formulada antes da busca), diferença", hip_dif, "%"),
        *[(f"Casamento concorrente: S19 = {c0['combinação']}", float(c0["diferença (%)"]), "%")
          for _, c0 in conc_s19.head(1).iterrows()],
        ("S19-RS: NEREUS", ner_s19 / 1e3, "mil t"),
        ("S19-RS: NEREUS + E4 RS → S19 (risco de dupla contagem do refino)", (ner_s19 + e4_rs_s19) / 1e3, "mil t"),
    ], columns=["item", "valor", "unidade"])

    hoje = dt.date.today().isoformat()
    leiame = [
        "Diagnóstico da conta-satélite de emissões do S38: NEREUS × SEEG 13.0 (Rodada 05, Seção 3.1).",
        f"Gerado por scripts/rodada05_diagnostico_seeg.py em {hoje}. Não muda a MIP: só lê o SEEG e a base NEREUS.",
        f"Fonte SEEG: SEEG 13.0, Dados nacionais, aba 'Dados' — {URL_SEEG}; SHA-256 {SHA_SEEG} (conferido nesta "
        "execução; registro em dados/externos/seeg/DOWNLOAD.txt). Leitura: " + origem + ". O cache tem os rótulos de "
        "texto sem os espaços das pontas (no bruto, 'Efluentes líquidos industriais ', 'Disposição final ' e "
        "'Alimentos ' têm espaço no fim); nenhuma soma muda.",
        "Fonte NEREUS: MIP original (dados/original/, intocável), linha de CO2e (t) da conta-satélite, por setor-região.",
        f"Critérios (Seção 3): ano {ANO}; Gás = '{GAS_AR5}'; 'Emissão/Remoção/Bunker' ∈ {{Emissão, Emissão NCI}}; "
        "Estado = 'Rio Grande do Sul' → RS; 'Não Alocado' → Não alocado (separado); demais 26 UF → RB (estados).",
        "S38 pelo critério do revisor: 'Atividade geral' ∈ {Geração de eletricidade, Produção de combustíveis}. "
        f"Recorte: dados/externos/seeg/{SAIDA_RECORTE.name} ({int((rec[COL_VALOR] != 0).sum())} linhas com valor ≠ 0).",
        "componentes_S38: a tabela da Seção 3 reproduzida (t e mil t), com a referência das instruções e a diferença "
        "(meu − instrução, mil t), e a conferência da regra da E4. componentes_S38_detalhe: a mesma soma no nível "
        "Atividade › Categoria › Sub-categoria.",
        "E4: valor_t é sempre o valor exato do SEEG. A referência 18088.8 mil t das instruções para o refino do RB é a "
        "soma das parcelas já arredondadas (17329.1 + 759.7); o exato é 18088.87 (18088.9 arredondado)."
        if (e4_conf.regiao.eq(RB) & e4_conf.setor_destino.eq("S19") & e4_conf.nota.str.startswith("a referência")).any()
        else "E4: valor_t é sempre o valor exato do SEEG; as referências das instruções conferem.",
        f"metricas: Seção 3.1 item 2. Bases principais = convenção do próprio SEEG (pivot 'Consulta Série Histórica' do "
        f"arquivo): '{BRUTA}' = Emissão; '{LIQUIDA}' = Emissão + Remoção. Variantes: '{BRUTA_NCI}' = Emissão + Emissão "
        f"NCI (o critério da Seção 3; no S38 não há NCI, então coincide com a bruta do SEEG); '{LIQUIDA_NCI}' = as "
        "quatro. NCI = emissões e remoções não contabilizadas no inventário (no arquivo: solos manejados — redução e "
        "aumento do estoque de C no solo, queima de pasto — e queimadas não associadas a desmatamento). Bunker "
        f"(combustível de navegação e aviação internacionais, todo em 'Não Alocado') fica fora, salvo na sensibilidade "
        f"'{LIQUIDA_BUNKER}'. Rótulo 'CO2e (t) GTP AR4' (736 linhas de Resíduos) tratado como GTP-AR4.",
        f"metricas, escopo: '{TOT_BR}' = todas as linhas (26 UF + RS + 'Não Alocado'; Bunker fora), comparado ao total "
        f"da NEREUS (RS + RB, o país todo); '{TOT_BR_UF}' é a sensibilidade sem o 'Não Alocado'. S38-RB: 'estados' "
        "(sem Não alocado) e 'estados + Não alocado'. Diferença = (SEEG − NEREUS)/NEREUS. A 'métrica mais perto' é "
        "condicional ao conjunto de bases testado (três conjuntos na tabela) e registra os empates exatos (GTP-AR2 = "
        "GTP-AR4).",
        "casamentos_RS: Seção 3.1 item 3. Espaço de busca: somas de 1, 2 ou 3 unidades do SEEG do RS (GWP-AR5, Emissão + "
        "Emissão NCI, o critério da Seção 3) dentro de cada tipo de unidade (U1 = Atividade geral; U2 = Setor de "
        "emissão › Categoria emissora; U3 = Atividade geral › Categoria emissora › Sub-categoria emissora); unidades "
        "com valor = 0 ficam fora. Tipos não são misturados (evita dupla contagem). Setores: S01, S02, S03, S19, S26, "
        "S39, S43 e S38 (controle).",
        "Sentido econômico: HIPÓTESE deste diagnóstico — cada unidade recebe 'sim', 'discutível' ou 'não' para cada "
        "setor (aba unidades_RS, com o motivo); a combinação recebe o pior nível entre as suas unidades. Critério SCN "
        "igual para todos os setores: 'Produção de combustíveis' é 'não' para o S38 (seria circular dar 'discutível' "
        "pela própria hipótese em teste); a hipótese do revisor é marcada na coluna 'hipótese formulada antes da "
        "busca'. MUT (alterações de uso da terra, carbono no solo, resíduos florestais, queimadas) é 'discutível' para "
        "S01, S02 e S03 em U2 e em U3.",
        "Contagens: 'esperado por acaso' (nº em ±5% ÷ 50) é só ordem de grandeza (densidade uniforme); por setor pode "
        f"errar várias vezes. Unidades desprezíveis (|valor| < {TOL_DESPREZIVEL}% do alvo) geram combinações quase "
        "repetidas: a coluna 'núcleo' tira essas unidades e a contagem 'núcleos distintos' mede os achados efetivos.",
        "Casamento concorrente: " + txt_conc,
        "CAUTELA: não temos a documentação da conta-satélite da NEREUS. Casamentos numéricos são indícios, não provas; "
        f"com {n_total} combinações testadas, coincidências abaixo de 0.1% são esperadas por acaso.",
        "GTP-AR2 e GTP-AR4: valores idênticos linha a linha no arquivo " + ("(conferido)." if gtp_iguais else
                                                                           "— NÃO conferiu nesta execução."),
        f"Recorte salvo sem as linhas de valor 0 ({int((rec[COL_VALOR] == 0).sum())} de {len(rec)}); as somas não mudam.",
        "Convenção: ponto decimal, sem separador de milhar.",
    ]

    SAIDA_XLSX.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(SAIDA_XLSX, engine="openpyxl") as xw:
        escrever_blocos(xw, "LEIAME", [("LEIAME", leiame)])
        escrever_blocos(xw, "componentes_S38", [("Tabela da Seção 3 reproduzida (GWP-AR5, bruta, 2019)", tab),
                                                ("Regra da E4 (Seção 3.2): valor exato × referência das instruções",
                                                 e4_conf)])
        escrever_blocos(xw, "componentes_S38_detalhe", [("Detalhe por Atividade › Categoria › Sub-categoria", detalhe)])
        escrever_blocos(xw, "metricas", [
            ("Escopo e bases", [
                f"Bases principais (convenção do SEEG): {BRUTA} = Emissão; {LIQUIDA} = Emissão + Remoção.",
                f"Variantes: {BRUTA_NCI} = Emissão + Emissão NCI (critério da Seção 3); {LIQUIDA_NCI} = as quatro; "
                f"{LIQUIDA_BUNKER} = sensibilidade.",
                f"{TOT_BR}: todas as linhas, inclusive 'Não Alocado' (onde está todo o Bunker); {TOT_BR_UF}: sem o "
                "'Não Alocado'. O alvo NEREUS dos dois é o total RS + RB.",
                "A 'métrica mais perto' vale só dentro do conjunto de bases indicado; empates exatos na coluna "
                "'empatadas'.",
            ]),
            ("Métrica e base mais perto da NEREUS, por agregado e conjunto de bases", mais_perto),
            ("Totais por métrica e base (SEEG 13.0, 2019) × NEREUS", longo),
            ("Valores de 'Emissão/Remoção/Bunker' no arquivo e onde entram", erb),
        ])
        escrever_blocos(xw, "casamentos_RS", [
            ("Teste direto: Atividade geral do SEEG correspondente a cada setor (GWP-AR5, bruta, RS)", natural),
            ("Espaço de busca (por setor)", cas["espaco"]),
            ("Contagens por setor", cas["contagens"]),
            ("Todas as combinações com |diferença| < 0.1%", cas["achados"] if not cas["achados"].empty
             else pd.DataFrame({"resultado": ["nenhuma"]})),
            ("Melhor combinação de cada setor (mesmo acima de 0.1%)", cas["melhores"]),
        ])
        escrever_blocos(xw, "unidades_RS", [("Unidades do SEEG no RS (GWP-AR5, bruta) e sentido econômico por setor "
                                             "(HIPÓTESE do diagnóstico)", unid)])
        escrever_blocos(xw, "resumo", [("Números-chave", resumo), ("Conclusão", conclusao)])

    # --- console ---
    pd.set_option("display.width", 220)
    pd.set_option("display.max_columns", 30)
    pd.set_option("display.max_colwidth", 90)
    print("\nTabela da Seção 3 (mil t):")
    cols = ["codigo", "RS (mil t)", "RB estados (mil t)", "Não alocado (mil t)", "RS: diferença (mil t)",
            "RB estados: diferença (mil t)", "Não alocado: diferença (mil t)"]
    print(tab[cols].round(2).to_string(index=False))
    print("\nE4:")
    print(e4[["regiao", "componente", "setor_destino", "valor_t"]].to_string(index=False))
    print("\nMétricas mais perto:")
    print(mais_perto.round(3).to_string(index=False))
    print("\nCasamentos — espaço de busca:")
    print(cas["espaco"].to_string(index=False))
    print(cas["contagens"].round(2).to_string(index=False))
    if not ach.empty:
        print(ach[["setor", "tipo de unidade", "combinação", "diferença (%)", "sentido econômico"]]
              .round(4).to_string(index=False))
    print("\nMelhores:")
    print(cas["melhores"][["critério", "setor", "combinação", "diferença (%)", "sentido econômico"]]
          .round(3).to_string(index=False))
    print("\nConclusão:")
    print("\n".join(conclusao))
    for p in (SAIDA_RECORTE, SAIDA_E4, SAIDA_XLSX):
        print(f"gravado: {p.relative_to(RAIZ)} ({p.stat().st_size} bytes)")
    print(f"tempo total: {time.time() - t0:.1f} s")


if __name__ == "__main__":
    main()
