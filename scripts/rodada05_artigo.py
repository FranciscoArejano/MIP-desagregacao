"""Rodada 05 — material do artigo (pasta artigo/): tabelas, figuras, números-chave e textos.

Uso (na raiz do repositório, depois de scripts/rodada05_simulacoes.py e scripts/rodada05_diagnostico_seeg.py):
    python scripts/rodada05_artigo.py

Tudo sai dos arquivos de resultado (resultados/rodada05/ e parametros/*_v04*); nenhum número é digitado à mão.
Saídas:
- artigo/tabelas/tabelas_artigo.xlsx (uma aba por tabela; células numéricas), tabelas_artigo.docx e tabelas_artigo.md;
- artigo/figuras/figura_*.png (300 dpi) e .pdf;
- artigo/texto/numeros_chave.md (+ numeros_chave.json, usado pelos textos) e os textos renderizados a partir dos
  modelos em artigo/texto/modelos/*.md: cada `{{chave}}` (ou `{{chave:casas}}`) é trocado pelo número-chave no padrão
  brasileiro; uma chave desconhecida é erro. Cada texto sai em .md e .docx (pandoc, com as equações nativas do Word).

Convenção: no material do artigo, milhar com ponto e decimal com vírgula (1.925,0); no .xlsx, números com formato de
exibição de uma casa decimal (três nas razões, fatores e multiplicadores).
"""
from __future__ import annotations

import json
import re
import sys
import textwrap
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
import yaml  # noqa: E402

from mip import referencia as ref  # noqa: E402
from mip.caminhos import PASTA_PARAMETROS, PASTA_RESULTADOS  # noqa: E402

RES = PASTA_RESULTADOS / "rodada05"
ART = RAIZ / "artigo"
TAB, FIG, TXT = ART / "tabelas", ART / "figuras", ART / "texto"
MODELOS = TXT / "modelos"
V = "v04"
SIM = RES / f"simulacoes_{V}.xlsx"
QUADRO = RES / f"quadro_subsetores_{V}.xlsx"
CHECAGENS = RES / f"checagens_{V}.csv"
DIAG = RES / "diagnostico_seeg_nereus.xlsx"
PARAM = PASTA_PARAMETROS / f"quadro_parametros_{V}.xlsx"

FONTE = ("Fonte: Elaboração própria a partir da matriz inter-regional de insumo-produto RS-RB (NEREUS/USP), com o setor "
         "S38 desagregado.")
FONTE_B = "Fonte: Elaboração própria a partir da matriz inter-regional de insumo-produto RS-RB (NEREUS/USP)."
PAINEL_A = "Painel A — MIP com o S38 desagregado (versão final)"
PAINEL_B = "Painel B — MIP original"

PRINCIPAIS = {"E1": "Energia elétrica agregada RS", "R1": "Cesta renovável RS", "F1": "Cesta fóssil RS"}
CESTAS = {"R1": "Renovável base", "R2": "Renovável balanceada", "R3": "Renovável infraestrutura-intensiva",
          "R4a": "Renovável energia/transporte-intensiva", "F1": "Fóssil base", "F2": "Fóssil balanceada",
          "F3": "Fóssil energia-intensiva"}
OITO = {"E1": "Energia elétrica agregada"} | CESTAS
RENOV, FOSS = ["R1", "R2", "R3", "R4a"], ["F1", "F2", "F3"]

# nomes curtos dos setores: os 13 do artigo (exatos) e a forma curta do nome da MIP para os demais
NOMES = {
    "S01": "Agricultura e apoio à agricultura", "S02": "Pecuária", "S03": "Produção florestal, pesca e aquicultura",
    "S04": "Extração de carvão mineral e de minerais não metálicos", "S05": "Extração de petróleo e gás",
    "S06": "Extração de minério de ferro", "S07": "Extração de minerais metálicos não ferrosos",
    "S08": "Abate e produtos de carne", "S09": "Fabricação e refino de açúcar", "S10": "Outros produtos alimentares",
    "S11": "Fabricação de bebidas", "S12": "Fabricação de produtos do fumo", "S13": "Fabricação de produtos têxteis",
    "S14": "Confecção de vestuário e acessórios", "S15": "Fabricação de calçados e artefatos de couro",
    "S16": "Fabricação de produtos da madeira", "S17": "Fabricação de celulose e papel",
    "S18": "Impressão e reprodução de gravações", "S19": "Refino de petróleo e coquerias",
    "S20": "Fabricação de biocombustíveis", "S21": "Químicos orgânicos e inorgânicos, resinas e elastômeros",
    "S22": "Defensivos, desinfestantes, tintas e químicos diversos", "S23": "Produtos de limpeza, cosméticos e higiene",
    "S24": "Produtos farmoquímicos e farmacêuticos", "S25": "Produtos de borracha e de material plástico",
    "S26": "Fabricação de produtos de minerais não metálicos", "S27": "Siderurgia e tubos de aço sem costura",
    "S28": "Metalurgia de metais não ferrosos e fundição",
    "S29": "Fabricação de produtos de metal, exceto máquinas e equipamentos",
    "S30": "Equipamentos de informática, eletrônicos e ópticos", "S31": "Fabricação de máquinas e equipamentos elétricos",
    "S32": "Máquinas e equipamentos mecânicos", "S33": "Automóveis, caminhões e ônibus",
    "S34": "Peças e acessórios para veículos", "S35": "Outros equipamentos de transporte",
    "S36": "Móveis e produtos de indústrias diversas", "S37": "Manutenção, reparação e instalação de máquinas",
    "S38": "Energia elétrica, gás natural e outras utilidades", "S39": "Água, esgoto e gestão de resíduos",
    "S40": "Construção", "S41": "Comércio e reparação de veículos", "S42": "Comércio por atacado e a varejo",
    "S43": "Transporte terrestre", "S44": "Transporte aquaviário", "S45": "Transporte aéreo",
    "S46": "Armazenamento, auxiliares dos transportes e correio", "S47": "Alojamento", "S48": "Alimentação",
    "S49": "Edição e edição integrada à impressão", "S50": "Televisão, rádio, cinema e gravação de som e imagem",
    "S51": "Telecomunicações", "S52": "Sistemas e outros serviços de informação",
    "S53": "Intermediação financeira e seguros", "S54": "Atividades imobiliárias",
    "S55": "Atividades jurídicas, contábeis, consultoria e sedes", "S56":
    "Serviços de arquitetura, engenharia, testes/análises técnicas e P&D",
    "S57": "Outras atividades profissionais, científicas e técnicas", "S58": "Aluguéis não imobiliários",
    "S59": "Outras atividades administrativas", "S60": "Vigilância, segurança e investigação",
    "S61": "Administração pública, defesa e seguridade social", "S62": "Educação pública", "S63": "Educação privada",
    "S64": "Saúde pública", "S65": "Saúde privada", "S66": "Atividades artísticas, criativas e de espetáculos",
    "S67": "Organizações associativas e outros serviços pessoais", "S68": "Serviços domésticos",
    "S38_HID": "Geração hidrelétrica", "S38_EOL": "Geração eólica", "S38_SOL": "Geração solar fotovoltaica",
    "S38_BIO": "Geração a biomassa", "S38_CAR": "Geração a carvão mineral", "S38_GN": "Geração a gás natural",
    "S38_PET": "Geração a derivados de petróleo", "S38_NUC": "Geração nuclear",
    "S38_TRA": "Transmissão de energia elétrica", "S38_DIS": "Distribuição e comercialização de energia elétrica",
    "S38_GAS": "Gás canalizado e outras utilidades",
}
SUBSETORES = ["S38_HID", "S38_EOL", "S38_SOL", "S38_BIO", "S38_CAR", "S38_GN", "S38_PET", "S38_NUC", "S38_TRA",
              "S38_DIS", "S38_GAS"]
TEC_B = [("S38_HID", "Hidrelétrica"), ("S38_EOL", "Eólica"), ("S38_SOL", "Solar fotovoltaica"),
         ("S38_BIO", "Biomassa"), ("S38_GN", "Gás natural"), ("S38_CAR", "Carvão mineral")]
VERSOES = {"v04": "Principal (versão final)", "v04_P2": "P2 — preços iguais entre as fontes",
           "v04_E2": "E2 — emissões na proporção do SEEG", "v04_E3": "E3 — só os níveis do SEEG",
           "v04_E4": "E4 — refino, petróleo e gás e carvão fora do S38",
           "v04_gas4dist": "gas4dist — gás do RB pelo preço de 4 distribuidoras",
           "v04_rap_media": "rap_media — transmissão pela RAP 2019",
           "v04_k060": "k060 — piso do combustível κ = 0,60", "v04_eta005": "eta005 — EOB mínimo η = 0,05"}
FUNCAO = {("E1", "S38"): "Expansão agregada do setor energético",
          ("R1", "S38"): "Componente energético agregado associado à implantação",
          ("R1", "S31"): "Equipamentos elétricos e componentes de conexão",
          ("R1", "S29"): "Estruturas metálicas e componentes físicos",
          ("R1", "S40"): "Obras civis, fundações e instalação",
          ("R1", "S43"): "Transporte de equipamentos e insumos",
          ("R1", "S56"): "Engenharia, projetos, testes e P&D",
          ("F1", "S19"): "Processamento e oferta de combustíveis fósseis",
          ("F1", "S38"): "Energia/gás e utilidades associadas",
          ("F1", "S43"): "Logística e transporte associados"}
# cores e estilo das figuras (os do artigo)
AZUL, LARANJA, GRADE = "#2c6e9e", "#da7e2e", "#e0e0e0"


# ---------------------------------------------------------------------------------------------
# formatação
# ---------------------------------------------------------------------------------------------
def br(v, casas: int = 1) -> str:
    """Número no padrão brasileiro: milhar com ponto, decimal com vírgula; negativo com '−'."""
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    if isinstance(v, str):
        return v
    s = f"{abs(float(v)):,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if float(v) < 0 and float(s.replace(".", "").replace(",", ".")) != 0:
        s = "−" + s
    return s


def pct(v, casas: int = 1) -> str:
    return br(v, casas) + "%"


class Tabela:
    """Uma tabela do artigo: um ou mais painéis (título do painel, DataFrame), casas decimais por coluna e notas.

    (Classe simples, e não dataclass, para o módulo poder ser carregado por importlib nos testes e nas conferências.)"""

    def __init__(self, numero: str, titulo: str, paineis: list, casas: dict | None = None, notas: list | None = None,
                 fonte: str = FONTE, aba: str = ""):
        self.numero, self.titulo, self.paineis = numero, titulo, paineis  # numero: "Tabela 1", "Quadro 1", "Tabela A2"
        self.casas = dict(casas or {})  # coluna → casas decimais (padrão 1)
        self.notas = list(notas or [])
        self.fonte, self.aba = fonte, aba

    def colunas(self) -> list:
        return list(self.paineis[0][1].columns)

    def celula(self, col, v, ponto: bool = False) -> str:
        """Célula formatada: padrão brasileiro (artigo) ou, com ponto=True, ponto decimal sem milhar (relatório)."""
        f = (lambda x, c: "" if (isinstance(x, float) and np.isnan(x)) else f"{float(x):.{c}f}") if ponto else br
        if isinstance(v, (int, np.integer)) and not isinstance(v, bool):
            return f(v, 0)
        if isinstance(v, (float, np.floating)):
            return f(v, self.casas.get(col, 1))
        return "" if v is None else str(v)


# ---------------------------------------------------------------------------------------------
# dados
# ---------------------------------------------------------------------------------------------
def carregar() -> dict:
    x = pd.ExcelFile(SIM)
    d = {aba: pd.read_excel(x, aba) for aba in x.sheet_names if aba != "LEIAME"}
    q = pd.ExcelFile(QUADRO)
    for aba in ("quadro", "d18_latentes", "variantes_resumo", "realocacao_E4", "niveis_emissoes", "reequilibrio_RAS",
                "pisos_vbp", "insumos_nao_energeticos", "passos_v03_v04"):
        d[f"q_{aba}"] = pd.read_excel(q, aba)
    d["checagens"] = pd.read_csv(CHECAGENS)
    d["parametros"] = pd.read_excel(PARAM, "quadro")
    d["viabilidade"] = pd.read_excel(PARAM, "viabilidade")
    with open(PASTA_PARAMETROS / f"especificacao_{V}.yaml", encoding="utf-8") as f:
        d["espec"] = yaml.safe_load(f)
    if DIAG.exists():
        dx = pd.ExcelFile(DIAG)
        d["diag"] = {aba: pd.read_excel(dx, aba) for aba in dx.sheet_names}
    return d


def imp(d, aba, cenario) -> pd.Series:
    t = d[aba]
    return t[t["cenario"] == cenario].iloc[0]


def ind(d, versao, cenario) -> pd.Series:
    """Indicadores de um cenário numa versão (A0 na original; A2 e B nas versões da v04)."""
    if versao == "original":
        return imp(d, "A0", cenario)
    if versao == "original_E4":
        return imp(d, "A0_E4", cenario)
    t = d["robustez_detalhe"]
    return t[(t["versao"] == versao) & (t["cenario"] == cenario)].iloc[0]


# ---------------------------------------------------------------------------------------------
# tabelas
# ---------------------------------------------------------------------------------------------
def quadro1(d) -> Tabela:
    rep = d["reparticao_S38"]
    e1 = rep[(rep["cenario"] == "A2_E1") & (rep["versao"] == V)].set_index("subsetor")["peso"]
    cen = d["cenarios"]
    f1 = cen[(cen["cenario"] == "A2_F1") & cen["setor"].isin(["S38_CAR", "S38_GN"])].set_index("setor")["valor_R$mi"]
    f1 = f1 / f1.sum()
    linhas = []
    for c in ("E1", "R1", "F1"):
        for s, v in ref.CENARIOS_RS[c].items():
            if s != "S38":
                desag = "igual"
            elif c == "E1":
                desag = (f"{NOMES['S38_DIS']} ({pct(100 * e1['S38_DIS'])}) e {NOMES['S38_GAS']} "
                         f"({pct(100 * e1['S38_GAS'])}), pela composição da demanda final")
            elif c in RENOV:
                desag = f"{NOMES['S38_EOL']} (S38_EOL)"
            else:
                desag = (f"{NOMES['S38_CAR']} ({pct(100 * f1['S38_CAR'])}) e {NOMES['S38_GN']} "
                         f"({pct(100 * f1['S38_GN'])}), pelo VBP de 2019")
            linhas.append({"Cenário": PRINCIPAIS[c].replace(" RS", ""), "Região": "RS", "Código": s, "Setor": NOMES[s],
                           "Participação no choque (%)": 100 * v / 1000, "Valor do choque (R$ milhões)": v,
                           "Função no cenário": FUNCAO[(c, s)], "Setor na MIP desagregada": desag})
    return Tabela("Quadro 1", "Composição dos vetores de choque dos cenários principais e atribuição do S38 na MIP "
                  "desagregada", [(None, pd.DataFrame(linhas))], notas=[
                      "Nota: choques de R$ 1.000 milhões na demanda final de setores do RS. Na MIP desagregada, o S38 "
                      "da energia elétrica agregada é repartido entre distribuição e gás canalizado na proporção das "
                      "vendas de cada subsetor à demanda final do RS; o das cestas renováveis vai para a geração eólica; "
                      "o das cestas fósseis, para a geração a carvão mineral e a gás natural do RS na proporção do VBP "
                      "de 2019 de cada uma."], aba="Quadro1")


def _t1_painel(d, versao, pref) -> pd.DataFrame:
    linhas = []
    for c, nome in PRINCIPAIS.items():
        x = ind(d, versao, f"{pref}_{c}")
        linhas.append({"Cenário": nome, "Produção bruta induzida (R$ milhões)": x["ΔX total (R$ mi)"],
                       "Emissões totais (t CO₂e)": x["CO₂e total (t)"],
                       "Intensidade carbônica (t CO₂e/R$ milhão)": x["IC (t CO₂e por R$ mi)"],
                       "Produção bruta por mil t CO₂e (R$ milhões)": x["Produção por mil t CO₂e (R$ mi)"]})
    t = pd.DataFrame(linhas)
    f, r = t.iloc[2], t.iloc[1]
    t.loc[len(t)] = {"Cenário": "Diferença de emissões incorporadas (fóssil − renovável)",
                     "Emissões totais (t CO₂e)": f["Emissões totais (t CO₂e)"] - r["Emissões totais (t CO₂e)"]}
    ic_f, ic_r = f["Intensidade carbônica (t CO₂e/R$ milhão)"], r["Intensidade carbônica (t CO₂e/R$ milhão)"]
    t.loc[len(t)] = {"Cenário": "Redução da intensidade carbônica (%)",
                     "Intensidade carbônica (t CO₂e/R$ milhão)": 100 * (ic_f - ic_r) / ic_f}
    return t


def tabela1(d) -> Tabela:
    return Tabela("Tabela 1", "Impactos econômicos e ambientais dos cenários energéticos",
                  [(PAINEL_A, _t1_painel(d, V, "A2")), (PAINEL_B, _t1_painel(d, "original", "A0"))],
                  notas=["Nota: choques de R$ 1.000 milhões na demanda final do RS. A redução da intensidade carbônica é "
                         "a da cesta renovável em relação à cesta fóssil, em % da intensidade da cesta fóssil."],
                  aba="Tabela1")


def _top10(d, aba, cenario, memoria=True) -> pd.DataFrame:
    top = d[aba]
    top = top[top["cenario"] == cenario].sort_values("#")
    x = imp(d, "A0" if cenario.startswith("A0") else "A2_sens", cenario)
    linhas = [{"Ranking": int(r["#"]), "Região": r["regiao"], "Código": r["codigo"], "Setor": NOMES[r["codigo"]],
               "Produção induzida (R$ milhões)": r["dX (R$ mi)"], "Emissões induzidas (t CO₂e)": r["CO2e induzido (t)"],
               "Participação (%)": r["% do CO2e total"]} for _, r in top.iterrows()]
    tot_x, tot_e = x["ΔX total (R$ mi)"], x["CO₂e total (t)"]
    sx, se = top["dX (R$ mi)"].sum(), top["CO2e induzido (t)"].sum()
    linhas.append({"Ranking": "", "Região": "", "Código": "", "Setor": "Demais setores",
                   "Produção induzida (R$ milhões)": tot_x - sx, "Emissões induzidas (t CO₂e)": tot_e - se,
                   "Participação (%)": 100 * (tot_e - se) / tot_e})
    linhas.append({"Ranking": "Total", "Região": "", "Código": "", "Setor": "Cesta renovável RS",
                   "Produção induzida (R$ milhões)": tot_x, "Emissões induzidas (t CO₂e)": tot_e,
                   "Participação (%)": 100.0})
    if memoria:
        p = d["participacao_S38"]
        m = p[(p["cenario"] == cenario) & (p["codigo"] == "S38 (total)")].iloc[0]
        rot = ("Memória: subsetores do S38 (RS e RB)" if not cenario.startswith("A0") else "Memória: S38 (RS e RB)")
        linhas.append({"Ranking": "", "Região": "", "Código": "", "Setor": rot,
                       "Produção induzida (R$ milhões)": m["dX"], "Emissões induzidas (t CO₂e)": m["co2e"],
                       "Participação (%)": m["pct_co2e_total"]})
    return pd.DataFrame(linhas)


def tabela2(d) -> Tabela:
    return Tabela("Tabela 2", "Principais setores responsáveis pelas emissões induzidas pela cesta renovável",
                  [(PAINEL_A, _top10(d, "top10_A2", "A2_R1"))], notas=[
                      "Nota: choque de R$ 1.000 milhões na demanda final do RS (cesta renovável base); o S38 da cesta "
                      "foi atribuído à geração eólica do RS. A linha de memória soma os 22 subsetores do S38 (11 em cada "
                      "região) e não entra no total."], aba="Tabela2")


def tabelaA2(d) -> Tabela:
    return Tabela("Tabela A2", "Principais setores responsáveis pelas emissões induzidas pela cesta renovável na MIP "
                  "original", [(PAINEL_B, _top10(d, "top10_A0", "A0_R1"))], fonte=FONTE_B,
                  notas=["Nota: choque de R$ 1.000 milhões na demanda final do RS (cesta renovável base). A linha de "
                         "memória soma o S38 das duas regiões e não entra no total."], aba="TabelaA2")


def _t3_painel(d, versao, pref) -> pd.DataFrame:
    reg = d["regioes"]
    linhas = []
    for c, nome in PRINCIPAIS.items():
        for r in ("RS", "RB"):
            x = reg[(reg["versao"] == versao) & (reg["cenario"] == f"{pref}_{c}") & (reg["regiao"] == r)].iloc[0]
            linhas.append({"Cenário": nome, "Região impactada": r, "Produção induzida (R$ milhões)": x["dX"],
                           "Emissões induzidas (t CO₂e)": x["co2e"], "Participação nas emissões (%)": x["pct_co2e"]})
    return pd.DataFrame(linhas)


def tabela3(d) -> Tabela:
    return Tabela("Tabela 3", "Spillovers ambientais entre o Rio Grande do Sul e o Restante do Brasil",
                  [(PAINEL_A, _t3_painel(d, V, "A2")), (PAINEL_B, _t3_painel(d, "original", "A0"))],
                  notas=["Nota: choques de R$ 1.000 milhões na demanda final do RS; RB = Restante do Brasil."],
                  aba="Tabela3")


def _t4_linha(x, nome) -> dict:
    return {"Cenário": nome, "Produção bruta induzida (R$ milhões)": x["ΔX total (R$ mi)"],
            "Emissões totais (t CO₂e)": x["CO₂e total (t)"],
            "Intensidade carbônica (t CO₂e/R$ milhão)": x["IC (t CO₂e por R$ mi)"],
            "Emissões indiretas (%)": x["% indiretas"], "Emissões no RB (%)": x["% CO₂e no RB"]}


def tabela4(d) -> Tabela:
    pa = pd.DataFrame([_t4_linha(ind(d, V, f"A2_{c}"), n) for c, n in CESTAS.items()])
    pb = pd.DataFrame([_t4_linha(ind(d, "original", f"A0_{c}"), n + (" (a)" if c == "R4a" else ""))
                       for c, n in CESTAS.items()])
    return Tabela("Tabela 4", "Sensibilidade dos resultados à composição dos vetores de choque",
                  [(PAINEL_A, pa), (PAINEL_B, pb)],
                  notas=["Nota: choques de R$ 1.000 milhões na demanda final do RS; vetores na Tabela A1.",
                         "(a) Vetor da Tabela A1 (45/15/10/10/15/5)."], aba="Tabela4")


def tabela5(d) -> Tabela:
    linhas = []
    opc_r = [("", "Geração eólica 100% (principal)"), ("_EOL50SOL50", "Geração eólica 50% + solar 50%"),
             ("_EOL67TRA33", "Geração eólica 2/3 + transmissão 1/3")]
    opc_f = [("", "Carvão mineral e gás natural ∝ VBP (principal)"), ("_CAR100", "Geração a carvão mineral 100%"),
             ("_GN100", "Geração a gás natural 100%")]
    for c, nome in CESTAS.items():
        for suf, rot in (opc_r if c in RENOV else opc_f):
            x = imp(d, "A2_sens", f"A2_{c}{suf}")
            linhas.append({"Cenário": nome, "Atribuição do S38": rot} | {k: v for k, v in _t4_linha(x, nome).items()
                                                                          if k != "Cenário"})
    return Tabela("Tabela 5", "Sensibilidade à tecnologia atribuída ao S38",
                  [(PAINEL_A, pd.DataFrame(linhas))],
                  notas=["Nota: choques de R$ 1.000 milhões na demanda final do RS. A atribuição muda só a parcela do "
                         "S38 de cada cesta; os demais setores do vetor não mudam. Solar: subsetor sem produção no RS em "
                         "2019 (latente), com os coeficientes técnicos do RB e a TUST do RS."], aba="Tabela5")


def tabela6(d) -> Tabela:
    tw = d["B_TWh"]
    tw = tw[tw["versao"] == V].set_index("subsetor (RS)")
    b1 = d["B_RS1bi"].set_index("cenario")
    linhas = []
    for s, nome in TEC_B:
        x = tw.loc[s]
        linhas.append({"Tecnologia": nome, "Preço médio (R$/MWh)": x["R$/MWh = choque de 1 TWh (R$ mi)"],
                       "Produção induzida por TWh (R$ milhões)": x["ΔX total (R$ mi)"],
                       "Emissões diretas por TWh (t CO₂e)": x["diretas"],
                       "Emissões indiretas por TWh (t CO₂e)": x["indiretas"],
                       "Emissões totais por TWh (t CO₂e)": x["t CO₂e por TWh (total)"], "g CO₂e/kWh": x["g CO₂e/kWh"],
                       "Emissões no RB (%)": x["% no RB"],
                       "Emissões por R$ 1 bilhão (t CO₂e)": b1.loc[f"B_{s.replace('S38_', '')}", "CO₂e total (t)"],
                       "Emissões evitadas por TWh em relação ao carvão (t CO₂e)":
                           x["evitadas por TWh se substituir o carvão (t)"]})
    linhas.append({"Tecnologia": "Distribuição (eletricidade da rede, mix de 2019)",
                   "Emissões por R$ 1 bilhão (t CO₂e)": b1.loc["B_DIS", "CO₂e total (t)"]})
    return Tabela("Tabela 6", "Emissões por tecnologia de geração do RS", [(PAINEL_A, pd.DataFrame(linhas))],
                  notas=["Nota: Choque de 1 TWh (ao preço médio de 2019) e de R$ 1 bilhão na demanda final de cada "
                         "subsetor de geração do RS. Emissões da operação (produção de energia elétrica); não incluem "
                         "bens de capital (implantação das usinas). Solar: subsetor sem produção no RS em 2019 "
                         "(latente), com os coeficientes técnicos do RB e a TUST do RS."], aba="Tabela6")


def tabela7(d) -> Tabela:
    p1, p2 = [], []
    for c, nome in OITO.items():
        a0, a1, a2 = ind(d, "original", f"A0_{c}"), imp(d, "A1", f"A1_{c}"), ind(d, V, f"A2_{c}")
        e0, e1, e2 = a0["CO₂e total (t)"], a1["CO₂e total (t)"], a2["CO₂e total (t)"]
        p1.append({"Cenário": nome, "Emissões A0 (t CO₂e)": e0, "Emissões A1 (t CO₂e)": e1,
                   "Emissões A2 (t CO₂e)": e2, "Variação A1/A0 (%)": 100 * (e1 - e0) / e0,
                   "Variação A2/A0 (%)": 100 * (e2 - e0) / e0})
        p2.append({"Cenário": nome, "Intensidade A0 (t CO₂e/R$ milhão)": a0["IC (t CO₂e por R$ mi)"],
                   "Intensidade A2 (t CO₂e/R$ milhão)": a2["IC (t CO₂e por R$ mi)"],
                   "Emissões indiretas A0 (%)": a0["% indiretas"], "Emissões indiretas A2 (%)": a2["% indiretas"],
                   "Emissões no RB A0 (%)": a0["% CO₂e no RB"], "Emissões no RB A2 (%)": a2["% CO₂e no RB"]})
    return Tabela("Tabela 7", "Efeito da desagregação: MIP original, S38 desagregado sem e com atribuição de tecnologia",
                  [("Parte 1 — Emissões totais", pd.DataFrame(p1)),
                   ("Parte 2 — Intensidade carbônica, emissões indiretas e emissões no RB", pd.DataFrame(p2))],
                  notas=["Nota: A0 = MIP original; A1 = MIP desagregada, com o S38 do vetor repartido entre distribuição e "
                         "gás canalizado pela composição da demanda final; A2 = MIP desagregada, com o S38 atribuído a "
                         "tecnologias (Quadro 1). Na energia elétrica agregada, A1 e A2 coincidem."], aba="Tabela7")


def tabelaA1(d) -> Tabela:
    setores = ["S38", "S31", "S29", "S40", "S43", "S56", "S19"]
    linhas = []
    for s in setores:
        lin = {"Código": s, "Setor": NOMES[s]}
        for c, n in CESTAS.items():
            lin[n] = 100 * ref.CENARIOS_RS[c].get(s, 0.0) / 1000
        linhas.append(lin)
    tot = {"Código": "", "Setor": "Total"} | {n: sum(ln[n] for ln in linhas) for n in CESTAS.values()}
    linhas.append(tot)
    cen = d["cenarios"]
    f1 = cen[(cen["cenario"] == "A2_F1") & cen["setor"].isin(["S38_CAR", "S38_GN"])].set_index("setor")["valor_R$mi"]
    f1 = 100 * f1 / f1.sum()
    return Tabela("Tabela A1", "Vetores de choque das cestas (participação de cada setor no choque, %)",
                  [(None, pd.DataFrame(linhas))],
                  notas=["Nota: choques de R$ 1.000 milhões na demanda final do RS. Renovável energia/transporte-intensiva: "
                         "vetor 45/15/10/10/15/5. No A2 (MIP desagregada), a parcela do S38 das cestas renováveis vai para "
                         f"a geração eólica do RS, e a das cestas fósseis para a geração a carvão mineral ({pct(f1['S38_CAR'])}) "
                         f"e a gás natural ({pct(f1['S38_GN'])}) do RS, na proporção do VBP de 2019."], aba="TabelaA1")


def tabelaA3(d) -> Tabela:
    q = d["q_quadro"]
    linhas = []
    for r in ("RS", "RB"):
        for s in SUBSETORES:
            x = q[(q["versao"] == V) & (q["regiao"] == r) & (q["subsetor"] == s)].iloc[0]
            lat = "latente" in str(x["alertas"])
            nome = NOMES[s] + (" (latente)" if lat else (" (sem produção)" if x["VBP (R$ mi)"] == 0 else ""))
            linhas.append({"Região": r, "Subsetor": nome, "VBP (R$ milhões)": x["VBP (R$ mi)"], "GWh": x["GWh"],
                           "R$/MWh": x["R$/MWh implícito"], "t CO₂e/MWh (E1)": x["t CO2e/MWh E1"],
                           "EOB/VBP": x["EOB/VBP"], "Multiplicador de produção": x["multiplicador de produção"],
                           "Multiplicador de emissões (t CO₂e/R$ milhão)": x["multiplicador de emissões (t/R$ mi)"]})
        o = q[(q["versao"] == "original") & (q["regiao"] == r)].iloc[0]
        linhas.append({"Região": r, "Subsetor": "S38 agregado (MIP original)", "VBP (R$ milhões)": o["VBP (R$ mi)"],
                       "EOB/VBP": o["EOB/VBP"], "Multiplicador de produção": o["multiplicador de produção"],
                       "Multiplicador de emissões (t CO₂e/R$ milhão)": o["multiplicador de emissões (t/R$ mi)"]})
    t = pd.DataFrame(linhas)
    t[["GWh", "R$/MWh", "t CO₂e/MWh (E1)"]] = t[["GWh", "R$/MWh", "t CO₂e/MWh (E1)"]].where(t["GWh"] > 0)
    return Tabela("Tabela A3", "Os 11 subsetores do S38 no RS e no RB: produção, preços, emissões e multiplicadores",
                  [(None, t)], casas={"t CO₂e/MWh (E1)": 3, "EOB/VBP": 3, "Multiplicador de produção": 3},
                  notas=["Nota: GWh de 2019 (BEN 2020 e ONS); R$/MWh = VBP ÷ GWh; t CO₂e/MWh pela variante E1. "
                         "Multiplicador de produção = soma da coluna da inversa de Leontief; multiplicador de emissões = "
                         "emissões totais (diretas e indiretas) por R$ milhão de demanda final do subsetor. Subsetores "
                         "latentes: coeficientes técnicos do mesmo subsetor no RB, com a TUST por MWh do RS."],
                  aba="TabelaA3")


def tabelaA4(d) -> Tabela:
    linhas = [
        ("Matriz inter-regional de insumo-produto RS-RB e conta-satélite de emissões (CO₂e)", "NEREUS/USP", "2019",
         "Base da desagregação; todas as restrições de agregação"),
        ("Geração de energia elétrica por fonte (serviço público), Brasil", "EPE, Balanço Energético Nacional 2020, "
         "Tabela 5.4", "2019 (BEN 2020)", "GWh por fonte do Brasil; RB = Brasil − RS"),
        ("Geração por usina", "ONS, dados abertos (geração por usina)", "2019",
         "GWh por fonte do RS; MWh no SIN (pesos da TUST da geração)"),
        ("Encargos de uso do sistema de transmissão", "ONS, Síntese da Apuração Mensal dos Serviços e Encargos (SEAMSE)",
         "dez/2018, jun/2019 e dez/2019", "VBP da transmissão (ano civil de 2019)"),
        ("Receita anual permitida da transmissão", "ANEEL, REH nº 2.408/2018 e decisão de 25/06/2019 (REH nº 2.565/2019; "
         "total por fonte secundária)", "ciclos 2018-2019 e 2019-2020", "Variante _rap_media"),
        ("Rateio dos encargos de uso entre geração e consumo", "ANEEL, Resolução Normativa nº 559/2013", "2013",
         "Parcela da transmissão paga pela geração (50%)"),
        ("Preços de energia por fonte", "ANEEL, resultado dos leilões de geração", "leilões até 2019, corrigidos pelo "
         "IPCA (IBGE, SIDRA, Tabela 1737)", "Participação de cada fonte no VBP da geração (preços P1)"),
        ("Remunerações e pessoal ocupado por atividade", "IBGE, Cadastro Central de Empresas (CEMPRE), SIDRA, Tabela "
         "6450", "2019", "Repartição de remunerações e ocupações entre geração, transmissão, distribuição e gás"),
        ("Demonstrações contábeis da Sulgás", "Sulgás (publicadas pela CAGE/Sefaz-RS)", "2019",
         "VBP do gás canalizado no RS; preço do gás usado no RB; vendas por segmento"),
        ("Demanda de gás natural por distribuidora", "MME, Boletim Mensal de Acompanhamento da Indústria de Gás Natural "
         "nº 154 (fonte: Abegás)", "dez/2019 (média de 2019)", "Volume não térmico do RB"),
        ("Demonstrações financeiras de distribuidoras de gás", "CVM, DFP de Comgás, CEG e Gasmig", "2019",
         "Variante _gas4dist"),
        ("Importações de energia elétrica e gás natural", "MDIC, Comex Stat", "2019", "Importações do S38 (energia)"),
        ("Emissões de gases de efeito estufa", "SEEG 13.0 (Observatório do Clima), GWP-AR5", "2019",
         "Níveis por fonte de geração (E1), SF₆, variantes E2 a E4, diagnóstico (Tabela A7)"),
    ]
    return Tabela("Tabela A4", "Fontes de dados da desagregação",
                  [(None, pd.DataFrame(linhas, columns=["Dado", "Fonte", "Ano ou versão", "Uso na desagregação"]))],
                  fonte="Fonte: Elaboração própria.", notas=["Nota: URLs, datas de acesso, páginas e SHA-256 dos arquivos "
                                                             "no registro de fontes do repositório (dados/externos/FONTES.md)."],
                  aba="TabelaA4")


def tabelaA5(d) -> Tabela:
    q = d["q_quadro"]
    e = d["espec"]
    d18 = d["q_d18_latentes"].set_index("subsetor")
    vr = d["q_variantes_resumo"].set_index("versao")
    tra = vr.loc[V, "VBP TRA-RS"] + vr.loc[V, "VBP TRA-RB"]
    tra_m = vr.loc["v04_rap_media", "VBP TRA-RS"] + vr.loc["v04_rap_media", "VBP TRA-RB"]
    gas_rb = vr.loc[V, "VBP GAS-RB"]
    gas_4d = vr.loc["v04_gas4dist", "VBP GAS-RB"]
    gas_rs = q[(q["versao"] == V) & (q["regiao"] == "RS") & (q["subsetor"] == "S38_GAS")]["VBP (R$ mi)"].iloc[0]
    k = e["piso_vbp"]["kappa"]
    eta = e["reequilibrio"]["eta"]
    comp = pd.read_csv(PASTA_PARAMETROS / f"componentes_{V}.csv")
    alfa = float(comp.loc[comp["componente"] == "fracao_transmissao_distribuicao", "valor"].iloc[0])
    lin = pd.read_csv(PASTA_PARAMETROS / f"regras_linha_{V}.csv")
    teto = float(lin.loc[lin["subsetor"] == "S38_GAS", "valor"].astype(float).max())
    linhas = [
        ("VBP da transmissão (Brasil)", f"R$ {br(tra)} milhões", "Encargos de uso faturados pelo ONS no ano civil de 2019 "
         "(SEAMSE), sem PIS/Pasep e Cofins; repartido entre RS e RB pela participação no VBP do S38",
         f"RAP 2019: R$ {br(tra_m)} milhões (_rap_media)"),
        ("Parcela da transmissão paga pelo consumo, via distribuição (α); a geração paga 1 − α", br(alfa, 1),
         "Rateio inicial dos encargos de uso: 50% geração e 50% consumo (ANEEL, REN nº 559/2013)", "—"),
        ("TUST da geração entre as fontes", "∝ MWh gerado no SIN", "Regra do encargo por uso da rede; sistemas isolados "
         "não pagam", "—"),
        ("TUST das colunas latentes (solar e derivados de petróleo no RS)",
         f"{br(d18.loc['S38_SOL', 'tau (R$/MWh)'], 2)} R$/MWh", "HIPÓTESE: a mesma TUST por MWh das fontes do RS no SIN, "
         "ao preço do espelho do RB", "TUST por real do espelho (registro no relatório)"),
        ("VBP do gás canalizado no RS", f"R$ {br(gas_rs)} milhões", "Receita líquida da Sulgás em 2019", "—"),
        ("VBP do gás canalizado no RB", f"R$ {br(gas_rb)} milhões", "HIPÓTESE: preço da Sulgás × volume não térmico do "
         "RB (MME/Abegás)", f"Preço médio de 4 distribuidoras: R$ {br(gas_4d)} milhões (_gas4dist)"),
        ("Participação máxima do gás canalizado na compra de S38 de cada comprador", br(teto, 2), "HIPÓTESE (teto)", "—"),
        ("Piso do VBP das térmicas (κ)", br(k, 2), "HIPÓTESE: combustível ≤ κ × VBP", "κ = 0,60 (_k060)"),
        ("Excedente operacional bruto mínimo (η)", br(eta, 2), "HIPÓTESE: EOB ≥ η × VBP em cada subsetor (ajuste RAS "
         "dos insumos não energéticos)", "η = 0,05 (_eta005)"),
        ("Remunerações e ocupações entre as fontes de geração", "∝ VBP − combustível", "HIPÓTESE: o combustível é "
         "repasse e não exige mão de obra proporcional", "—"),
        ("Preços da energia por fonte (P1)", "leilões", "HIPÓTESE: preços relativos dos leilões ajustados à produção "
         "da MIP", "Preços iguais (P2)"),
        ("Emissões das usinas (E1)", "níveis do SEEG + resíduo", "Níveis por fonte do SEEG; o resíduo (base − níveis) "
         "vai para distribuição e gás canalizado ∝ VBP", "E2 (proporcional ao SEEG), E3 (só níveis), E4 (realocação)"),
        ("Subsetores sem produção no RS (solar e derivados de petróleo)", "espelho do RB", "HIPÓTESE: coeficientes "
         "técnicos do mesmo subsetor no RB", "—"),
    ]
    return Tabela("Tabela A5", "Parâmetros e hipóteses da desagregação",
                  [(None, pd.DataFrame(linhas, columns=["Parâmetro", "Valor", "Fonte ou justificativa",
                                                        "Sensibilidade testada"]))],
                  fonte="Fonte: Elaboração própria.", aba="TabelaA5")


def tabelaA6(d) -> Tabela:
    cen = [("A2_E1", OITO["E1"]), ("A2_R1", CESTAS["R1"]), ("A2_F1", CESTAS["F1"])] + \
          [(f"B_TWh_{s.replace('S38_', '')}", f"1 TWh — {n}") for s, n in TEC_B]
    linhas = []
    for c, nome in cen:
        b = ind(d, V, c)
        for v, rot in VERSOES.items():
            x = ind(d, v, c)
            linhas.append({"Cenário": nome, "Versão": rot, "Emissões (t CO₂e)": x["CO₂e total (t)"],
                           "Desvio das emissões (%)": 100 * (x["CO₂e total (t)"] - b["CO₂e total (t)"]) / b["CO₂e total (t)"],
                           "Intensidade (t CO₂e/R$ milhão)": x["IC (t CO₂e por R$ mi)"],
                           "Desvio da intensidade (%)": 100 * (x["IC (t CO₂e por R$ mi)"] - b["IC (t CO₂e por R$ mi)"])
                           / b["IC (t CO₂e por R$ mi)"]})
    p1 = pd.DataFrame(linhas)
    linhas = []
    for c in ("E1", "R1", "F1"):
        b, x = ind(d, "original", f"A0_{c}"), ind(d, "original_E4", f"A0_{c}")
        for rot, y in (("MIP original", b), ("MIP original com a E4", x)):
            linhas.append({"Cenário": OITO[c] if c == "E1" else CESTAS[c], "Versão": rot,
                           "Emissões (t CO₂e)": y["CO₂e total (t)"],
                           "Desvio das emissões (%)": 100 * (y["CO₂e total (t)"] - b["CO₂e total (t)"]) / b["CO₂e total (t)"],
                           "Intensidade (t CO₂e/R$ milhão)": y["IC (t CO₂e por R$ mi)"],
                           "Desvio da intensidade (%)": 100 * (y["IC (t CO₂e por R$ mi)"] - b["IC (t CO₂e por R$ mi)"])
                           / b["IC (t CO₂e por R$ mi)"]})
    return Tabela("Tabela A6", "Robustez: resultados da versão principal e das variantes da desagregação",
                  [("Painel A — MIP com o S38 desagregado: versão principal e 8 variantes", p1),
                   ("Painel B — MIP original sem e com a realocação de emissões da E4 (A0_E4)", pd.DataFrame(linhas))],
                  notas=["Nota: desvio em relação à versão principal (Painel A) ou à MIP original (Painel B). Choques de "
                         "R$ 1.000 milhões (cenários) e de 1 TWh ao preço médio de cada versão (tecnologias). Na E4, "
                         "emissões do refino, da exploração de petróleo e gás e da mineração de carvão do SEEG saem do "
                         "S38 e vão para os setores S19, S05 e S04 da mesma região."], aba="TabelaA6")


def _com_cabecalho(df: pd.DataFrame, marca: str) -> pd.DataFrame:
    """Aba com linhas de título antes do cabeçalho: usa como cabeçalho a primeira linha que contém `marca`."""
    if marca in map(str, df.columns):
        return df
    for i in range(min(len(df), 10)):
        if marca in map(str, df.iloc[i].tolist()):
            out = df.iloc[i + 1:].copy()
            out.columns = [str(c) for c in df.iloc[i].tolist()]
            return out.reset_index(drop=True)
    raise SystemExit(f"cabeçalho com '{marca}' não encontrado")


def tabelaA7(d) -> Tabela | None:
    if "diag" not in d:
        return None
    c = _com_cabecalho(d["diag"]["componentes_S38"], "RS (mil t)")
    comp = next(x for x in c.columns if x.startswith("componente"))
    regs = [("RS (mil t)", "RS"), ("RB estados (mil t)", "RB (estados)"), ("Não alocado (mil t)", "Não alocado")]
    t = c[[comp] + [a for a, _ in regs]].copy()
    fim = t.index[t[comp].astype(str).str.startswith("S38 na base NEREUS")]
    if not len(fim):
        raise SystemExit("diagnóstico: linha 'S38 na base NEREUS' não encontrada na aba componentes_S38")
    t = t.loc[:fim[0]].dropna(subset=[comp])  # a aba continua com a conferência da E4, que não entra na tabela
    t = t.rename(columns={comp: "Componente (SEEG: atividade › categoria › subcategoria)"}
                 | {a: f"{b} (mil t CO₂e)" for a, b in regs})
    for a, b in regs:
        t[f"{b} (mil t CO₂e)"] = pd.to_numeric(t[f"{b} (mil t CO₂e)"], errors="coerce")
    return Tabela("Tabela A7", "Emissões do S38 na base NEREUS e componentes do SEEG 13.0 (2019, GWP-AR5)",
                  [(None, t.reset_index(drop=True))], fonte="Fonte: Elaboração própria a partir da base NEREUS/USP e do "
                                                            "SEEG 13.0 (Observatório do Clima).",
                  notas=["Nota: emissões de 'Emissão' e 'Emissão NCI' (não há NCI nessas duas atividades). RB = estados "
                         "exceto o RS; 'Não alocado' = sem estado. No RS, a soma das duas atividades é compatível com o "
                         "S38 da base NEREUS, o que é um indício fraco de que a conta-satélite inclua no S38 o refino e a "
                         "produção de carvão mineral: o RB não fecha pelo mesmo critério, coincidências são esperadas com "
                         "o número de combinações testadas (Apêndice B) e há uma combinação concorrente que atribui o "
                         "refino ao S19. Não há documentação da conta-satélite que permita confirmar."], aba="TabelaA7")


def tabelaA8(d) -> Tabela:
    c = d["checagens"]
    p = c[c["versao"] == V].set_index("checagem")
    reag = max(float(p.loc[f"reagregacao_{n}", "valor"]) for n in ("Z", "F", "X", "total_linha", "V", "V_df", "X_df"))
    sat = max(float(p.loc[f"satelite_{n}_preservado", "valor"]) for n in ("ocupacoes", "co2e", "agua", "energia"))
    vr = d["q_variantes_resumo"]
    a0 = d["A0"].set_index("cenario")
    dif_a0 = max(abs(a0.loc[f"A0_{e.cenario}", "CO₂e total (t)"] - e.co2e_total) for e in ref.ESPERADO.itertuples())
    eob_min = float(vr.loc[vr["versao"] == V, "EOB/VBP mínimo"].iloc[0])
    nao_neg = p.loc["nao_negatividade", "status"] == "OK"
    e = d["espec"]
    alerta = c[c["status"] == "ALERTA"]
    n_mult = int((alerta["checagem"] == "multiplicador_producao_subsetores").sum())
    sem_ne = sorted(set(c["versao"]) - set(alerta.loc[alerta["checagem"] == "insumos_nao_energeticos_minimo", "versao"]))
    n_ne = c["versao"].nunique() - len(sem_ne)
    if sem_ne != [f"{V}_eta005", f"{V}_k060"]:
        raise SystemExit(f"Tabela A8: versões sem o alerta do T12 mudaram: {sem_ne}")
    kappa_sens, eta_sens = (yaml.safe_load(open(PASTA_PARAMETROS / f"especificacao_{V}{suf}.yaml", encoding="utf-8"))[a][b]
                            for suf, a, b in (("_k060", "piso_vbp", "kappa"), ("_eta005", "reequilibrio", "eta")))
    q = d["q_quadro"]
    qm = q[(q["versao"] == V) & q["alertas"].fillna("").str.contains("multiplicador")]
    mult = [f"{NOMES[x.subsetor].lower()} do {x.regiao}" + (", latente" if x["VBP (R$ mi)"] == 0 else "")
            for _, x in qm.iterrows()]
    linhas = [
        ("Re-agregação", "A soma dos 11 subsetores de cada região reproduz a matriz original (fluxos intermediários, "
         "demanda final, produção, valor adicionado e importações)", reag, "R$ milhões"),
        ("Contas-satélite", "A soma dos subsetores reproduz as emissões, as ocupações, a água e a energia do S38 original",
         sat, "unidades da conta"),
        ("Balanço", "Em cada subsetor, vendas = VBP e compras + valor adicionado = VBP (salvo o resíduo da matriz "
         "original, repartido pelo VBP)", float(p.loc["balanco_coluna_subsetores", "valor"]), "R$ milhões"),
        ("Não negatividade", "Nenhum fluxo, valor adicionado ou emissão negativo novo" if nao_neg else
         "HÁ NEGATIVOS NOVOS (ver checagens)", float(p.loc["nao_negatividade", "valor"]), "número de células"),
        ("Excedente operacional", f"EOB ≥ {br(100 * eob_min, 0)}% do VBP em todos os subsetores com produção", eob_min,
         "fração do VBP"),
        ("Modelo", "A inversa de Leontief é não negativa, com diagonal ≥ 1, também com as colunas latentes",
         float(p.loc["L_nao_negativa", "valor"]), "menor elemento de L"),
        ("Reprodução do artigo", "Os nove cenários do artigo, rodados na MIP original, reproduzem os valores de "
         "referência; arredondados a uma casa, os Painéis B das Tabelas 1, 3 e 4 (exceto a linha R4a da Tabela 4) e a "
         "Tabela A2 são idênticos aos publicados", dif_a0, "t CO₂e (maior diferença nas emissões totais)"),
        ("Variantes", "As 9 especificações (principal e 8 variantes) passam nas mesmas checagens; nas variantes E3 e E4 "
         "as emissões preservam só o total", float(len(vr)), "especificações"),
        ("Alertas", f"Sem falhas. Alertas registrados: multiplicador de produção acima de {br(e['alerta_multiplicador_producao'], 1)} "
         f"em {n_mult} das 9 especificações (na principal: {'; '.join(mult)}); insumos não energéticos da geração a gás natural abaixo de "
         f"{br(100 * e['alerta_insumos_nao_energeticos'], 0)}% do VBP em {n_ne} (todas, exceto κ = {br(kappa_sens, 2)} e "
         f"η = {br(eta_sens, 2)})", float((c["status"] == "FALHA").sum()), "falhas"),
    ]
    return Tabela("Tabela A8", "Testes de consistência da desagregação",
                  [(None, pd.DataFrame(linhas, columns=["Teste", "Resultado", "Valor", "Unidade"]))],
                  casas={"Valor": 6}, fonte="Fonte: Elaboração própria.", aba="TabelaA8",
                  notas=["Nota: 'Valor' é a maior diferença absoluta encontrada (ou a contagem indicada); tolerância de "
                         "10⁻⁶ nas igualdades."])


def construir_tabelas(d: dict) -> list[Tabela]:
    ts = [quadro1(d), tabela1(d), tabela2(d), tabela3(d), tabela4(d), tabela5(d), tabela6(d), tabela7(d),
          tabelaA1(d), tabelaA2(d), tabelaA3(d), tabelaA4(d), tabelaA5(d), tabelaA6(d), tabelaA7(d), tabelaA8(d)]
    return [t for t in ts if t is not None]


# ---------------------------------------------------------------------------------------------
# escrita das tabelas: .xlsx, .md e .docx
# ---------------------------------------------------------------------------------------------
def escrever_xlsx(tabs: list[Tabela], caminho: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    wb.remove(wb.active)
    for t in tabs:
        ws = wb.create_sheet(t.aba[:31])
        ws.append([f"{t.numero}: {t.titulo}"])
        ws.cell(row=1, column=1).font = Font(bold=True)
        for nome, df in t.paineis:
            ws.append([])
            if nome:
                ws.append([nome])
                ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            ws.append(list(df.columns))
            for c in range(1, len(df.columns) + 1):
                ws.cell(row=ws.max_row, column=c).font = Font(bold=True)
                ws.cell(row=ws.max_row, column=c).alignment = Alignment(wrap_text=True, vertical="top")
            for _, r in df.iterrows():
                vals = [None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v)
                        for v in r.tolist()]
                ws.append(vals)
                for c, (col, v) in enumerate(zip(df.columns, vals), start=1):
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        casas = t.casas.get(col, 1)
                        ws.cell(row=ws.max_row, column=c).number_format = "#,##0" + ("." + "0" * casas if casas else "")
        ws.append([])
        for n in t.notas + [t.fonte]:
            ws.append([n])
            ws.cell(row=ws.max_row, column=1).font = Font(italic=True)
        ws.column_dimensions["A"].width = 40
        for col in "BCDEFGHIJKL":
            ws.column_dimensions[col].width = 18
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


def tabela_md(t: Tabela, ponto: bool = False) -> str:
    linhas = [f"**{t.numero}:** {t.titulo}", ""]
    for nome, df in t.paineis:
        if nome:
            linhas += [f"*{nome}*", ""]
        cols = list(df.columns)
        linhas.append("| " + " | ".join(cols) + " |")
        num = [pd.api.types.is_numeric_dtype(df[c]) and not pd.api.types.is_bool_dtype(df[c]) for c in cols]
        linhas.append("|" + "|".join("---:" if n else "---" for n in num) + "|")
        for _, r in df.iterrows():
            linhas.append("| " + " | ".join(t.celula(c, r[c], ponto).replace("|", "\\|") for c in cols) + " |")
        linhas.append("")
    linhas += [f"*{n}*" for n in t.notas + [t.fonte]]
    return "\n".join(linhas) + "\n"


def _bordas(tabela) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tblPr = tabela._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")  # oitavos de ponto: 4 = 0,5 pt
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)
    mar = OxmlElement("w:tblCellMar")  # margens internas estreitas (0,08 cm) para caber em A4 retrato
    for lado in ("left", "right"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), "45")
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def _fonte_run(run, tamanho, negrito=False, italico=False):
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(tamanho)
    run.bold = negrito
    run.italic = italico


def documento_base():
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(12)
    return doc


def escrever_docx(tabs: list[Tabela], caminho: Path) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = documento_base()
    for i, t in enumerate(tabs):
        if i:
            doc.add_paragraph()
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        _fonte_run(p.add_run(f"{t.numero}:"), 12, negrito=True)
        _fonte_run(p.add_run(f" {t.titulo}"), 12)
        grupos = []  # painéis consecutivos com as mesmas colunas vão para a mesma tabela do Word
        for nome, df in t.paineis:
            if grupos and list(grupos[-1][0][1].columns) == list(df.columns):
                grupos[-1].append((nome, df))
            else:
                grupos.append([(nome, df)])
        for g in grupos:
            cols = list(g[0][1].columns)
            tab = doc.add_table(rows=1 + sum(len(df) + (1 if nome else 0) for nome, df in g), cols=len(cols))
            tab.alignment = WD_TABLE_ALIGNMENT.CENTER
            _bordas(tab)
            for c, col in enumerate(cols):
                cel = tab.cell(0, c)
                _fonte_run(cel.paragraphs[0].add_run(str(col)), 8, negrito=True)
            _larguras(tab, [str(c) for c in cols], [[t.celula(c, v) for c, v in zip(cols, r)] for _, df in g
                                                     for r in df.itertuples(index=False)])
            lin = 1
            for nome, df in g:
                if nome:
                    cel = tab.cell(lin, 0).merge(tab.cell(lin, len(cols) - 1))
                    for extra in cel.paragraphs[1:]:  # a fusão junta um parágrafo vazio de cada célula
                        extra._element.getparent().remove(extra._element)
                    _fonte_run(cel.paragraphs[0].add_run(nome), 8, negrito=True)
                    lin += 1
                for _, r in df.iterrows():
                    for c, col in enumerate(cols):
                        cel = tab.cell(lin, c)
                        par = cel.paragraphs[0]
                        v = r[col]
                        if isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, bool):
                            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        _fonte_run(par.add_run(t.celula(col, v)), 8)
                    lin += 1
            if g is not grupos[-1]:
                doc.add_paragraph()
        for n in t.notas + [t.fonte]:
            q = doc.add_paragraph()
            q.paragraph_format.space_after = Pt(0)
            _fonte_run(q.add_run(n), 10, italico=True)
    for tab in doc.tables:  # células compactas: sem espaço entre parágrafos, entrelinha simples
        for row in tab.rows:
            for cel in row.cells:
                for par in cel.paragraphs:
                    par.paragraph_format.space_before = Pt(0)
                    par.paragraph_format.space_after = Pt(0)
                    par.paragraph_format.line_spacing = 1.0
    caminho.parent.mkdir(parents=True, exist_ok=True)
    doc.save(caminho)


def _larguras(tab, cabecalho: list, linhas: list, total_cm: float = 16.0) -> None:
    """Largura das colunas: cada uma tem pelo menos a da maior palavra (números não quebram) e cresce com o tamanho
    médio do texto; o conjunto ocupa a largura útil da página A4 (Times New Roman 8: ~0,15 cm por caractere)."""
    from docx.shared import Cm

    cm_car, folga = 0.145, 0.22
    n = len(cabecalho)
    minimo, desejado = [], []
    for j in range(n):
        textos = [str(lin[j]) for lin in linhas if j < len(lin)]
        palavras = [w for x in textos + [cabecalho[j]] for w in re.split(r"[\s/-]+", x) if w]  # quebra em hífen e barra
        maior = max([len(w) for w in palavras] or [3])
        medio = (sum(len(x) for x in textos) / len(textos)) if textos else 3
        minimo.append(cm_car * maior + folga)
        desejado.append(max(minimo[-1], cm_car * min(medio * 1.1, 45) + folga, cm_car * len(cabecalho[j]) / 3 + folga))
    if sum(desejado) > total_cm and sum(minimo) < total_cm:  # corta só o que passa do mínimo
        extra = [d - m for d, m in zip(desejado, minimo)]
        f = (total_cm - sum(minimo)) / sum(extra)
        larg = [m + e * f for m, e in zip(minimo, extra)]
    else:
        larg = [d * total_cm / sum(desejado) for d in desejado]
    tab.autofit = False
    for j, col in enumerate(tab.columns):  # grade da tabela (lida pelo LibreOffice) e largura de cada célula (Word)
        col.width = Cm(larg[j])
    for row in tab.rows:
        for j, cel in enumerate(row.cells[:n]):
            cel.width = Cm(larg[j])


# ---------------------------------------------------------------------------------------------
# figuras
# ---------------------------------------------------------------------------------------------
def _estilo():
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8.5, "axes.edgecolor": "#888888",
                         "axes.labelcolor": "#222222", "xtick.color": "#222222", "ytick.color": "#222222",
                         "figure.facecolor": "white", "axes.facecolor": "white", "savefig.facecolor": "white",
                         "pdf.fonttype": 42})
    return plt


def _eixos(ax, eixo_valor="x"):
    for lado in ("top", "right"):
        ax.spines[lado].set_visible(False)
    ax.grid(axis=eixo_valor, color=GRADE, linewidth=0.8)
    ax.set_axisbelow(True)


def _salvar(fig, nome) -> list[Path]:
    FIG.mkdir(parents=True, exist_ok=True)
    saidas = [FIG / f"{nome}.png", FIG / f"{nome}.pdf"]
    fig.savefig(saidas[0], dpi=300, bbox_inches="tight", metadata={"Software": None})
    fig.savefig(saidas[1], bbox_inches="tight", metadata={"CreationDate": None, "ModDate": None, "Producer": None})
    return saidas


def figura_top10(d, aba, cenario, nome, plt) -> list[Path]:
    top = d[aba]
    top = top[top["cenario"] == cenario].sort_values("#", ascending=False)
    fig, ax = plt.subplots(figsize=(7, 4.4))
    _eixos(ax, "x")
    rot = ["\n".join(textwrap.wrap(f"{NOMES[c]} ({r})", 30)) for c, r in zip(top["codigo"], top["regiao"])]
    cores = [AZUL if r == "RS" else LARANJA for r in top["regiao"]]
    y = np.arange(len(top))
    ax.barh(y, top["% do CO2e total"], color=cores, height=0.62, edgecolor="white", linewidth=0.6)
    for yi, v in zip(y, top["% do CO2e total"]):
        ax.text(v + 0.4, yi, pct(v), va="center", fontsize=7.5, color="#222222")
    ax.set_yticks(y, rot, fontsize=7.2)
    ax.set_xlabel("Participação nas emissões induzidas pela cesta renovável (%)")
    ax.set_xlim(0, top["% do CO2e total"].max() * 1.15)
    ax.xaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: br(v, 0)))
    from matplotlib.patches import Patch

    ax.legend(handles=[Patch(color=AZUL, label="Rio Grande do Sul (RS)"), Patch(color=LARANJA, label="Restante do Brasil (RB)")],
              frameon=False, loc="lower right", fontsize=8)
    fig.tight_layout()
    s = _salvar(fig, nome)
    plt.close(fig)
    return s


def figura2(d, plt) -> list[Path]:
    di = d["diretas_indiretas"].set_index(["versao", "cenario"])
    fig, ax = plt.subplots(figsize=(7, 4.4))
    _eixos(ax, "y")
    larg = 0.36
    x = np.arange(3)
    for desl, versao, pref, rot in ((-0.2, "original", "A0", "MIP original"), (0.2, V, "A2", "MIP desagregada")):
        dir_ = np.array([di.loc[(versao, f"{pref}_{c}"), "pct_diretas"] for c in PRINCIPAIS])
        ind_ = 100 - dir_
        ax.bar(x + desl, dir_, larg, color=AZUL, edgecolor="white", linewidth=0.6)
        ax.bar(x + desl, ind_, larg, bottom=dir_, color=LARANJA, edgecolor="white", linewidth=0.6)
        for xi, a, b in zip(x + desl, dir_, ind_):
            for base, v in ((0, a), (a, b)):
                if v >= 4:
                    ax.text(xi, base + v / 2, pct(v), ha="center", va="center", color="white", fontweight="bold",
                            fontsize=7.5)
            ax.text(xi, 101.5, rot, ha="center", va="bottom", fontsize=7, color="#222222")
    ax.set_xticks(x, list(PRINCIPAIS.values()))
    ax.set_ylim(0, 108)
    ax.set_yticks(np.arange(0, 101, 20))
    ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: br(v, 0) + "%"))
    ax.set_ylabel("Participação nas emissões induzidas (%)")
    from matplotlib.patches import Patch

    ax.legend(handles=[Patch(color=AZUL, label="Emissões diretas"), Patch(color=LARANJA, label="Emissões indiretas")],
              frameon=False, loc="upper center", bbox_to_anchor=(0.5, -0.1), ncol=2, fontsize=8)
    fig.tight_layout()
    s = _salvar(fig, "figura_2_diretas_indiretas")
    plt.close(fig)
    return s


def figura3(d, plt) -> list[Path]:
    tw = d["B_TWh"]
    tw = tw[tw["versao"] == V].set_index("subsetor (RS)")
    fig, axs = plt.subplots(1, 2, figsize=(7, 4.4), gridspec_kw={"width_ratios": [1, 1.25]})
    grupos = ((axs[0], [("S38_GN", "Gás natural"), ("S38_CAR", "Carvão mineral")], "Fósseis"),
              (axs[1], [("S38_HID", "Hidrelétrica"), ("S38_SOL", "Solar fotovoltaica"), ("S38_EOL", "Eólica"),
                        ("S38_BIO", "Biomassa")], "Demais fontes"))
    for ax, itens, titulo in grupos:
        _eixos(ax, "x")
        y = np.arange(len(itens))
        dd = np.array([tw.loc[s, "diretas"] for s, _ in itens]) / 1000
        ii = np.array([tw.loc[s, "indiretas"] for s, _ in itens]) / 1000
        ax.barh(y, dd, 0.55, color=AZUL, edgecolor="white", linewidth=0.6)
        ax.barh(y, ii, 0.55, left=dd, color=LARANJA, edgecolor="white", linewidth=0.6)
        for yi, v in zip(y, dd + ii):
            ax.text(v * 1.02, yi, br(v, 1), va="center", fontsize=7.5)
        ax.set_yticks(y, [n for _, n in itens])
        ax.set_xlim(0, (dd + ii).max() * 1.22)
        ax.xaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: br(v, 0)))
        ax.set_xlabel("mil t CO₂e por TWh")
        ax.text(0, 1.02, titulo, transform=ax.transAxes, fontsize=8.5, color="#222222")
    from matplotlib.patches import Patch

    fig.legend(handles=[Patch(color=AZUL, label="Emissões diretas"), Patch(color=LARANJA, label="Emissões indiretas")],
               frameon=False, loc="lower center", ncol=2, fontsize=8, bbox_to_anchor=(0.5, -0.02))
    fig.tight_layout(rect=(0, 0.06, 1, 1))
    s = _salvar(fig, "figura_3_emissoes_por_TWh")
    plt.close(fig)
    return s


ROTULO_FIG4 = {"E1": "Energia\nelétrica\nagregada", "R1": "Renovável\nbase", "R2": "Renovável\nbalanceada",
               "R3": "Renovável\ninfraestrutura-\nintensiva", "R4a": "Renovável\nenergia/\ntransporte-\nintensiva",
               "F1": "Fóssil\nbase", "F2": "Fóssil\nbalanceada", "F3": "Fóssil\nenergia-\nintensiva"}


def figura4(d, plt) -> list[Path]:
    fig, axs = plt.subplots(1, 2, figsize=(7, 4.4), gridspec_kw={"width_ratios": [6, 3]})
    grupos = ((axs[0], ["E1"] + RENOV, "Energia elétrica agregada e cestas renováveis"),
              (axs[1], FOSS, "Cestas fósseis"))
    larg = 0.38
    for ax, cs, titulo in grupos:
        _eixos(ax, "y")
        x = np.arange(len(cs))
        a0 = np.array([ind(d, "original", f"A0_{c}")["CO₂e total (t)"] for c in cs]) / 1000
        a2 = np.array([ind(d, V, f"A2_{c}")["CO₂e total (t)"] for c in cs]) / 1000
        ax.bar(x - larg / 2 - 0.01, a0, larg, color=AZUL, edgecolor="white", linewidth=0.6)
        ax.bar(x + larg / 2 + 0.01, a2, larg, color=LARANJA, edgecolor="white", linewidth=0.6)
        for xi, v in zip(x - larg / 2 - 0.01, a0):
            ax.text(xi, v, br(v, 0), ha="center", va="bottom", fontsize=6.5)
        for xi, v in zip(x + larg / 2 + 0.01, a2):
            ax.text(xi, v, br(v, 0), ha="center", va="bottom", fontsize=6.5)
        ax.set_xticks(x, [ROTULO_FIG4[c] for c in cs], fontsize=6.2)
        ax.set_ylim(0, max(a0.max(), a2.max()) * 1.12)
        ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: br(v, 0)))
        ax.text(0, 1.02, titulo, transform=ax.transAxes, fontsize=8.5, color="#222222")
    axs[0].set_ylabel("Emissões totais (mil t CO₂e por R$ 1 bilhão)")
    from matplotlib.patches import Patch

    fig.legend(handles=[Patch(color=AZUL, label="MIP original"), Patch(color=LARANJA, label="MIP desagregada")],
               frameon=False, loc="lower center", ncol=2, fontsize=8, bbox_to_anchor=(0.5, -0.02))
    fig.tight_layout(rect=(0, 0.06, 1, 1))
    s = _salvar(fig, "figura_4_emissoes_por_cesta")
    plt.close(fig)
    return s


def figuras(d) -> list[Path]:
    plt = _estilo()
    return (figura_top10(d, "top10_A2", "A2_R1", "figura_1_top10_cesta_renovavel", plt)
            + figura_top10(d, "top10_A0", "A0_R1", "figura_A1_top10_cesta_renovavel_original", plt)
            + figura2(d, plt) + figura3(d, plt) + figura4(d, plt))


# ---------------------------------------------------------------------------------------------
# números-chave e textos
# ---------------------------------------------------------------------------------------------
def numeros_chave(d: dict) -> dict:
    """Números que provavelmente serão citados: valor bruto, casas, unidade e origem (arquivo › aba › linha)."""
    N = {}
    sim, qd, dg = f"resultados/rodada05/simulacoes_{V}.xlsx", f"resultados/rodada05/quadro_subsetores_{V}.xlsx", \
        "resultados/rodada05/diagnostico_seeg_nereus.xlsx"

    def put(chave, valor, casas, unidade, origem, descricao):
        N[chave] = {"valor": float(valor), "casas": casas, "unidade": unidade, "origem": origem, "descricao": descricao}

    for versao, pref, aba in (("original", "A0", "A0"), (V, "A2", "robustez_detalhe")):
        for c in ["E1"] + RENOV + FOSS:
            x = ind(d, versao, f"{pref}_{c}")
            org = f"{sim} › {aba} › {pref}_{c}" + (f" (versão {versao})" if versao == V else "")
            k = f"{pref.lower()}_{c.lower()}"
            put(f"{k}_co2e", x["CO₂e total (t)"], 1, "t CO₂e", org, f"emissões totais, {OITO[c]} ({pref})")
            put(f"{k}_dx", x["ΔX total (R$ mi)"], 1, "R$ milhões", org, f"produção bruta induzida, {OITO[c]} ({pref})")
            put(f"{k}_ic", x["IC (t CO₂e por R$ mi)"], 1, "t CO₂e/R$ milhão", org, f"intensidade carbônica, {OITO[c]} ({pref})")
            put(f"{k}_pct_rb", x["% CO₂e no RB"], 1, "%", org, f"emissões no RB, {OITO[c]} ({pref})")
            put(f"{k}_pct_ind", x["% indiretas"], 1, "%", org, f"emissões indiretas, {OITO[c]} ({pref})")
            put(f"{k}_pct_s38", x["% das emissões no S38 (subsetores)"], 1, "%", org,
                f"participação do S38 (ou dos subsetores) nas emissões, {OITO[c]} ({pref})")
    for c in ["E1"] + RENOV + FOSS:
        x = imp(d, "A1", f"A1_{c}")
        put(f"a1_{c.lower()}_co2e", x["CO₂e total (t)"], 1, "t CO₂e", f"{sim} › A1 › A1_{c}", f"emissões totais, {OITO[c]} (A1)")
        put(f"var_a1_a0_{c.lower()}", 100 * (x["CO₂e total (t)"] / N[f"a0_{c.lower()}_co2e"]["valor"] - 1), 1, "%",
            f"{sim} › A1 e A0 › {c}", f"variação das emissões A1/A0, {OITO[c]}")
        put(f"var_a2_a0_{c.lower()}", 100 * (N[f"a2_{c.lower()}_co2e"]["valor"] / N[f"a0_{c.lower()}_co2e"]["valor"] - 1), 1,
            "%", f"{sim} › A2 e A0 › {c}", f"variação das emissões A2/A0, {OITO[c]}")
    for pref in ("a0", "a2"):
        put(f"{pref}_dif_fossil_renov", N[f"{pref}_f1_co2e"]["valor"] - N[f"{pref}_r1_co2e"]["valor"], 1, "t CO₂e",
            f"{sim} › {pref.upper()} › F1 − R1", f"diferença de emissões fóssil − renovável ({pref.upper()})")
        put(f"{pref}_reducao_ic", 100 * (1 - N[f"{pref}_r1_ic"]["valor"] / N[f"{pref}_f1_ic"]["valor"]), 1, "%",
            f"{sim} › {pref.upper()} › F1 e R1", f"redução da intensidade carbônica da cesta renovável em relação à fóssil ({pref.upper()})")
    top = d["top10_A2"]
    t1 = top[(top["cenario"] == "A2_R1") & (top["#"] == 1)].iloc[0]
    put("a2_r1_top1_pct", t1["% do CO2e total"], 1, "%", f"{sim} › top10_A2 › A2_R1 #1",
        f"participação do 1º emissor da cesta renovável no A2 ({t1['regiao']} {t1['codigo']})")
    tw = d["B_TWh"]
    tw = tw[tw["versao"] == V].set_index("subsetor (RS)")
    b1 = d["B_RS1bi"].set_index("cenario")
    for s, nome in TEC_B:
        k = s.replace("S38_", "").lower()
        org = f"{sim} › B_TWh › {s} (v04)"
        put(f"twh_{k}_t", tw.loc[s, "t CO₂e por TWh (total)"], 1, "t CO₂e/TWh", org, f"emissões por TWh, {nome}")
        put(f"twh_{k}_gkwh", tw.loc[s, "g CO₂e/kWh"], 1, "g CO₂e/kWh", org, f"g CO₂e/kWh, {nome}")
        put(f"twh_{k}_preco", tw.loc[s, "R$/MWh = choque de 1 TWh (R$ mi)"], 1, "R$/MWh", org, f"preço médio, {nome}")
        put(f"twh_{k}_pct_rb", tw.loc[s, "% no RB"], 1, "%", org, f"emissões no RB por TWh, {nome}")
        put(f"twh_{k}_evit_car", tw.loc[s, "evitadas por TWh se substituir o carvão (t)"], 1, "t CO₂e/TWh", org,
            f"emissões evitadas por TWh em relação ao carvão, {nome}")
        put(f"twh_{k}_evit_gn", tw.loc[s, "evitadas por TWh se substituir o gás natural (t)"], 1, "t CO₂e/TWh", org,
            f"emissões evitadas por TWh em relação ao gás natural, {nome}")
        put(f"bi_{k}_t", b1.loc[f"B_{s.replace('S38_', '')}", "CO₂e total (t)"], 1, "t CO₂e", f"{sim} › B_RS1bi › B_{k.upper()}",
            f"emissões por R$ 1 bilhão, {nome}")
    put("bi_dis_t", b1.loc["B_DIS", "CO₂e total (t)"], 1, "t CO₂e", f"{sim} › B_RS1bi › B_DIS",
        "emissões por R$ 1 bilhão, distribuição (mix de 2019)")
    # robustez: faixa entre as variantes
    for c in ("A2_E1", "A2_R1", "A2_F1") + tuple(f"B_TWh_{s.replace('S38_', '')}" for s, _ in TEC_B):
        vals = {v: ind(d, v, c)["CO₂e total (t)"] for v in VERSOES}
        k = c.lower()
        put(f"rob_{k}_min", min(vals.values()), 1, "t CO₂e", f"{sim} › robustez_detalhe › {c}",
            f"menor valor entre as 9 versões ({min(vals, key=vals.get)}), {c}")
        put(f"rob_{k}_max", max(vals.values()), 1, "t CO₂e", f"{sim} › robustez_detalhe › {c}",
            f"maior valor entre as 9 versões ({max(vals, key=vals.get)}), {c}")
        for v in ("v04_E3", "v04_E4", "v04_P2", "v04_k060", "v04_rap_media"):
            put(f"dev_{v.replace('v04_', '').lower()}_{k}", 100 * (vals[v] / vals[V] - 1), 1, "%",
                f"{sim} › robustez_detalhe › {c}", f"desvio das emissões da {v} em relação à principal, {c}")
    for c in ("E1", "R1", "F1"):
        x = imp(d, "A0_E4", f"A0_{c}")
        put(f"a0e4_{c.lower()}_var", x["variação de CO₂e total (t) (%)"], 1, "%", f"{sim} › A0_E4 › A0_{c}",
            f"variação das emissões com a E4 na MIP original, {OITO[c]}")
    # MIP: transmissão, D18, gás, subsetores
    vr = d["q_variantes_resumo"].set_index("versao")
    put("vbp_tra_total", vr.loc[V, "VBP TRA-RS"] + vr.loc[V, "VBP TRA-RB"], 1, "R$ milhões",
        f"{qd} › variantes_resumo › v04", "VBP da transmissão (Brasil, D17)")
    put("vbp_tra_rs", vr.loc[V, "VBP TRA-RS"], 1, "R$ milhões", f"{qd} › variantes_resumo › v04", "VBP da transmissão, RS")
    put("vbp_tra_media", vr.loc["v04_rap_media", "VBP TRA-RS"] + vr.loc["v04_rap_media", "VBP TRA-RB"], 1, "R$ milhões",
        f"{qd} › variantes_resumo › v04_rap_media", "VBP da transmissão na variante _rap_media")
    put("vbp_gas_rb", vr.loc[V, "VBP GAS-RB"], 1, "R$ milhões", f"{qd} › variantes_resumo › v04", "VBP do gás canalizado, RB")
    d18 = d["q_d18_latentes"].set_index("subsetor")
    put("tau_rs", d18.loc["S38_SOL", "tau (R$/MWh)"], 2, "R$/MWh", f"{qd} › d18_latentes", "TUST por MWh das fontes do RS (τ_RS)")
    for s in ("S38_SOL", "S38_PET"):
        k = s.replace("S38_", "").lower()
        put(f"p_{k}_rb", d18.loc[s, "p (R$/MWh)"], 1, "R$/MWh", f"{qd} › d18_latentes › {s}", f"R$/MWh do espelho {s}-RB")
        put(f"coef_tust_{k}_antes", d18.loc[s, "coeficiente de TUST antes"], 4, "R$/R$", f"{qd} › d18_latentes › {s}",
            f"coeficiente de TUST da coluna latente {s}-RS sem a D18")
        put(f"coef_tust_{k}_depois", d18.loc[s, "coeficiente de TUST depois"], 4, "R$/R$", f"{qd} › d18_latentes › {s}",
            f"coeficiente de TUST da coluna latente {s}-RS com a D18")
    q = d["q_quadro"]
    for r in ("RS", "RB"):
        for s in SUBSETORES:
            x = q[(q["versao"] == V) & (q["regiao"] == r) & (q["subsetor"] == s)].iloc[0]
            k = f"{r.lower()}_{s.replace('S38_', '').lower()}"
            put(f"vbp_{k}", x["VBP (R$ mi)"], 1, "R$ milhões", f"{qd} › quadro › v04 {r} {s}", f"VBP {r} {s}")
            if pd.notna(x["multiplicador de emissões (t/R$ mi)"]):
                put(f"mult_em_{k}", x["multiplicador de emissões (t/R$ mi)"], 1, "t CO₂e/R$ milhão",
                    f"{qd} › quadro › v04 {r} {s}", f"multiplicador de emissões {r} {s}")
                put(f"mult_prod_{k}", x["multiplicador de produção"], 3, "", f"{qd} › quadro › v04 {r} {s}",
                    f"multiplicador de produção {r} {s}")
            if pd.notna(x["t CO2e/MWh E1"]):
                put(f"fator_e1_{k}", x["t CO2e/MWh E1"], 3, "t CO₂e/MWh", f"{qd} › quadro › v04 {r} {s}",
                    f"fator de emissão E1 {r} {s}")
    niv = d["q_niveis_emissoes"]
    niv = niv[niv["versao"] == V].set_index("regiao")
    for r in ("RS", "RB"):
        put(f"residuo_{r.lower()}", niv.loc[r, "residuo"], 1, "t CO₂e", f"{qd} › niveis_emissoes › v04 {r}",
            f"resíduo de emissões do S38 (E1), {r}")
        put(f"residuo_{r.lower()}_pct", 100 * niv.loc[r, "residuo"] / niv.loc[r, "total_base"], 1, "%",
            f"{qd} › niveis_emissoes › v04 {r}", f"resíduo de emissões em % do CO₂e do S38, {r}")
        put(f"s38_co2e_{r.lower()}", niv.loc[r, "total_base"], 1, "t CO₂e", f"{qd} › niveis_emissoes › v04 {r}",
            f"emissões do S38 na base NEREUS, {r}")
    rea = d["q_realocacao_E4"]
    for r in ("RS", "RB"):
        x = rea[(rea["regiao"] == r) & ~rea["setor_destino"].astype(str).str.startswith("(")]
        put(f"e4_movido_{r.lower()}", x["valor"].sum(), 1, "t CO₂e", f"{qd} › realocacao_E4 › {r}",
            f"emissões movidas do S38 na E4, {r}")
    for r in ("RS", "RB"):
        movido = N[f"e4_movido_{r.lower()}"]["valor"]
        put(f"e4_sobra_{r.lower()}", N[f"residuo_{r.lower()}"]["valor"] - movido, 1, "t CO₂e",
            f"{qd} › niveis_emissoes e realocacao_E4 › {r}", f"resíduo de emissões que fica em DIS e GAS na E4, {r}")
    # parâmetros (das especificações e dos CSVs de parâmetros)
    e = d["espec"]
    put("param_kappa", e["piso_vbp"]["kappa"], 2, "", f"parametros/especificacao_{V}.yaml › piso_vbp.kappa", "κ (piso do VBP das térmicas)")
    put("param_eta", e["reequilibrio"]["eta"], 2, "", f"parametros/especificacao_{V}.yaml › reequilibrio.eta", "η (EOB mínimo)")
    for suf, chave, campo in (("_k060", "param_kappa_sens", ("piso_vbp", "kappa")), ("_eta005", "param_eta_sens", ("reequilibrio", "eta"))):
        with open(PASTA_PARAMETROS / f"especificacao_{V}{suf}.yaml", encoding="utf-8") as f:
            es = yaml.safe_load(f)
        put(chave, es[campo[0]][campo[1]], 2, "", f"parametros/especificacao_{V}{suf}.yaml › {'.'.join(campo)}",
            f"{campo[1]} na variante {suf}")
    comp = pd.read_csv(PASTA_PARAMETROS / f"componentes_{V}.csv")
    put("param_alfa", comp.loc[comp["componente"] == "fracao_transmissao_distribuicao", "valor"].astype(float).iloc[0], 1,
        "", f"parametros/componentes_{V}.csv › fracao_transmissao_distribuicao", "parcela da transmissão paga pela distribuição (α)")
    lin = pd.read_csv(PASTA_PARAMETROS / f"regras_linha_{V}.csv")
    g = lin.loc[lin["subsetor"] == "S38_GAS", "valor"].astype(float)
    put("param_teto_gas", g.max(), 2, "", f"parametros/regras_linha_{V}.csv › S38_GAS (máximo)",
        "participação máxima do gás canalizado nas compras de S38 de um comprador (teto)")
    for r in ("RS", "RB"):
        o = q[(q["versao"] == "original") & (q["regiao"] == r)].iloc[0]
        put(f"vbp_s38_{r.lower()}", o["VBP (R$ mi)"], 1, "R$ milhões", f"{qd} › quadro › original {r}", f"VBP do S38 na MIP original, {r}")
        put(f"mult_em_s38_{r.lower()}", o["multiplicador de emissões (t/R$ mi)"], 1, "t CO₂e/R$ milhão",
            f"{qd} › quadro › original {r}", f"multiplicador de emissões do S38 na MIP original, {r}")
        ger = q[(q["versao"] == V) & (q["regiao"] == r) & (q["tipo"] == "geracao")]
        put(f"vbp_ger_{r.lower()}", ger["VBP (R$ mi)"].sum(), 1, "R$ milhões", f"{qd} › quadro › v04 {r} geração",
            f"VBP da geração (8 subsetores), {r}")
        put(f"gwh_ger_{r.lower()}", ger["GWh"].sum(), 1, "GWh", f"{qd} › quadro › v04 {r} geração", f"GWh de 2019, {r}")
    ne = d["q_insumos_nao_energeticos"]
    ne = ne[ne["versao"] == V].set_index(["regiao", "subsetor"])
    for r in ("RS", "RB"):
        put(f"t12_gn_{r.lower()}", ne.loc[(r, "S38_GN"), "insumos_nao_energeticos / VBP"], 3, "fração do VBP",
            f"{qd} › insumos_nao_energeticos › v04 {r} S38_GN", f"insumos não energéticos ÷ VBP da geração a gás, {r} (T12)")
    put("t12_min_k060", vr.loc["v04_k060", "insumos não energéticos / VBP mínimo"], 3, "fração do VBP",
        f"{qd} › variantes_resumo › v04_k060", "menor razão insumos não energéticos ÷ VBP com κ = 0,60")
    for chave, campo, casas, descricao in (
            ("param_alerta_ne", "alerta_insumos_nao_energeticos", 2, "limiar do ALERTA do T12 (insumos não energéticos ÷ VBP)"),
            ("param_alerta_mult", "alerta_multiplicador_producao", 1, "limiar do alerta do multiplicador de produção")):
        put(chave, e[campo], casas, "", f"parametros/especificacao_{V}.yaml › {campo}", descricao)
    vd = pd.read_csv(RAIZ / "dados" / "externos" / "valores_documentais_v03.csv").set_index("parametro")
    m = re.search(r"×\s*([0-9.]+)\s*m³/kWh", str(vd.loc["theta_RS"].to_dict()))
    if m is None:
        raise SystemExit("valores_documentais_v03.csv › theta_RS: consumo específico (m³/kWh) não encontrado")
    put("param_consumo_gas_rs", float(m.group(1)), 2, "m³/kWh", "dados/externos/valores_documentais_v03.csv › theta_RS",
        "consumo específico de gás das térmicas usado no θ_RS (HIPÓTESE)")
    for s in ("S38_SOL", "S38_PET"):
        k = s.replace("S38_", "").lower()
        put(f"eob_impl_{k}_rs", 100 * d18.loc[s, "EOB/VBP implícito da latente"], 1, "% do VBP",
            f"{qd} › d18_latentes › {s}", f"EOB/VBP implícito da coluna latente {s}-RS com a D18 (espelho − Δ TUST)")
        put(f"eob_espelho_{k}_rb", 100 * d18.loc[s, "EOB/VBP do espelho"], 1, "% do VBP",
            f"{qd} › d18_latentes › {s}", f"EOB/VBP do espelho {s}-RB")
    put("mult_em_rs_sol_sem_d18", d18.loc["S38_SOL", "multiplicador de emissões (t/R$ mi) sem D18"], 1, "t CO₂e/R$ milhão",
        f"{qd} › d18_latentes › S38_SOL", "multiplicador de emissões do SOL-RS latente sem a D18")
    put("mult_prod_rs_sol_sem_d18", d18.loc["S38_SOL", "multiplicador de produção sem D18"], 3, "",
        f"{qd} › d18_latentes › S38_SOL", "multiplicador de produção do SOL-RS latente sem a D18")
    sol = d["d18_SOL"].set_index("versao")
    put("twh_sol_sem_d18", sol.loc[f"{V} sem D18", "t CO₂e por TWh (total)"], 1, "t CO₂e/TWh", f"{sim} › d18_SOL",
        "emissões por TWh da solar sem a D18 (TUST por real do espelho)")
    chk = d["checagens"]
    chk = chk[chk["versao"] == V].set_index("checagem")
    put("reag_max", max(float(chk.loc[f"reagregacao_{n}", "valor"]) for n in ("Z", "F", "X", "total_linha", "V")), 9,
        "R$ milhões", f"resultados/rodada05/checagens_{V}.csv › reagregacao_*", "maior diferença da re-agregação")
    a0 = d["A0"].set_index("cenario")
    put("a0_dif_artigo", max(abs(a0.loc[f"A0_{x.cenario}", "CO₂e total (t)"] - x.co2e_total) for x in ref.ESPERADO.itertuples()),
        2, "t CO₂e", f"{sim} › A0 × src/mip/referencia.py", "maior diferença entre o A0 e os valores publicados")
    if "diag" in d and "resumo" in d["diag"]:
        rs = _com_cabecalho(d["diag"]["resumo"], "item")
        rs = rs.dropna(subset=["valor"])
        val = {str(i): float(v) for i, v in zip(rs["item"], rs["valor"]) if isinstance(v, (int, float, np.number))}

        def diag(chave, prefixo, casas, unidade, descricao, exato=False):
            k = next((i for i in val if (i == prefixo if exato else i.startswith(prefixo))), None)
            if k is None:
                raise SystemExit(f"diagnóstico: '{prefixo}' não encontrado na aba resumo")
            put(chave, val[k], casas, unidade, f"{dg} › resumo › {k}", descricao)

        diag("diag_s38_rs_seeg", "S38-RS: SEEG", 1, "mil t CO₂e", "S38-RS pelo SEEG (geração de eletricidade + produção de combustíveis)")
        diag("diag_s38_rs_nereus", "S38-RS: NEREUS", 1, "mil t CO₂e", "S38-RS na base NEREUS")
        diag("diag_s38_rs_dif", "S38-RS: diferença", 2, "%", "diferença SEEG − NEREUS no S38-RS")
        diag("diag_s38_rb_seeg", "S38-RB: SEEG", 1, "mil t CO₂e", "S38-RB pelo SEEG (estados)")
        diag("diag_s38_rb_nereus", "S38-RB: NEREUS", 1, "mil t CO₂e", "S38-RB na base NEREUS")
        diag("diag_s38_rb_dif", "S38-RB: diferença", 2, "%", "diferença SEEG − NEREUS no S38-RB")
        diag("diag_nao_alocado", "Não alocado", 1, "mil t CO₂e", "'Não alocado' das duas atividades")
        diag("diag_testes", "Casamentos: total de testes", 0, "combinações", "combinações testadas no casamento de setores")
        diag("diag_casamentos", "Casamentos: |dif| < 0.1% (todos)", 0, "combinações", "combinações com diferença < 0,1%")
        diag("diag_acaso", "Casamentos: |dif| < 0.1% esperados", 0, "combinações", "combinações < 0,1% esperadas por acaso")
        diag("diag_total_br_liq", "Total Brasil: SEEG GWP-AR5 líquida", 1, "%", "total Brasil, SEEG AR5 líquida − NEREUS")
        diag("diag_total_br_bruta", "Total Brasil: SEEG GWP-AR5 bruta", 1, "%", "total Brasil, SEEG AR5 bruta − NEREUS")
        diag("diag_nucleos", "Casamentos: |dif| < 0.1%, núcleos distintos", 0, "combinações",
             "núcleos distintos entre as combinações < 0,1% (sem unidades desprezíveis)")
        diag("diag_sentido_sim", "Casamentos: |dif| < 0.1% com sentido 'sim'", 0, "combinações",
             "combinações < 0,1% com sentido econômico (critério SCN)")
        diag("diag_discutivel", "Casamentos: |dif| < 0.1% 'discutível'", 0, "combinações",
             "combinações < 0,1% com sentido discutível (critério SCN)")
        diag("diag_concorrente_dif", "Casamento concorrente", 3, "%", "diferença do casamento concorrente S19-RS ≈ refino + química")
        diag("diag_s19_rs_nereus", "S19-RS: NEREUS", 1, "mil t CO₂e", "S19-RS na base NEREUS", exato=True)
        diag("diag_s19_rs_e4", "S19-RS: NEREUS + E4", 1, "mil t CO₂e", "S19-RS na E4 (base NEREUS + refino do SEEG do RS)")
    return N


def escrever_numeros(N: dict) -> None:
    TXT.mkdir(parents=True, exist_ok=True)
    with open(TXT / "numeros_chave.json", "w", encoding="utf-8") as f:
        json.dump(N, f, ensure_ascii=False, indent=1, sort_keys=True)
    linhas = ["# Números-chave", "",
              "Gerado por `scripts/rodada05_artigo.py` a partir dos arquivos de resultado. Cada número tem o valor bruto, o "
              "valor no padrão brasileiro e a origem (arquivo › aba › linha). Nos textos, a chave entre chaves duplas "
              "(`{{chave}}`) é trocada pelo valor no padrão brasileiro.", "",
              "| Chave | Descrição | Valor bruto | Padrão brasileiro | Unidade | Origem |", "|---|---|---:|---:|---|---|"]
    for k in sorted(N):
        x = N[k]
        esc = lambda t: str(t).replace("|", "\\|").replace("$", "\\$")  # noqa: E731 (R$ não vira LaTeX no pandoc)
        linhas.append(f"| `{k}` | {esc(x['descricao'])} | {x['valor']!r} | {br(x['valor'], x['casas'])} | {esc(x['unidade'])} | "
                      f"{esc(x['origem'])} |")
    (TXT / "numeros_chave.md").write_text("\n".join(linhas) + "\n", encoding="utf-8")


PADRAO = re.compile(r"\{\{\s*([a-z0-9_]+)\s*(?::\s*(\d+))?\s*\}\}")


def renderizar(texto: str, N: dict, nome: str) -> str:
    faltam = sorted({m.group(1) for m in PADRAO.finditer(texto) if m.group(1) not in N})
    if faltam:
        raise SystemExit(f"{nome}: chaves desconhecidas {faltam} (ver artigo/texto/numeros_chave.md)")
    return PADRAO.sub(lambda m: br(N[m.group(1)]["valor"], int(m.group(2)) if m.group(2) else N[m.group(1)]["casas"]),
                      texto)


def md_para_docx(md: Path, docx: Path) -> None:
    """Markdown (com LaTeX) → .docx pelo pandoc (equações nativas do Word); depois A4, margens e Times New Roman."""
    import pypandoc
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    pypandoc.convert_file(str(md), "docx", outputfile=str(docx), extra_args=["--from=markdown+tex_math_dollars"])
    doc = Document(str(docx))
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5)
    for st in doc.styles:
        try:
            if st.font is not None:
                st.font.name = "Times New Roman"
                rpr = st.element.get_or_add_rPr()
                rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
                rpr.get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
                rpr.get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
                st.font.size = Pt(12)
                st.font.color.rgb = None
        except (AttributeError, ValueError, KeyError):
            continue
    for tab in doc.tables:
        for row in tab.rows:
            for cel in row.cells:
                for p in cel.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8)
    doc.save(str(docx))


def textos(N: dict) -> list[Path]:
    saidas = []
    if not MODELOS.exists():
        print("[artigo] sem modelos de texto em artigo/texto/modelos/ (textos não gerados)")
        return saidas
    for m in sorted(MODELOS.glob("*.md")):
        md = TXT / m.name
        md.write_text(renderizar(m.read_text(encoding="utf-8"), N, m.name), encoding="utf-8")
        saidas.append(md)
    for md in sorted(TXT.glob("*.md")):
        dx = md.with_suffix(".docx")
        md_para_docx(md, dx)
        saidas.append(dx)
    return saidas


def main():
    d = carregar()
    tabs = construir_tabelas(d)
    TAB.mkdir(parents=True, exist_ok=True)
    escrever_xlsx(tabs, TAB / "tabelas_artigo.xlsx")
    (TAB / "tabelas_artigo.md").write_text(
        "# Tabelas do artigo\n\nGerado por `scripts/rodada05_artigo.py`; números no padrão brasileiro.\n\n"
        + "\n".join(tabela_md(t) for t in tabs), encoding="utf-8")
    escrever_docx(tabs, TAB / "tabelas_artigo.docx")
    (RES / "tabelas_artigo_ponto_decimal.md").write_text(  # as mesmas tabelas na convenção do repositório (relatório)
        "# Tabelas do artigo — ponto decimal, sem separador de milhar (convenção do repositório)\n\nGerado por "
        "`scripts/rodada05_artigo.py`.\n\n" + "\n".join(tabela_md(t, ponto=True) for t in tabs), encoding="utf-8")
    figs = figuras(d)
    N = numeros_chave(d)
    escrever_numeros(N)
    txt = textos(N)
    print(f"[artigo] {len(tabs)} tabelas; {len(figs)} arquivos de figura; {len(N)} números-chave; {len(txt)} arquivos de texto")


if __name__ == "__main__":
    main()
